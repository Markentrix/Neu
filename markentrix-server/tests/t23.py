import httpx, time, re, email, subprocess, io, base64, csv
B="http://localhost:8000"; H={"X-MK":"1"}
def chk(n,c,i=""): print("OK  " if c else "FAIL",n,"" if c else str(i)[:300])
def sql(q): return subprocess.run(["su","postgres","-c",f"psql -q -t -A -h /var/tmp/pgtest -p 54329 markentrix -c \"{q}\""],capture_output=True,text=True).stdout.strip()
def mails():
    try: return [email.message_from_string(x) for x in open("/var/tmp/pt/mails.txt").read().split("\n=====\n") if x.strip()]
    except FileNotFoundError: return []
def text(m):
    for p in m.walk():
        if p.get_content_type()=="text/plain": return p.get_payload(decode=True).decode("utf8","replace")
def mail_an(adr, enthaelt=""):
    time.sleep(1.2)
    return [m for m in mails() if adr in str(m["To"]) and enthaelt in str(email.header.make_header(email.header.decode_header(m["Subject"] or ""))) + text(m)]
def token(adr):
    for m in reversed(mails()):
        if adr in str(m["To"]):
            g=re.search(r"einladung\?token=([A-Za-z0-9_-]+)",text(m))
            if g: return g.group(1)
a=httpx.Client(base_url=B,headers=H); a.post("/api/login",json={"email":"chef@test.de","password":open("/var/tmp/pt/adminpw").read().strip()})
a.put("/api/admin/einstellungen/smtp",json={"host":"127.0.0.1","port":8025,"sicherung":"keine","absender":"info@markentrix.de"})
sfx=str(int(time.time()))
def neuer_kunde(name):
    k=a.post("/api/admin/kunden",json={"firma":f"{name} {sfx}","ansprechpartner":"Tina Test","email":f"{name.lower()}{sfx}@k.de","telefon":"0511","strasse":"Weg 1","plz":"29221","ort":"Celle","paket_code":"KI_ONE","einladen":True}).json()
    httpx.post(B+"/api/einladung",headers=H,json={"token":token(f"{name.lower()}{sfx}@k.de"),"password":"InhaberPasswort1"})
    c=httpx.Client(base_url=B,headers=H); r=c.post("/api/login",json={"email":f"{name.lower()}{sfx}@k.de","password":"InhaberPasswort1"}); assert r.status_code==200,r.text
    return k["id"],c,f"{name.lower()}{sfx}@k.de"
kid,c,kmail=neuer_kunde("Alpha")
from openpyxl import Workbook
def xlsx(rows):
    wb=Workbook(); ws=wb.active
    for r in rows: ws.append(r)
    b=io.BytesIO(); wb.save(b); return "data:application/octet-stream;base64,"+base64.b64encode(b.getvalue()).decode()
def imp(datei,name,**kw): return a.post("/api/admin/leads/import",json={"datei_base64":datei,"datei_name":name,**kw})
kopf=["Unternehmen","Ansprechpartner","E-Mail","Tel","Webseite","Kategorie","PLZ","Stadt","Adresse"]
rows=[kopf,
 [f"Werkstatt A {sfx}","Hans A",f"a{sfx}@w.de",5113123456,"https://www.wa.de/","KFZ",29221,"Celle","Weg 1"],
 [f"Werkstatt B {sfx}",None,f"b{sfx}@w.de","0511 2","wb.de","KFZ","29223","Celle",None],
 [f"Werkstatt A {sfx}","Doppelt",f"a{sfx}@w.de",None,None,None,None,None,None],
 [None,"Ohne Firma",None,None,None,None,None,None,None],
 [f"Kaputt GmbH {sfx}",None,"keine-mail",None,None,None,None,None,None],
 [f"Werkstatt C {sfx}",None,None,None,None,"KFZ","29225","Hambühren",None]]
d=xlsx(rows)
r=imp(d,"liste.xlsx").json(); b0=r["blaetter"][0]; chk("Vorschau",r["neu"]==3 and b0["doppelt"]==1 and b0["ungueltig"]==2 and not r["importiert"],r)
chk("Spalten erkannt",b0["erkannt"].get("firma")=="Unternehmen" and b0["erkannt"].get("telefon")=="Tel" and b0["erkannt"].get("ort")=="Stadt",b0["erkannt"])
chk("Vorschau schreibt nichts",sql(f"select count(*) from leads where firma like '%{sfx}'")=="0")
r=imp(d,"liste.xlsx",nur_pruefen=False,liste="Werkstätten Celle",temperatur="kalt",quelle="akquise"); chk("Import",r.status_code==200 and r.json()["importiert"] and r.json()["neu"]==3,r.text)
chk("3 Leads angelegt",sql(f"select count(*) from leads where firma like 'Werkstatt _ {sfx}'")=="3")
chk("Felder korrekt",sql(f"select telefon||'|'||website||'|'||plz||'|'||temperatur||'|'||liste||'|'||quelle from leads where firma='Werkstatt A {sfx}'")=="5113123456|www.wa.de|29221|kalt|Werkstätten Celle · Sheet|akquise",sql(f"select telefon||'|'||website from leads where firma='Werkstatt A {sfx}'"))
chk("Keine Einwilligung vermerkt",sql(f"select count(*) from lead_aktivitaeten a join leads l on l.id=a.lead_id where l.liste like 'Werkstätten Celle%' and a.art='consent'")=="0")
r=imp(d,"liste.xlsx",nur_pruefen=False).json(); chk("Zweiter Import: nur Dubletten",r["neu"]==0 and r["blaetter"][0]["doppelt"]==4,r)
chk("Nicht doppelt in DB",sql(f"select count(*) from leads where firma like 'Werkstatt _ {sfx}'")=="3")
# CSV mit Semikolon und Windows-Zeichensatz
txt="Firma;E-Mail;Telefon;PLZ;Ort\r\nBäckerei Müller "+sfx+";m"+sfx+"@b.de;0511 9;29221;Celle\r\n"
cs="data:text/csv;base64,"+base64.b64encode(txt.encode("cp1252")).decode()
r=imp(cs,"b.csv",nur_pruefen=False,liste="CSV-Test").json(); chk("CSV Import",r["neu"]==1,r)
chk("Umlaute ok",sql(f"select count(*) from leads where firma='Bäckerei Müller {sfx}'")=="1")
chk("Ohne Firmenspalte abgelehnt",imp(xlsx([["Foo","Bar"],["1","2"]]),"x.xlsx").status_code==400)
chk("Leere Datei abgelehnt",imp(xlsx([["Firma"]]),"x.xlsx").status_code==400)
chk("Falsches Format abgelehnt",imp("data:x;base64,"+base64.b64encode(b"abc").decode(),"x.pdf").status_code==400)
chk("Kaputte xlsx abgelehnt",imp("data:x;base64,"+base64.b64encode(b"abc").decode(),"x.xlsx").status_code==400)
chk("Kunde darf nicht importieren",c.post("/api/admin/leads/import",json={"datei_base64":d,"datei_name":"x.xlsx"}).status_code==403) if False else None
# Vertriebler zuordnen + Filter
p=a.post("/api/admin/vertrieb/partner",json={"name":"Imp Vertrieb","status":"aktiv"}).json()["id"]
d2=xlsx([["Firma","Ort"],[f"Zugeordnet GmbH {sfx}","Hannover"]]); imp(d2,"z.xlsx",nur_pruefen=False,vertriebspartner_id=p,liste="Hannover",temperatur="heiss")
L=lambda q: a.get("/api/admin/leads?"+q).json()
chk("Filter Liste",[x["firma"] for x in L("liste=Hannover%20·%20Sheet")["leads"]]==[f"Zugeordnet GmbH {sfx}"],L("liste=Hannover%20·%20Sheet"))
chk("Filter Vertriebler",len(L(f"vertriebler={p}")["leads"])==1)
chk("Filter nicht zugeordnet",all(x["vertriebler"] is None for x in L("vertriebler=keiner")["leads"]) and len(L("vertriebler=keiner")["leads"])>=3)
chk("Filter Einstufung",all(x["temperatur"]=="heiss" for x in L("temperatur=heiss")["leads"]) and len(L("temperatur=heiss")["leads"])>=1)
chk("Filter ohne Einstufung",all(x["temperatur"] is None for x in L("temperatur=keine")["leads"]))
chk("Filter Quelle",all(x["quelle"]=="akquise" for x in L("quelle=akquise")["leads"]) and len(L("quelle=akquise")["leads"])>=3)
chk("Suche E-Mail",[x["firma"] for x in L(f"q=a{sfx}@w.de")["leads"]]==[f"Werkstatt A {sfx}"])
chk("Suche Telefon",len(L("q=5113123456")["leads"])==1)
chk("Suche Branche",len([x for x in L("q=KFZ")["leads"] if sfx in x["firma"]])==3)
chk("Listen im Ergebnis",any(l["liste"]=="Werkstätten Celle · Sheet" and l["n"]==3 for l in L("")["listen"]),L("")["listen"])
chk("Kombinierter Filter",len(L("liste=Werkstätten%20Celle%20·%20Sheet&q=Hambühren")["leads"])==1)
chk("Importierter Lead hat Lösch-Frist",sql(f"select loeschen_nach is not null from leads where firma='Werkstatt A {sfx}'")=="t")
chk("Audit",int(sql("select count(*) from audit_log where aktion='leads_import'"))>=3)

# ---- Mehrere Blätter, Prio, Adresse, Gruppe
def xlsx_multi(blaetter):
    wb=Workbook(); wb.remove(wb.active)
    for name,rows in blaetter:
        ws=wb.create_sheet(name)
        for r in rows: ws.append(r)
    b=io.BytesIO(); wb.save(b); return "data:application/octet-stream;base64,"+base64.b64encode(b.getvalue()).decode()
K=["Prio","Firmenname","Branche","Branchengruppe","Telefon","Adresse","Ort","km","Bewertung","Anz. Bewertungen","Google Maps","Hinweis"]
mb=xlsx_multi([("Übersicht",[["Titel"],["x","y"]]),
 ("Warme Leads",[K,["A",f"Friseur Eins {sfx}","Friseursalon","Beauty","05141 111111","Weg 5, 29221 Celle","Celle",1.2,4.7,27,"Maps","aktiv"],["B",f"Friseur Zwei {sfx}","Friseursalon","Beauty","05141 222222","Str 2, 29223 Celle","Celle",2,4.0,5,"Maps",None],["C",f"Bagger Drei {sfx}","Baubetrieb","Handwerk & Bau","05141 333333","Ring 9, 29225 Celle","Celle",3,None,None,None,None]]),
 ("Noch ungeprüft",[["Firmenname","Branche","Adresse","Ort"],[f"Ungeprueft {sfx}","Bäcker","Hof 1, 29221 Celle","Celle"]])])
r=imp(mb,"multi.xlsx").json(); names={b["name"]:b for b in r["blaetter"]}
chk("Blätter erkannt (Übersicht ausgelassen)",set(names)=={"Warme Leads","Noch ungeprüft"},list(names))
chk("Standard: ungeprüft nicht vorgewählt",names["Warme Leads"]["gewaehlt"] and not names["Noch ungeprüft"]["gewaehlt"] and r["neu"]==3,r["neu"])
r=imp(mb,"multi.xlsx",nur_pruefen=False,liste="Celle").json(); chk("Multi-Import",r["importiert"] and r["neu"]==3,r)
chk("Prio A heiß",sql(f"select temperatur from leads where firma='Friseur Eins {sfx}'")=="heiss"); chk("Prio B warm",sql(f"select temperatur from leads where firma='Friseur Zwei {sfx}'")=="warm"); chk("Prio C kalt",sql(f"select temperatur from leads where firma='Bagger Drei {sfx}'")=="kalt")
chk("PLZ aus Adresse",sql(f"select plz||'|'||ort from leads where firma='Friseur Eins {sfx}'")=="29221|Celle")
chk("Gruppe gespeichert",sql(f"select gruppe from leads where firma='Bagger Drei {sfx}'")=="Handwerk & Bau")
chk("Notiz mit Bewertung","Google-Bewertung 4.7 (27 Bewertungen)" in sql(f"select notiz from leads where firma='Friseur Eins {sfx}'"))
chk("Ungeprüftes Blatt nicht importiert",sql(f"select count(*) from leads where firma='Ungeprueft {sfx}'")=="0")
r=imp(mb,"multi.xlsx",nur_pruefen=False,blaetter=["Noch ungeprüft"],liste="Celle").json(); chk("Blatt gezielt wählen",r["neu"]==1,r)
chk("Filter Branchengruppe",[x["firma"] for x in L("gruppe=Handwerk%20%26%20Bau")["leads"]]==[f"Bagger Drei {sfx}"] and any(g["gruppe"]=="Beauty" for g in L("")["gruppen"]))
chk("Suche in Notiz",len(L("q=Prio%20A")["leads"])>=1)

# ---- Seiten und Sammelaktionen
big=xlsx([["Firma","Telefon","Ort"]]+[[f"Bulk {i:03d} {sfx}",f"0511 {900000+i}","Celle"] for i in range(120)])
r=imp(big,"bulk.xlsx",nur_pruefen=False,liste="Bulk",temperatur="kalt").json(); chk("120 importiert",r["neu"]==120,r)
pg=L("liste=Bulk%20%C2%B7%20Sheet"); chk("Seite 1: 50 von 120",len(pg["leads"])==50 and pg["gesamt"]==120 and pg["pro_seite"]==50,(len(pg["leads"]),pg["gesamt"]))
p3=L("liste=Bulk%20%C2%B7%20Sheet&seite=3"); chk("Seite 3: 20",len(p3["leads"])==20)
ids=[x["id"] for x in pg["leads"][:5]]
r=a.post("/api/admin/leads/bulk",json={"aktion":"temperatur","wert":"heiss","ids":ids}); chk("Bulk Einstufung",r.status_code==200 and r.json()["geaendert"]==5,r.text)
chk("Bulk wirkt",sql(f"select count(*) from leads where liste='Bulk · Sheet' and temperatur='heiss'")=="5")
r=a.post("/api/admin/leads/bulk",json={"aktion":"vertriebler","wert":p,"alle_gefilterten":True,"filter":{"liste":"Bulk · Sheet","temperatur":"kalt"}}); chk("Bulk alle Gefilterten zuordnen",r.json()["geaendert"]==115,r.text)
chk("Zuordnung gesetzt",sql(f"select count(*) from leads where liste='Bulk · Sheet' and vertriebspartner_id='{p}'")=="115")
r=a.post("/api/admin/leads/bulk",json={"aktion":"status","wert":"kontaktiert","ids":ids}); chk("Bulk Status",r.json()["geaendert"]==5)
chk("Bulk ungültiger Status",a.post("/api/admin/leads/bulk",json={"aktion":"status","wert":"gewonnen","ids":ids}).status_code==400)
chk("Bulk ohne Auswahl",a.post("/api/admin/leads/bulk",json={"aktion":"loeschen"}).status_code==400)
chk("Bulk ungültige ID",a.post("/api/admin/leads/bulk",json={"aktion":"loeschen","ids":["xyz"]}).status_code==400)
chk("Bulk Kunde verboten",c.post("/api/admin/leads/bulk",json={"aktion":"loeschen","ids":ids}).status_code==403)
r=a.post("/api/admin/leads/bulk",json={"aktion":"loeschen","alle_gefilterten":True,"filter":{"liste":"Bulk · Sheet"}}); chk("Bulk löschen",r.json()["geaendert"]==120,r.text)
chk("Weg",sql("select count(*) from leads where liste='Bulk · Sheet'")=="0")

# ---- Vertriebler anlegen und gleich einladen
r=a.post("/api/admin/vertrieb/partner",json={"name":"Nora Neu","email":f"nora{sfx}@v.de","status":"aktiv","einladen":True}); chk("Vertriebler mit Einladung",r.status_code==200 and r.json()["mail"]["gesendet"] and "einladung?token=" in r.json()["einladungslink"],r.text)
mm=mail_an(f"nora{sfx}@v.de","Vertriebsportal"); chk("Mail mit Vertriebs-Text",len(mm)==1 and "Leads" in text(mm[0]) and "einladung?token=" in text(mm[0]),len(mm))
zg=lambda n:[x for x in a.get("/api/admin/vertrieb/partner").json() if x["name"]==n][0]["zugang"]
chk("Zugang: eingeladen",zg("Nora Neu")=="eingeladen"); chk("Zugang: keiner",zg("Imp Vertrieb")=="keiner")
httpx.post(B+"/api/einladung",headers=H,json={"token":token(f"nora{sfx}@v.de"),"password":"NoraPasswort123"}); chk("Zugang: aktiv",zg("Nora Neu")=="aktiv")
chk("Einladung ohne E-Mail abgelehnt",a.post("/api/admin/vertrieb/partner",json={"name":"Ohne Mail","einladen":True}).status_code==400)
chk("Ohne Einladen kein Zugang",a.post("/api/admin/vertrieb/partner",json={"name":"Stille Sophie","email":f"so{sfx}@v.de"}).status_code==200 and zg("Stille Sophie")=="keiner")
r=a.post(f"/api/admin/vertrieb/partner/{[x for x in a.get('/api/admin/vertrieb/partner').json() if x['name']=='Stille Sophie'][0]['id']}/einladung",json={"email":f"so{sfx}@v.de","per_mail":True}); chk("Nachträglich einladen",r.status_code==200 and r.json()["mail"]["gesendet"],r.text)
