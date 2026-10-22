import httpx, time, re, email, subprocess, io
B="http://localhost:8000"; H={"X-MK":"1"}
def chk(n,c,i=""): print("OK  " if c else "FAIL",n,"" if c else str(i)[:300])
def sql(q): return subprocess.run(["su","postgres","-c",f"psql -q -t -A -h /var/tmp/pgtest -p 54329 markentrix -c \"{q}\""],capture_output=True,text=True).stdout.strip()
def mails():
    try: return [email.message_from_string(x) for x in open("/var/tmp/pt/mails.txt").read().split("\n=====\n") if x.strip()]
    except FileNotFoundError: return []
def text(m):
    for p in m.walk():
        if p.get_content_type()=="text/plain": return p.get_payload(decode=True).decode("utf8","replace")
def mail_an(adr, enthaelt=""):
    time.sleep(1.2)
    return [m for m in mails() if adr in str(m["To"]) and enthaelt in str(email.header.make_header(email.header.decode_header(m["Subject"] or ""))) + text(m)]
def token(adr):
    for m in reversed(mails()):
        if adr in str(m["To"]):
            g=re.search(r"einladung\?token=([A-Za-z0-9_-]+)",text(m))
            if g: return g.group(1)
a=httpx.Client(base_url=B,headers=H); a.post("/api/login",json={"email":"chef@test.de","password":open("/var/tmp/pt/adminpw").read().strip()})
a.put("/api/admin/einstellungen/smtp",json={"host":"127.0.0.1","port":8025,"sicherung":"keine","absender":"info@markentrix.de"})
sfx=str(int(time.time())+305)
a.put("/api/admin/einstellungen/firma",json={"name":"Markentrix","inhaber":"Christian Wien","strasse":"Teststr. 1","plz":"29221","ort":"Celle","email":"info@markentrix.de","kleinunternehmer":True,"mwst_satz":19})
def neuer_kunde(name):
    k=a.post("/api/admin/kunden",json={"firma":f"{name} {sfx}","ansprechpartner":"Tina Test","email":f"{name.lower()}{sfx}@k.de","telefon":"0511","strasse":"Weg 1","plz":"29221","ort":"Celle","paket_code":"KI_ONE","einladen":True}).json()
    httpx.post(B+"/api/einladung",headers=H,json={"token":token(f"{name.lower()}{sfx}@k.de"),"password":"InhaberPasswort1"})
    c=httpx.Client(base_url=B,headers=H); r=c.post("/api/login",json={"email":f"{name.lower()}{sfx}@k.de","password":"InhaberPasswort1"}); assert r.status_code==200,r.text
    return k["id"],c,f"{name.lower()}{sfx}@k.de"


import datetime, json, os, sys, zoneinfo
TZ=zoneinfo.ZoneInfo("Europe/Berlin")
heute=datetime.date.today()
kA,cA,mA=neuer_kunde("Absa"); kB,cB,mB=neuer_kunde("Absb")
sys.path.insert(0,"/home/claude/markentrix-server/portal")
os.environ.update(DATABASE_URL="postgresql://mk_app:apppw@localhost:54329/markentrix?host=/var/tmp/pgtest",PORTAL_SECRET="testgeheimnis123",PORTAL_JOBS="0")
from app import db, jobs, webbetrieb
from app.db import tx
db.open_pool()
def cfg():
    with tx("system") as cur: return jobs.einstellungen_lesen(cur)
import openpyxl
heute=datetime.datetime.now(TZ).date()
tage=[(heute+datetime.timedelta(days=i)).isoformat() for i in (3,4,5,6,7,8)]
dom=f"absage{sfx}.de"
sql(f"INSERT INTO websites (kunden_id, domain, art, status) VALUES ('{kA}','{dom}','live','live')")
W=httpx.Client(base_url=B,headers={"Host":dom}); API="/siteapi/api"
def post(pfad,body): return W.post(API+pfad,headers={"Origin":"https://"+dom,"Content-Type":"application/json","User-Agent":"Mozilla/5.0 Test"},content=json.dumps(body).encode())
D="/api/kunde/website-dienste"
cA.put(D+"/mail",json={"aktiv":True,"empfaenger":f"chef@{dom}","betreff":"x","name":"Salon"})
r=cA.put(D+"/kalender",json={"aktiv":True,"dauer":30,"tage":30,"von":"08:00","bis":"18:00","wochentage":[1,2,3,4,5,6,7],"vorlauf_std":0,"ort":"Salon","empfaenger":"","zeiten":"08:00, 09:00, 10:00, 11:00","kapazitaet":0,"freigabe":False}); chk("Kalender: Sofortbuchung",r.status_code==200,r.text)
r=cA.put(D+"/paypal",json={"aktiv":True,"waehrung":"EUR","me_link":"https://paypal.me/salona"}); chk("PayPal.Me eintragen",r.status_code==200,r.text)

# ---- Leistungen mit Zusatzfragen
L=lambda **kw:{**{"dauer_min":60,"puffer_min":0},**kw}
r=cA.post(D+"/leistungen",json=L(name="Schnitt",preis="45 €",dauer_min=30,felder=[{"label":"Haarlänge","pflicht":True},{"label":"Haarzustand","pflicht":False}])); chk("Leistung mit Zusatzfragen",r.status_code==200,r.text)
r=cA.post(D+"/leistungen",json=L(name="Färben",preis="60 €",dauer_min=30)); chk("Zweite Leistung",r.status_code==200,r.text)
r=cA.post(D+"/leistungen",json=L(name="Balayage",preis="180 €",dauer_min=60,freigabe=True,anzahlung_art="betrag",anzahlung_wert=50)); chk("Leistung mit Anzahlung",r.status_code==200,r.text)
r=cA.post(D+"/leistungen",json=L(name="Falsch",felder=[{"label":"Name"}])); chk("Reservierter Feldname abgelehnt",r.status_code==400,r.text)
r=cA.post(D+"/leistungen",json=L(name="Doppelt",felder=[{"label":"Farbe"},{"label":"farbe"}])); chk("Doppelte Zusatzfrage abgelehnt",r.status_code==400,r.text)
ls=W.get(API+"/leistungen",headers={"Origin":"https://"+dom}).json()["leistungen"]; sc={x["key"]:x for x in ls}
chk("API liefert Zusatzfragen",[f["key"] for f in sc["schnitt"]["felder"]]==["haarlaenge","haarzustand"] and sc["schnitt"]["felder"][0]["pflicht"] is True,sc.get("schnitt"))
sid=sql(f"select site_id from website_dienste where kunden_id='{kA}'")
def buchen(tag,uhr,leistung,mail,**extra):
    return post("/anfrage",{"typ":"termin","name":"Peter Besucher","email":mail,"telefon":"0511 123456","datum":tag,"uhrzeit":uhr,"leistung":leistung,"einwilligung":"ja",**extra})
r=buchen(tage[0],"08:00","schnitt","p0@besucher.de"); chk("Pflichtfeld fehlt: 422",r.status_code==422 and "haarlaenge" in r.json().get("errors",{}),r.text)
r=buchen(tage[0],"08:00","schnitt","p0@besucher.de",haarlaenge="lang",haarzustand="trocken"); chk("Mit Pflichtfeld gebucht",r.status_code==200,r.text)
m=mail_an(f"chef@{dom}"); chk("Besitzer-Mail zeigt Beschriftung der Zusatzfrage",bool(m) and "Haarlänge: lang" in text(m[-1]),[text(x)[:200] for x in m[-1:]])
# ---- mehrere Leistungen
r=buchen(tage[1],"08:00","schnitt,faerben","p1@besucher.de",haarlaenge="kurz"); chk("Zwei Leistungen in einem Termin",r.status_code==200,r.text)
row=sql(f"select (extract(epoch from ende-start)/60)::int || '|' || preis_betrag from website_termine where email='p1@besucher.de'"); chk("Dauer und Preis addiert",row=="60|105.00",row)
r=buchen(tage[1],"08:30","schnitt,faerben","p1b@besucher.de",haarlaenge="kurz"); chk("Zweite Buchung überschneidet sich nicht",r.status_code==422,r.text)
r=buchen(tage[1],"09:00","schnitt,gibtsnicht","p1c@besucher.de",haarlaenge="kurz"); chk("Unbekannte Leistung in Liste abgelehnt",r.status_code==422,r.text)
s=W.get(API+"/slots",params={"tag":tage[1],"leistung":"schnitt,faerben"},headers={"Origin":"https://"+dom}).json(); fr={x["zeit"]:x["frei"] for x in s["slots"]}
chk("Slots für beide Leistungen zusammen",fr.get("08:00") is False and fr.get("09:00") is True,fr)
# ---- Absage mit Frist, Erstattung
r=buchen(tage[2],"08:00","balayage","p2@besucher.de"); chk("Freigabe-Leistung wird Wunsch",r.status_code==200,r.text)
bid=sql(f"select id from website_termine where email='p2@besucher.de'")
r=cA.post(D+f"/buchungen/{bid}/bestaetigen",json={"anzahlung":""}); chk("Bestätigen",r.status_code==200,r.text)
r=cA.post(D+f"/buchungen/{bid}/anzahlung-erhalten",json={}); chk("Anzahlung erhalten",r.status_code==200,r.text)
tok=sql(f"select token from website_termine where id='{bid}'")
P=httpx.Client(base_url=B)
r=P.get(f"/site/{sid}/termin/absagen",params={"t":tok}); chk("Absageseite nennt kostenlose Absage und Erstattung",r.status_code==200 and "kostenlose Absage ist noch möglich" in r.text and "erstattet" in r.text,r.text[:400])
chk("Link absagen allein ändert nichts",sql(f"select status from website_termine where id='{bid}'")=="gebucht")
r=P.post(f"/site/{sid}/termin/absagen",data={"t":tok}); chk("Absage per Klick",r.status_code==200 and "erstattet" in r.text,r.text[:300])
chk("Als kostenlos gespeichert",sql(f"select storno_kostenlos::text || abgesagt_von from website_termine where id='{bid}'")=="truebesucher")
m=mail_an(f"chef@{dom}","abgesagt"); chk("Besitzer bekommt Erstattungshinweis",bool(m) and "zurückzahlen" in text(m[-1]),[text(x)[:300] for x in m[-1:]])
er=cA.get(D+"/erstattungen").json()["erstattungen"]; chk("Erstattung offen gelistet",len(er)==1 and er[0]["betrag"]=="50.00",er)
r=cA.post(D+f"/buchungen/{bid}/erstattet",json={}); chk("Als erstattet markieren",r.status_code==200,r.text)
chk("Danach nichts mehr offen",cA.get(D+"/erstattungen").json()["erstattungen"]==[])
# zu spät
r=buchen(tage[3],"08:00","schnitt","p3@besucher.de",haarlaenge="lang"); b3=sql(f"select id from website_termine where email='p3@besucher.de'"); t3=sql(f"select token from website_termine where id='{b3}'")
sql(f"update website_termine set start=now()+interval '5 hours', ende=now()+interval '5.5 hours', anzahlung_betrag=20, anzahlung='20.00', anzahlung_erhalten_am=now() where id='{b3}'")
r=P.get(f"/site/{sid}/termin/absagen",params={"t":t3}); chk("Zu spät: Seite nennt vollen Betrag",r.status_code==200 and "volle Betrag" in r.text and "nicht erstattet" in r.text,r.text[:500])
chk("Zu spät: kein Umbuchen-Angebot","verschieben" not in r.text)
r=P.get(f"/site/{sid}/termin/verschieben",params={"t":t3}); chk("Zu spät: Umbuchen gesperrt","möglich" in r.text and 'name="neu"' not in r.text,r.text[:300])
r=P.post(f"/site/{sid}/termin/absagen",data={"t":t3}); chk("Zu späte Absage geht trotzdem durch",r.status_code==200 and "volle Betrag" in r.text,r.text[:300])
chk("Als nicht kostenlos gespeichert",sql(f"select storno_kostenlos::text from website_termine where id='{b3}'")=="false")
chk("Keine Erstattung offen",cA.get(D+"/erstattungen").json()["erstattungen"]==[])
m=mail_an(f"chef@{dom}","abgesagt"); chk("Besitzer-Mail: voller Betrag fällig",any("volle Betrag ist fällig" in text(x) for x in m),[text(x)[:200] for x in m[-2:]])
# Betrieb sagt ab
r=buchen(tage[4],"08:00","schnitt","p4@besucher.de",haarlaenge="lang"); b4=sql(f"select id from website_termine where email='p4@besucher.de'")
cA.post(D+f"/buchungen/{b4}/anzahlung-erhalten",json={}); sql(f"update website_termine set anzahlung_betrag=15, anzahlung='15.00' where id='{b4}'")
r=cA.post(D+f"/buchungen/{b4}/ablehnen",json={"grund":"Krank"}); chk("Betrieb sagt ab: Erstattung nötig",r.status_code==200 and r.json().get("erstatten") is True,r.text)
m=mail_an("p4@besucher.de"); chk("Besucher: Anzahlung wird erstattet",bool(m) and "erstatten wir" in text(m[-1]),[text(x)[:300] for x in m[-1:]])
chk("Liste zeigt Betriebs-Absage",[e["betrag"] for e in cA.get(D+"/erstattungen").json()["erstattungen"]]==["15.00"])
# ---- Umbuchen
r=buchen(tage[5],"08:00","schnitt","p5@besucher.de",haarlaenge="lang"); b5=sql(f"select id from website_termine where email='p5@besucher.de'"); t5=sql(f"select token from website_termine where id='{b5}'")
sql(f"update website_termine set erinnerung=2 where id='{b5}'")
r=P.get(f"/site/{sid}/termin/verschieben",params={"t":t5}); neu=re.findall(r'name="neu" value="([^"]+)"',r.text); chk("Verschieben-Seite bietet Zeiten",r.status_code==200 and len(neu)>3,r.text[:300])
ziel=neu[-1].replace("&#x27;","'")
r=P.post(f"/site/{sid}/termin/verschieben",data={"t":t5,"neu":ziel}); chk("Termin verschoben",r.status_code==200 and "jetzt am" in r.text,r.text[:300])
row=sql(f"select (start at time zone 'utc')::text || '|' || erinnerung || '|' || (verschoben_am is not null)::text from website_termine where id='{b5}'"); chk("Neue Zeit gespeichert, Erinnerung zurückgesetzt",row.endswith("|0|true") and ziel[:16].replace("T"," ") in row.replace("+00:00",""),(row,ziel))
chk("Kalendereintrag mitgezogen",sql(f"select to_char(t.beginn at time zone 'utc','YYYY-MM-DD HH24:MI') from termine t join website_termine w on w.termin_id=t.id where w.id='{b5}'")==sql(f"select to_char(start at time zone 'utc','YYYY-MM-DD HH24:MI') from website_termine where id='{b5}'"))
m=mail_an("p5@besucher.de","neuer Termin"); chk("Besucher: Mail mit Kalender-Anhang",bool(m) and any(p.get_content_type()=="text/calendar" for p in m[-1].walk()),len(m))
m=mail_an(f"chef@{dom}","verschoben"); chk("Besitzer wird informiert",bool(m))
r=P.post(f"/site/{sid}/termin/verschieben",data={"t":t5,"neu":"2020-01-01T10:00:00+00:00"}); chk("Nicht freie Zeit abgelehnt","nicht mehr frei" in r.text,r.text[:200])
r=P.post(f"/site/{sid}/termin/verschieben",data={"t":t5,"neu":"quatsch"}); chk("Ungültige Zeit abgelehnt","gibt es nicht" in r.text,r.text[:200])
# ---- Quittung
r=cA.get(D+f"/buchungen/{bid}/quittung.pdf"); chk("Quittung auch für abgesagten Termin mit Anzahlung abrufbar",r.status_code==200 and r.content[:4]==b"%PDF",r.status_code)
r=buchen(tage[5],"09:00","balayage","p6@besucher.de"); b6=sql(f"select id from website_termine where email='p6@besucher.de'")
r=cA.get(D+f"/buchungen/{b6}/quittung.pdf"); chk("Ohne erhaltene Anzahlung keine Quittung",r.status_code==409,r.status_code)
cA.post(D+f"/buchungen/{b6}/anzahlung-erhalten",json={})
r=cA.get(D+f"/buchungen/{b6}/quittung.pdf"); chk("Quittung als PDF",r.status_code==200 and r.content[:4]==b"%PDF" and "inline" in r.headers.get("content-disposition",""),r.status_code)
r=cA.post(D+f"/buchungen/{b6}/quittung-senden",json={}); chk("Quittung senden",r.status_code==200,r.text)
m=mail_an("p6@besucher.de","Quittung"); chk("Mail mit PDF-Anhang",bool(m) and any(p.get_content_type()=="application/pdf" for p in m[-1].walk()),len(m))
chk("Zeitpunkt gespeichert",sql(f"select (quittung_am is not null)::text from website_termine where id='{b6}'")=="true")
r=cB.get(D+f"/buchungen/{b6}/quittung.pdf"); chk("Fremder Betrieb kommt nicht an die Quittung",r.status_code in (404,409),r.status_code)
# ---- Excel
r=cA.get(D+"/uebersicht.xlsx?zeitraum=monat"); wb=openpyxl.load_workbook(io.BytesIO(r.content)); ws=wb["Termine"]
kopf=[c.value for c in ws[1]]; stati=[row[kopf.index("Status")].value for row in ws.iter_rows(min_row=2)]; er_sp=[row[kopf.index("Erstattung")].value for row in ws.iter_rows(min_row=2)]
r2=cA.get(D+"/uebersicht.xlsx?zeitraum=eigen&von="+heute.isoformat()+"&bis="+(heute+datetime.timedelta(days=20)).isoformat()); ws2=openpyxl.load_workbook(io.BytesIO(r2.content))["Termine"]
st2=[row[kopf.index("Status")].value for row in ws2.iter_rows(min_row=2)]; er2=[row[kopf.index("Erstattung")].value for row in ws2.iter_rows(min_row=2)]
chk("Excel: zu spät abgesagt und Erstattungsspalte",any("zu spät" in str(x) for x in st2+stati) and ("erstattet" in er2+er_sp or "offen" in er2+er_sp),(st2,er2))
print("fertig")
