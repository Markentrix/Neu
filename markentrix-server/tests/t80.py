import httpx, time, re, email, subprocess, io
B="http://localhost:8000"; H={"X-MK":"1"}
def chk(n,c,i=""): print("OK  " if c else "FAIL",n,"" if c else str(i)[:300])
def sql(q): return subprocess.run(["su","postgres","-c",f"psql -q -t -A -h /var/tmp/pgtest -p 54329 markentrix -c \"{q}\""],capture_output=True,text=True).stdout.strip()
def mails():
    try: return [email.message_from_string(x) for x in open("/var/tmp/pt/mails.txt").read().split("\n=====\n") if x.strip()]
    except FileNotFoundError: return []
def text(m):
    for p in m.walk():
        if p.get_content_type()=="text/plain": return p.get_payload(decode=True).decode("utf8","replace")
def mail_an(adr, enthaelt=""):
    time.sleep(1.2)
    return [m for m in mails() if adr in str(m["To"]) and enthaelt in str(email.header.make_header(email.header.decode_header(m["Subject"] or ""))) + text(m)]
def token(adr):
    for m in reversed(mails()):
        if adr in str(m["To"]):
            g=re.search(r"einladung\?token=([A-Za-z0-9_-]+)",text(m))
            if g: return g.group(1)
a=httpx.Client(base_url=B,headers=H); a.post("/api/login",json={"email":"chef@test.de","password":open("/var/tmp/pt/adminpw").read().strip()})
a.put("/api/admin/einstellungen/smtp",json={"host":"127.0.0.1","port":8025,"sicherung":"keine","absender":"info@markentrix.de"})
sfx=str(int(time.time())+297)
a.put("/api/admin/einstellungen/firma",json={"name":"Markentrix","inhaber":"Christian Wien","strasse":"Teststr. 1","plz":"29221","ort":"Celle","email":"info@markentrix.de","kleinunternehmer":True,"mwst_satz":19})
def neuer_kunde(name):
    k=a.post("/api/admin/kunden",json={"firma":f"{name} {sfx}","ansprechpartner":"Tina Test","email":f"{name.lower()}{sfx}@k.de","telefon":"0511","strasse":"Weg 1","plz":"29221","ort":"Celle","paket_code":"KI_ONE","einladen":True}).json()
    httpx.post(B+"/api/einladung",headers=H,json={"token":token(f"{name.lower()}{sfx}@k.de"),"password":"InhaberPasswort1"})
    c=httpx.Client(base_url=B,headers=H); r=c.post("/api/login",json={"email":f"{name.lower()}{sfx}@k.de","password":"InhaberPasswort1"}); assert r.status_code==200,r.text
    return k["id"],c,f"{name.lower()}{sfx}@k.de"


import datetime, json, os, sys, zoneinfo
TZ=zoneinfo.ZoneInfo("Europe/Berlin")
heute=datetime.date.today()
kA,cA,mA=neuer_kunde("Anza"); kB,cB,mB=neuer_kunde("Anzb")
sys.path.insert(0,"/home/claude/markentrix-server/portal")
os.environ.update(DATABASE_URL="postgresql://mk_app:apppw@localhost:54329/markentrix?host=/var/tmp/pgtest",PORTAL_SECRET="testgeheimnis123",PORTAL_JOBS="0")
from app import db, jobs, webbetrieb
from app.db import tx
db.open_pool()
def cfg():
    with tx("system") as cur: return jobs.einstellungen_lesen(cur)
import openpyxl
heute=datetime.datetime.now(TZ).date()
tage=[(heute+datetime.timedelta(days=i)).isoformat() for i in (3,4,5,6,7,8)]
dom=f"anzahlung{sfx}.de"
sql(f"INSERT INTO websites (kunden_id, domain, art, status) VALUES ('{kA}','{dom}','live','live')")
W=httpx.Client(base_url=B,headers={"Host":dom}); API="/siteapi/api"
def post(pfad,body): return W.post(API+pfad,headers={"Origin":"https://"+dom,"Content-Type":"application/json","User-Agent":"Mozilla/5.0 Test"},content=json.dumps(body).encode())
D="/api/kunde/website-dienste"
cA.put(D+"/mail",json={"aktiv":True,"empfaenger":f"chef@{dom}","betreff":"x","name":"Salon"})
r=cA.put(D+"/kalender",json={"aktiv":True,"dauer":30,"tage":30,"von":"08:00","bis":"18:00","wochentage":[1,2,3,4,5,6,7],"vorlauf_std":0,"ort":"Salon","empfaenger":"","zeiten":"08:00, 09:00, 10:00, 11:00","kapazitaet":0,"freigabe":False}); chk("Kalender: Sofortbuchung",r.status_code==200,r.text)
r=cA.put(D+"/paypal",json={"aktiv":True,"waehrung":"EUR","me_link":"https://paypal.me/salona"}); chk("PayPal.Me eintragen",r.status_code==200,r.text)
# ---- Leistungen mit Anzahlung
L=lambda **kw:{**{"dauer_min":60,"puffer_min":0},**kw}
r=cA.post(D+"/leistungen",json=L(name="Prozent ohne Preis",anzahlung_art="prozent",anzahlung_wert=20)); chk("Prozent ohne Preis abgelehnt",r.status_code==400,r.text)
r=cA.post(D+"/leistungen",json=L(name="Zu viel Prozent",preis="100 €",anzahlung_art="prozent",anzahlung_wert=150)); chk("Über 100 Prozent abgelehnt",r.status_code==400,r.text)
r=cA.post(D+"/leistungen",json=L(name="Betrag fehlt",anzahlung_art="betrag")); chk("Anzahlung ohne Zahl abgelehnt",r.status_code==400,r.text)
cA.post(D+"/leistungen",json=L(name="Schnitt",preis="45 €",dauer_min=30))
cA.post(D+"/leistungen",json=L(name="Balayage",preis="ab 180 €",dauer_min=60,freigabe=True,anzahlung_art="betrag",anzahlung_wert=50))
cA.post(D+"/leistungen",json=L(name="Politur",preis="200 €",dauer_min=60,anzahlung_art="prozent",anzahlung_wert=25))
t=cA.get(D+"/team").json(); ls={x["name"]:x for x in t["leistungen"]}
chk("Berechnete Anzahlungen",ls["Balayage"]["anzahlung"]=="50.00" and ls["Politur"]["anzahlung"]=="50.00" and ls["Schnitt"]["anzahlung"]=="",ls)
chk("Einstellungen im Team-Ergebnis",t["einstellung"]["storno_std"]==24 and t["einstellung"]["frist_std"]==0 and t["einstellung"]["paypal_bereit"] is True,t["einstellung"])
pl=W.get(API+"/leistungen").json()["leistungen"]; pk={x["key"]:x for x in pl}
chk("Schnittstelle: freigabe und anzahlung je Leistung",pk["balayage"]["freigabe"] is True and pk["balayage"]["anzahlung"]=="50.00" and pk["schnitt"]["freigabe"] is False and pk["schnitt"]["anzahlung"]=="" and pk["politur"]["anzahlung"]=="50.00",pl)
w=sql(f"select wert from wissensbasis where kunden_id='{kA}' and bereich='leistungen' and quelle='system'")
chk("Wissensdatenbank bekommt Leistungen automatisch",("Balayage" in w and "Anzahlung 50,00 €" in w and "Dauer etwa 30 Minuten" in w and "24 Stunden" in w),w)
v=W.get(API+"/verfuegbarkeit",params={"leistung":"balayage"}).json(); chk("selfBooking false bei Leistung mit Freigabe",v["selfBooking"] is False,v)
v=W.get(API+"/verfuegbarkeit",params={"leistung":"schnitt"}).json(); chk("selfBooking true bei normaler Leistung",v["selfBooking"] is True,v)
# ---- Buchungen
T=lambda tag,uhr,leistung,mail:{"typ":"termin","name":"Mia Kunde","email":mail,"telefon":"0511 1","datum":tag,"uhrzeit":uhr,"leistung":leistung,"einwilligung":"ja","haarlaenge":"lang","haarzustand":"normal"}
r=post("/anfrage",T(tage[0],"08:00","schnitt",f"schnitt{sfx}@besucher.de")); chk("Schnitt gebucht",r.status_code==200,r.text); r1=r.json()["ref"]
chk("Schnitt: sofort gebucht, keine Anzahlung",sql(f"select status||coalesce(anzahlung,'-') from website_termine where ref='{r1}'")=="gebucht-")
m=mail_an(f"schnitt{sfx}@besucher.de","Termin"); chk("Schnitt: Mail ohne Zahlungslink",len(m)>=1 and "paypal.me" not in text(m[0]),[text(x) for x in m])
r=post("/anfrage",T(tage[1],"09:00","politur",f"politur{sfx}@besucher.de")); chk("Politur gebucht",r.status_code==200,r.text); r2=r.json()["ref"]
chk("Politur: gebucht mit Anzahlung 50",sql(f"select status||anzahlung||anzahlung_betrag||preis_betrag from website_termine where ref='{r2}'")=="gebucht50.0050.00200.00")
m=mail_an(f"politur{sfx}@besucher.de","paypal.me"); chk("Politur: Zahlungslink sofort in der Mail, mit Absageregel",len(m)==1 and "https://paypal.me/salona/50.00EUR" in text(m[0]) and "24 Stunden" in text(m[0]) and "erstattet" in text(m[0]),[text(x) for x in m])
r=post("/anfrage",T(tage[2],"10:00","balayage",f"balayage{sfx}@besucher.de")); chk("Balayage angefragt",r.status_code==200,r.text); r3=r.json()["ref"]
chk("Balayage: nur vorgemerkt (Freigabe der Leistung)",sql(f"select status||anzahlung from website_termine where ref='{r3}'")=="reserviert50.00")
m=mail_an(f"balayage{sfx}@besucher.de","Eingangsbestätigung"); chk("Balayage: Zahlungslink schon in der Eingangsbestätigung",len(m)==1 and "paypal.me/salona/50.00EUR" in text(m[0]),[text(x) for x in m])
m=mail_an(f"chef@{dom}",r3); chk("Inhaber wird auf Zahlungseingang hingewiesen",len(m)==1 and "Anzahlung erhalten" in text(m[0]) and "Haarlänge: lang" in text(m[0]),[text(x) for x in m])
b=cA.get(D+"/buchungen").json()["buchungen"]; bm={x["ref"]:x for x in b}
chk("Portal zeigt Anzahlung, Preis und offenen Zahlungsstand",bm[r2]["anzahlung"]=="50.00" and bm[r2]["preis"]=="200.00" and bm[r2]["anzahlung_erhalten_am"] is None,bm[r2])
r=cA.post(D+f"/buchungen/{bm[r2]['id']}/bezahlt"); chk("Alles bezahlt geht bei bestätigtem Termin",r.status_code==200,r.text)
chk("Bezahlt setzt auch Anzahlung erhalten",sql(f"select (anzahlung_erhalten_am is not null)::text||(bezahlt_am is not null)::text from website_termine where ref='{r2}'")=="truetrue")
r=cA.post(D+f"/buchungen/{bm[r3]['id']}/bezahlt"); chk("Wartender Termin: Alles bezahlt abgelehnt",r.status_code==409,r.text)
r=cA.post(D+f"/buchungen/{bm[r3]['id']}/anzahlung-erhalten"); chk("Anzahlung erhalten markieren",r.status_code==200 and sql(f"select (anzahlung_erhalten_am is not null)::text from website_termine where ref='{r3}'")=="true",r.text)
r=cB.post(D+f"/buchungen/{bm[r3]['id']}/anzahlung-erhalten"); chk("Anderer Kunde kann nichts markieren",r.status_code==409,r.text)
r=cA.post(D+f"/buchungen/{bm[r3]['id']}/bestaetigen",json={"anzahlung":""}); chk("Bestätigen mit vorbelegter Anzahlung",r.status_code==200,r.text)
m=mail_an(f"balayage{sfx}@besucher.de","bestätigt"); chk("Bestätigung mit Zahlungslink über 50 €",len(m)>=1 and "paypal.me/salona/50.00EUR" in text(m[-1]),[text(x) for x in m])
# ---- Frist
r=cA.put(D+"/anzahlung",json={"frist_std":24,"storno_std":12}); chk("Frist und Absageregel speichern",r.status_code==200,r.text)
r=post("/anfrage",T(tage[3],"08:00","politur",f"frist{sfx}@besucher.de")); r4=r.json()["ref"]
r=post("/anfrage",T(tage[3],"09:00","politur",f"fristok{sfx}@besucher.de")); r5=r.json()["ref"]
m=mail_an(f"frist{sfx}@besucher.de","paypal.me"); chk("Mail nennt Frist und neue Absageregel",len(m)==1 and "innerhalb von 24 Stunden" in text(m[0]) and "bis 12 Stunden" in text(m[0]),[text(x) for x in m])
bm=cA.get(D+"/buchungen").json()["buchungen"]; idok=[x["id"] for x in bm if x["ref"]==r5][0]; cA.post(D+f"/buchungen/{idok}/anzahlung-erhalten")
sql(f"update website_termine set erstellt_am=now()-interval '30 hours' where ref in ('{r4}','{r5}')")
print("Frist-Lauf:",webbetrieb.anzahlung_frist(cfg())); time.sleep(1.2)
chk("Nach der Frist ohne Zahlung: Termin freigegeben",sql(f"select status from website_termine where ref='{r4}'")=="abgesagt")
chk("Bezahlter Termin bleibt",sql(f"select status from website_termine where ref='{r5}'")=="gebucht")
chk("Besucher wird informiert",len(mail_an(f"frist{sfx}@besucher.de","freigegeben"))>=1)
chk("Zeit ist wieder buchbar",post("/anfrage",T(tage[3],"08:00","politur",f"neu{sfx}@besucher.de")).status_code==200)
# ---- Excel
r=cA.get(D+f"/uebersicht.xlsx?zeitraum=eigen&von={heute.isoformat()}&bis={(heute+datetime.timedelta(days=30)).isoformat()}"); chk("Excel-Datei kommt",r.status_code==200 and r.content[:2]==b"PK" and "spreadsheetml" in r.headers["content-type"],r.status_code)
wb=openpyxl.load_workbook(io.BytesIO(r.content)); chk("Drei Blätter",wb.sheetnames==["Termine","Zusammenfassung","Gutscheine"],wb.sheetnames)
ws=wb["Termine"]; zeilen=list(ws.iter_rows(min_row=2,values_only=True)); refs={z[13]:z for z in zeilen}
chk("Termine stehen mit Preis, Anzahlung und Zahlungsstand drin",len(zeilen)>=5 and refs[r2][8]==200 and refs[r2][9]==50 and refs[r2][11] is not None and refs[r2][12]==0,zeilen[:3])
chk("Abgesagter Termin ist als abgesagt markiert",refs[r4][7]=="abgesagt")
chk("Offener Rest wird gerechnet (Anzahlung erhalten, nicht komplett bezahlt)",refs[r5][12]==150,refs[r5])
zs={z[0]:z[1] for z in wb["Zusammenfassung"].iter_rows(min_row=4,max_row=11,values_only=True)}
chk("Zusammenfassung rechnet",zs["Abgesagt"]>=1 and zs["Anzahlungen fällig (€)"]>=200 and zs["davon erhalten (€)"]>=100 and zs["Komplett bezahlt (Termine)"]==1,zs)
chk("Monat und Woche gehen",cA.get(D+"/uebersicht.xlsx?zeitraum=monat").status_code==200 and cA.get(D+"/uebersicht.xlsx?zeitraum=woche").status_code==200)
chk("Ungültiger Zeitraum: 400",cA.get(D+"/uebersicht.xlsx?zeitraum=eigen&von=x&bis=y").status_code==400)
chk("Anderer Kunde: leere Tabelle",len(list(openpyxl.load_workbook(io.BytesIO(cB.get(D+"/uebersicht.xlsx?zeitraum=monat").content))["Termine"].iter_rows(min_row=2,values_only=True)))==0)
# ---- Preise-Baustein
sid=cA.get(D).json()["site_id"]
r=httpx.get(B+f"/site/{sid}/leistungen",headers={"Origin":"https://"+dom}); chk("Preisliste für den Baustein",r.status_code==200 and [x["name"] for x in r.json()["leistungen"]]==["Schnitt","Balayage","Politur"] and r.json()["leistungen"][0]["preis"]=="45 €",r.text)
chk("Baustein im Skript vorhanden","data-mk" in open("/home/claude/markentrix-server/portal/static/js/site.js").read() and "mk-preise" in open("/home/claude/markentrix-server/portal/static/js/site.js").read())
# ---- Änderungen mit Freigabe
A=f"/api/admin/kunden/{kA}"
r=a.put(A+"/leistungs-freigabe",json={"freigabe":True}); chk("Freigabepflicht einschalten",r.status_code==200,r.text)
r=cA.post(D+"/leistungen",json=L(name="Politur Plus",preis="300 €",dauer_min=90)); chk("Neue Leistung wird nur vorgemerkt",r.status_code==200 and r.json().get("wartet") is True,r.text)
chk("Noch nicht live",[x["name"] for x in httpx.get(B+f"/site/{sid}/leistungen",headers={"Origin":"https://"+dom}).json()["leistungen"]]==["Schnitt","Balayage","Politur"])
r=cA.put(D+f"/leistungen/{ls['Schnitt']['id']}",json=L(name="Schnitt",schluessel="schnitt",preis="55 €",dauer_min=30)); chk("Preisänderung wird vorgemerkt",r.json().get("wartet") is True,r.text)
r=cA.delete(D+f"/leistungen/{ls['Politur']['id']}"); chk("Löschwunsch wird vorgemerkt",r.json().get("wartet") is True,r.text)
av=a.get(A+"/leistungs-aenderungen").json(); chk("Admin sieht 3 offene Wünsche mit Altwert",av["freigabe"] is True and len([x for x in av["aenderungen"] if x["status"]=="offen"])==3 and any(x["alt"] and x["alt"]["preis"]=="45 €" for x in av["aenderungen"]),av)
chk("Aufgaben fürs Team angelegt",int(sql(f"select count(*) from aufgaben where kunden_id='{kA}' and auto_schluessel like 'leistung-%' and status='offen'"))==3)
ids={x["name"]+x["aktion"]:x["id"] for x in av["aenderungen"]}
r=a.post(A+f"/leistungs-aenderungen/{ids['Politur Plusneu']}/uebernehmen"); chk("Admin übernimmt neue Leistung",r.status_code==200,r.text)
r=a.post(A+f"/leistungs-aenderungen/{ids['Schnittaendern']}/ablehnen",json={"notiz":"Bitte erst mit Chef klären"}); chk("Admin lehnt Preisänderung ab",r.status_code==200,r.text)
r=a.post(A+f"/leistungs-aenderungen/{ids['Politurloeschen']}/uebernehmen"); chk("Admin übernimmt Löschung",r.status_code==200,r.text)
pl=[(x["name"],x["preis"]) for x in httpx.get(B+f"/site/{sid}/leistungen",headers={"Origin":"https://"+dom}).json()["leistungen"]]
chk("Website zeigt jetzt den neuen Stand, abgelehnter Preis bleibt",pl==[("Schnitt","45 €"),("Balayage","ab 180 €"),("Politur Plus","300 €")],pl)
w=sql(f"select wert from wissensbasis where kunden_id='{kA}' and bereich='leistungen' and quelle='system'"); chk("Wissensdatenbank folgt",("Politur Plus" in w and "- Politur:" not in w),w)
r=a.post(A+f"/leistungs-aenderungen/{ids['Politur Plusneu']}/uebernehmen"); chk("Doppelt bearbeiten: 409",r.status_code==409,r.text)
ca=cA.get(D+"/leistungen/aenderungen").json()["aenderungen"]; chk("Kunde sieht Antwort des Teams",any(x["status"]=="abgelehnt" and "Chef" in x["notiz"] for x in ca),ca)
r=cA.post(D+"/leistungen",json=L(name="Nochmal",dauer_min=20)); aid=r.json()["id"]; r=cA.delete(D+f"/leistungen/aenderungen/{aid}"); chk("Kunde zieht Wunsch zurück",r.status_code==200,r.text)
r=cB.post(A+f"/leistungs-aenderungen/{aid}/uebernehmen"); chk("Kunde darf nicht freigeben",r.status_code==403,r.text)
r=a.put(A+"/leistungs-freigabe",json={"freigabe":False}); r=cA.post(D+"/leistungen",json=L(name="Sofort",dauer_min=20)); chk("Ohne Freigabepflicht sofort live",r.status_code==200 and "wartet" not in r.json(),r.text)
chk("Lesbare Feldnamen in der Mail an den Inhaber","Haarzustand: normal" in text(mail_an(f"chef@{dom}",r3)[0]))
