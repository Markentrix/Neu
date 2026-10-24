"""Website-Dienste: Kunden-Websites (von Markentrix gebaut und gehostet) nutzen Kontaktformular, Terminbuchung, PayPal und Bewertungen.

Drei Teile:
 - /site/<site-id>/…           öffentlich, wird vom Skript der Website aufgerufen (Herkunftsprüfung, Begrenzung, Lockfeld)
 - /api/kunde/website-dienste  der Inhaber trägt seine Daten ein (Zugangsdaten nur schreibbar, nie lesbar)
 - /api/admin/…                Status ohne Geheimnisse"""
import base64
import hashlib
import hmac
import json
import re
import secrets
import uuid
from datetime import date, datetime, timedelta, timezone
from html import escape
from urllib.parse import urlparse

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Request
from fastapi.responses import HTMLResponse, JSONResponse, Response
from pydantic import BaseModel, Field, ValidationError

from . import config, einstellungen, fotos, mailvorlage, push, quittung, webbetrieb
from . import sitedienste as sd
from .db import tx
from .security import EMAIL_RE, audit, client_ip, current_user, rate_limit, require_admin, require_inhaber

router = APIRouter()
TZ = sd.TZ
ROOT = urlparse(config.BASE_URL).hostname or ""
ROOT = ROOT[7:] if ROOT.startswith("portal.") else ROOT


# =============================================================== öffentlich (vom Skript der Website)
def _site(site_id: str) -> dict:
    if not re.match(r"^mk_[A-Za-z0-9]{6,40}$", site_id):
        raise HTTPException(404, "Unbekannte Website.")
    with tx("system") as cur:
        cur.execute("SELECT * FROM website_dienste WHERE site_id=%s", (site_id,))
        z = cur.fetchone()
        if not z:
            raise HTTPException(404, "Unbekannte Website.")
        cur.execute("SELECT status FROM kunden WHERE id=%s", (z["kunden_id"],))
        k = cur.fetchone()
        if not k or k["status"] == "beendet":      # nach Vertragsende nimmt die Website keine Anfragen und Buchungen mehr an
            raise HTTPException(404, "Unbekannte Website.")
        z["_hosts"] = sd.erlaubte_hosts(cur, z["kunden_id"], ROOT)
    return z


def _cors(request: Request, z: dict) -> dict:
    """Nur Seiten der Domains dieses Kunden dürfen die Schnittstelle aufrufen."""
    origin = request.headers.get("origin", "")
    host = (urlparse(origin).hostname or "").lower() if origin else ""
    if not host or host not in z["_hosts"]:
        raise HTTPException(403, "Diese Seite darf den Dienst nicht nutzen.")
    return {"Access-Control-Allow-Origin": origin, "Vary": "Origin", "Access-Control-Allow-Methods": "GET, POST, OPTIONS",
            "Access-Control-Allow-Headers": "Content-Type", "Access-Control-Max-Age": "600", "Cache-Control": "no-store"}


def _antwort(daten, kopf: dict, status: int = 200) -> JSONResponse:
    return JSONResponse(daten, status_code=status, headers=kopf)


async def _korper(request: Request) -> dict:
    roh = await request.body()
    if len(roh) > 20_000:
        raise HTTPException(413, "Die Eingabe ist zu lang.")
    try:
        d = json.loads(roh or b"{}")
    except ValueError:
        raise HTTPException(400, "Ungültige Anfrage.")
    if not isinstance(d, dict):
        raise HTTPException(400, "Ungültige Anfrage.")
    return d


def _pruefen(modell, daten: dict):
    try:
        return modell(**daten)
    except ValidationError:
        raise HTTPException(400, "Bitte füllen Sie alle Felder korrekt aus.")


@router.options("/site/{site_id}/{pfad:path}", include_in_schema=False)
def vorabfrage(site_id: str, pfad: str, request: Request):
    z = _site(site_id)
    return Response(status_code=204, headers=_cors(request, z))


@router.get("/site/{site_id}/config", include_in_schema=False)
def oeff_config(site_id: str, request: Request):
    z = _site(site_id)
    kopf = _cors(request, z)
    rate_limit("siteconf:" + client_ip(request), 120, 300)
    return _antwort({
        "kontakt": bool(z["mail_aktiv"]),
        "termin": {"aktiv": bool(z["kal_aktiv"]), "dauer": z["kal_dauer"], "ort": z["kal_ort"] or ""},
        "paypal": {"aktiv": bool(z["pp_aktiv"] and (z["pp_client_id"] or z["pp_me_link"])), "client_id": z["pp_client_id"] or "", "waehrung": z["pp_waehrung"],
                   "me_link": z["pp_me_link"] or ""},
        "bewertungen": bool(z["bew_aktiv"]),
        "analytics": {"aktiv": bool(z["ga_aktiv"] and z["ga_id"]), "id": z["ga_id"] or ""},
        "fotos": {"aktiv": bool(z["foto_aktiv"]), "max": fotos.MAX_ANZAHL, "tage": z["foto_tage"]},
    }, kopf)


async def _foto_korper(request: Request) -> bytes:
    try:
        if int(request.headers.get("content-length") or 0) > fotos.MAX_BYTES + 1024:
            raise HTTPException(413, "Das Foto ist zu groß (höchstens 6 MB).")
    except ValueError:
        raise HTTPException(400, "Ungültige Anfrage.")
    return (await request.body())[:fotos.MAX_BYTES + 1]


@router.post("/site/{site_id}/foto", include_in_schema=False)
def oeff_foto(site_id: str, request: Request, roh: bytes = Depends(_foto_korper)):
    """Ein Foto als Rohdaten (JPG, PNG, WebP). Antwort: Kennung, die beim Absenden von Kontaktformular bzw. Termin im Feld „fotos“ mitgeschickt wird."""
    z = _site(site_id)
    kopf = _cors(request, z)
    rate_limit("fotoup:" + client_ip(request), 30, 600)
    rate_limit("fotoupsite:" + site_id, 300, 86400)
    if z.get("verzug_pause"):
        raise HTTPException(503, "Dieser Dienst ist vorübergehend nicht verfügbar. Bitte rufen Sie uns direkt an.")
    if not z["foto_aktiv"]:
        raise HTTPException(409, "Foto-Upload ist nicht eingeschaltet.")
    with tx("system") as cur:
        token = fotos.speichern(cur, z["kunden_id"], roh)
    return _antwort({"ok": True, "token": token}, kopf)


class KontaktIn(BaseModel):
    name: str = Field(min_length=2, max_length=120)
    email: str = Field(max_length=200)
    telefon: str = Field(default="", max_length=60)
    nachricht: str = Field(min_length=3, max_length=3000)
    einwilligung: bool = False
    hp: str = Field(default="", max_length=200)          # Lockfeld: Menschen sehen es nicht, Bots füllen es aus
    seite: str = Field(default="", max_length=300)


def _empfaenger(text: str) -> list[str]:
    return [x.strip() for x in re.split(r"[,;\s]+", text or "") if x.strip()]


@router.post("/site/{site_id}/kontakt", include_in_schema=False)
def oeff_kontakt(site_id: str, request: Request, daten: dict = Depends(_korper)):
    z = _site(site_id)
    kopf = _cors(request, z)
    rate_limit("kontakt:" + client_ip(request), 6, 600)
    rate_limit("kontaktsite:" + site_id, 150, 86400)
    if z.get("verzug_pause"):
        raise HTTPException(503, "Dieser Dienst ist vorübergehend nicht verfügbar. Bitte rufen Sie uns direkt an.")
    if not z["mail_aktiv"]:
        raise HTTPException(409, "Das Kontaktformular ist nicht aktiv.")
    d = _pruefen(KontaktIn, daten)
    if d.hp:
        return _antwort({"ok": True}, kopf)               # Bot: so tun als ob, nichts senden
    email = d.email.strip().lower()
    if not EMAIL_RE.match(email):
        raise HTTPException(400, "Die E-Mail-Adresse ist ungültig.")
    if not d.einwilligung:
        raise HTTPException(400, "Bitte stimmen Sie der Verarbeitung Ihrer Angaben zu.")
    kid = z["kunden_id"]
    an = [a for a in _empfaenger(z["mail_empfaenger"]) if EMAIL_RE.match(a)][:5]
    kontakt = email + (", " + d.telefon.strip() if d.telefon.strip() else "")
    with tx("system") as cur:
        cur.execute("INSERT INTO anfragen (kunden_id, name, kontakt, nachricht, quelle) VALUES (%s,%s,%s,%s,'website') RETURNING id",
                    (kid, d.name.strip(), kontakt, d.nachricht.strip()))
        aid = cur.fetchone()["id"]
        foto_ids = fotos.verknuepfen(cur, kid, fotos.kennungen(daten.get("fotos")), z["foto_tage"], aid) if z["foto_aktiv"] else []
        anhang = fotos.fuer_mail(cur, kid, foto_ids) or None
        ids = push.kunde(cur, kid)
    push.senden(ids, "Neue Anfrage über die Website", f"{d.name.strip()} hat das Kontaktformular genutzt.", "/app#anfragen", "anfrage")
    betreff = z["mail_betreff"] or "Neue Anfrage über Ihre Website"
    text = (f"Neue Anfrage über das Kontaktformular Ihrer Website.\n\nName: {d.name.strip()}\nE-Mail: {email}\nTelefon: {d.telefon.strip() or '–'}\n\n"
            f"Nachricht:\n{d.nachricht.strip()}\n\n"
            + (f"Fotos: {len(foto_ids)} Stück " + ("im Anhang und " if anhang else "") + f"im Kundenportal unter „Anfragen“. Sie werden nach {z['foto_tage']} Tagen automatisch gelöscht.\n\n" if foto_ids else "")
            + f"—\nSie können direkt auf diese E-Mail antworten, die Antwort geht an {email}.")
    if not an:
        sd.ereignis(kid, "kontakt", False, "Kein Empfänger eingetragen, die Anfrage liegt nur im Portal unter „Anfragen“.")
        return _antwort({"ok": True}, kopf)
    fehler = ""
    for a in an:
        ok, via = sd.kundenmail(kid, a, betreff, text, reply_to=email, anhang=anhang, art="kontakt")
        if not ok:
            fehler = via
    if fehler:
        sd.ereignis(kid, "kontakt", False, "Versand fehlgeschlagen, Anfrage liegt im Portal. " + fehler[:150])
    else:
        sd.ereignis(kid, "kontakt", True, "gesendet")
    return _antwort({"ok": True}, kopf)


@router.get("/site/{site_id}/leistungen", include_in_schema=False)
def oeff_leistungen(site_id: str, request: Request):
    """Preisliste für den Baustein data-mk="preise": immer der aktuelle Stand aus dem Portal."""
    z = _site(site_id)
    kopf = _cors(request, z)
    rate_limit("siteleist:" + client_ip(request), 60, 300)
    with tx("system") as cur:
        cur.execute("SELECT name, beschreibung, preis, dauer_min FROM website_leistungen WHERE kunden_id=%s AND aktiv ORDER BY sort, name", (z["kunden_id"],))
        l = [{"name": r["name"], "beschreibung": r["beschreibung"] or "", "preis": r["preis"] or "", "dauer": r["dauer_min"]} for r in cur.fetchall()]
    return _antwort({"leistungen": l}, kopf)


@router.get("/site/{site_id}/termine", include_in_schema=False)
def oeff_termine(site_id: str, request: Request):
    z = _site(site_id)
    kopf = _cors(request, z)
    rate_limit("sitetermine:" + client_ip(request), 60, 300)
    if not z["kal_aktiv"]:
        return _antwort({"aktiv": False, "tage": []}, kopf)
    with tx("system") as cur:
        tage: dict = {}
        for s in sd.freie_slots(cur, z):
            lokal = s.astimezone(TZ)
            tage.setdefault(lokal.strftime("%Y-%m-%d"), []).append({"start": s.isoformat(), "uhr": lokal.strftime("%H:%M")})
    return _antwort({"aktiv": True, "dauer": z["kal_dauer"], "ort": z["kal_ort"] or "", "tage": [{"datum": d, "slots": s} for d, s in tage.items()]}, kopf)


class TerminIn(BaseModel):
    start: str = Field(max_length=40)
    name: str = Field(min_length=2, max_length=120)
    email: str = Field(max_length=200)
    telefon: str = Field(default="", max_length=60)
    nachricht: str = Field(default="", max_length=1500)
    einwilligung: bool = False
    hp: str = Field(default="", max_length=200)


def _termin_mails(kid, g: dict, z: dict, firma: str, an: list, foto_ids: list | None = None):
    lokal = g["start"].astimezone(TZ).strftime("%d.%m.%Y um %H:%M Uhr")
    ort = z["kal_ort"] or ""
    absagen = f"{config.BASE_URL}/site/{z['site_id']}/termin/absagen?t={g['token']}"
    anhang = ("termin.ics", sd.ics(str(g["id"]), g["start"], g["ende"], "Termin bei " + firma, ort), "text/calendar")
    wer = " · ".join(x for x in (g.get("leistung_name"), ("bei " + g["mitarbeiter_name"]) if g.get("mitarbeiter_name") else "") if x)
    sd.kundenmail(kid, g["email"], f"Ihr Termin bei {firma} am {lokal}",
                  f"Guten Tag {g['name']},\n\nIhr Termin ist gebucht: {lokal}.{(chr(10) + wer) if wer else ''}\n{ort}\n{sd.zahlungsabsatz(z, g.get('anzahlung'))}\nDie Kalendereinladung hängt an. Falls Sie nicht können, sagen Sie bitte hier ab:\n{absagen}\n\nBis bald!\n\nViele Grüße\n{firma}",
                  anhang=anhang, art="termin")
    bilder: list = []
    if foto_ids and an:
        with tx("system") as cur:
            bilder = fotos.fuer_mail(cur, kid, foto_ids)
    foto_text = (f"Fotos: {len(foto_ids)} Stück " + ("im Anhang und " if bilder else "") + f"im Kundenportal unter „Heute“. Sie werden nach {z['foto_tage']} Tagen automatisch gelöscht.\n\n") if foto_ids else ""
    for a in an:
        sd.kundenmail(kid, a, f"Neuer Termin: {g['name']} am {lokal}",
                      f"{g['name']} ({g['email']}, {g.get('telefon') or 'keine Nummer'}) hat über Ihre Website einen Termin gebucht: {lokal}.\n\n{g.get('nachricht') or ''}\n\n{foto_text}"
                      f"Der Termin steht im Portal unter „Heute“ und in der Einladung im Anhang.", reply_to=g["email"],
                      anhang=[("termin.ics", sd.ics(str(g["id"]), g["start"], g["ende"], "Termin: " + g["name"], ort), "text/calendar")] + bilder, art="termin")


@router.post("/site/{site_id}/termin", include_in_schema=False)
def oeff_termin(site_id: str, request: Request, bg: BackgroundTasks, daten: dict = Depends(_korper)):
    z = _site(site_id)
    kopf = _cors(request, z)
    rate_limit("sitebuchen:" + client_ip(request), 8, 3600)
    rate_limit("sitebuchensite:" + site_id, 100, 86400)
    if z.get("verzug_pause"):
        raise HTTPException(503, "Dieser Dienst ist vorübergehend nicht verfügbar. Bitte rufen Sie uns direkt an.")
    if not z["kal_aktiv"]:
        raise HTTPException(409, "Die Terminbuchung ist nicht aktiv.")
    if z["kal_freigabe"]:
        raise HTTPException(409, "Dieser Betrieb bestätigt Termine persönlich. Bitte nutzen Sie das Kontaktformular für Ihre Terminanfrage.")
    d = _pruefen(TerminIn, daten)
    if d.hp:
        return _antwort({"ok": True, "wann": ""}, kopf)
    email = d.email.strip().lower()
    if not EMAIL_RE.match(email):
        raise HTTPException(400, "Die E-Mail-Adresse ist ungültig.")
    if not d.einwilligung:
        raise HTTPException(400, "Bitte stimmen Sie der Verarbeitung Ihrer Angaben zu.")
    try:
        start = datetime.fromisoformat(d.start).astimezone(timezone.utc)
    except ValueError:
        raise HTTPException(400, "Ungültiger Termin.")
    kid = z["kunden_id"]
    voll = "Dieser Termin ist leider nicht mehr frei. Bitte wählen Sie einen anderen."
    with tx("system") as cur:
        cur.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s, 0))", (str(kid),))
        platz = sd.freier_platz(cur, z, start)
        if not platz:
            raise HTTPException(409, voll)
        rid = platz[1]
        ende = start + timedelta(minutes=z["kal_dauer"])
        cur.execute("SELECT firma FROM kunden WHERE id=%s", (kid,))
        firma = cur.fetchone()["firma"]
        kontakt = email + (", " + d.telefon.strip() if d.telefon.strip() else "")
        cur.execute("""INSERT INTO termine (kunden_id, titel, beginn, ende, kontakt_name, kontakt_info, status, notiz)
                       VALUES (%s,%s,%s,%s,%s,%s,'bestaetigt',%s) RETURNING id""",
                    (kid, "Termin über die Website", start, ende, d.name.strip(), kontakt, d.nachricht.strip() or None))
        tid = cur.fetchone()["id"]
        token = secrets.token_urlsafe(24)
        try:
            cur.execute("""INSERT INTO website_termine (kunden_id, termin_id, start, ende, name, email, telefon, nachricht, token, mitarbeiter_id)
                           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id, start, ende, name, email, telefon, nachricht, token""",
                        (kid, tid, start, ende, d.name.strip(), email, d.telefon.strip() or None, d.nachricht.strip() or None, token, rid))
        except Exception:
            raise HTTPException(409, voll)         # gleichzeitig gebucht: der Rest der Transaktion wird verworfen
        g = cur.fetchone()
        foto_ids = fotos.verknuepfen(cur, kid, fotos.kennungen(daten.get("fotos")), z["foto_tage"], None, g["id"]) if z["foto_aktiv"] else []
        an = [a for a in _empfaenger(z["kal_empfaenger"] or z["mail_empfaenger"]) if EMAIL_RE.match(a)][:5]
        ids = push.kunde(cur, kid)
    wann = start.astimezone(TZ).strftime("%d.%m.%Y um %H:%M Uhr")
    push.senden(ids, "Neuer Termin über die Website", f"{d.name.strip()}: {wann}", "/app#heute", "termin")
    sd.ereignis(kid, "termin", True, "gebucht")
    bg.add_task(_termin_mails, kid, g, z, firma, an, foto_ids)
    return _antwort({"ok": True, "wann": wann}, kopf)


def _seite(text: str) -> HTMLResponse:
    return _seite_html("Termin", f"<p style=\"margin:0\">{escape(text)}</p>")


def _seite_html(titel: str, inhalt: str, kopf: str = "") -> HTMLResponse:
    return HTMLResponse('<!doctype html><html lang="de"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>' + escape(titel) + '</title>'
                        '<body style="margin:0;background:#f3f4f6;font:16px system-ui,-apple-system,Segoe UI,Roboto,sans-serif;color:#0B1020;line-height:1.5">'
                        '<main style="max-width:560px;margin:6vh auto;padding:0 16px"><div style="background:#fff;border-radius:12px;padding:24px 26px;box-shadow:0 1px 3px rgba(0,0,0,.08)">'
                        + kopf + inhalt + '</div></main></body></html>')


def _seitenkopf(site_id: str, firma: str, hat_logo: bool) -> str:
    """Kopf der Absage- und Umbuchungsseiten: Logo des Betriebs (falls hinterlegt), sonst der Firmenname."""
    if hat_logo:
        return f'<img src="/site/{escape(site_id)}/logo" alt="{escape(firma, quote=True)}" style="max-height:56px;max-width:220px;display:block;margin:0 0 16px">'
    return f'<div style="font-weight:700;font-size:18px;margin:0 0 14px">{escape(firma)}</div>'


@router.get("/site/{site_id}/logo", include_in_schema=False)
def oeff_logo(site_id: str, request: Request):
    """Logo des Betriebs für die Absage-Seiten (dasselbe Bild steht in den Mails an die Gäste)."""
    rate_limit("sitelogo:" + client_ip(request), 120, 300)
    if not re.match(r"^mk_[A-Za-z0-9]{6,40}$", site_id):
        raise HTTPException(404, "Nicht gefunden.")
    with tx("system") as cur:
        cur.execute("SELECT k.logo, k.logo_typ FROM website_dienste d JOIN kunden k ON k.id=d.kunden_id WHERE d.site_id=%s AND k.status<>'beendet'", (site_id,))
        r = cur.fetchone()
    if not r or not r["logo"] or r["logo_typ"] not in ("image/png", "image/jpeg"):
        raise HTTPException(404, "Nicht gefunden.")
    return Response(bytes(r["logo"]), media_type=r["logo_typ"], headers={"Cache-Control": "public, max-age=3600", "X-Content-Type-Options": "nosniff"})


def _knopf_stil() -> str:
    return "font:inherit;padding:10px 18px;border-radius:8px;border:1px solid #0B1020;background:#0B1020;color:#fff;cursor:pointer;margin:4px 6px 4px 0"


def _buchung_zu_token(cur, site_id: str, t: str):
    cur.execute("""SELECT w.*, d.site_id, d.storno_std, d.mail_empfaenger, d.kal_empfaenger, d.kal_ort, d.kal_dauer, k.firma, (k.logo IS NOT NULL) AS hat_logo
                   FROM website_termine w JOIN website_dienste d ON d.kunden_id=w.kunden_id JOIN kunden k ON k.id=w.kunden_id WHERE w.token=%s AND d.site_id=%s""", (t, site_id))
    return cur.fetchone()


def _regel(g: dict) -> tuple[bool, float]:
    """(kostenlos möglich, Stunden bis zum Termin)."""
    std = (g["start"] - datetime.now(timezone.utc)).total_seconds() / 3600
    return std >= g["storno_std"], std


def _anz_text(g: dict) -> str:
    return str(g["anzahlung_betrag"] if g["anzahlung_betrag"] is not None else (g["anzahlung"] or "")).replace(".", ",")


@router.get("/site/{site_id}/termin/komme", include_in_schema=False)
def oeff_komme(site_id: str, t: str, request: Request):
    rate_limit("sitekomme:" + client_ip(request), 30, 600)
    if not re.match(r"^mk_[A-Za-z0-9]{6,40}$", site_id) or not 10 <= len(t) <= 100:
        return _seite("Dieser Link ist ungültig.")
    with tx("system") as cur:
        g = _buchung_zu_token(cur, site_id, t)
    if not g or g["status"] != "gebucht" or g["start"] <= datetime.now(timezone.utc):
        return _seite("Dieser Termin ist nicht mehr aktiv.")
    lokal = g["start"].astimezone(TZ).strftime("%d.%m.%Y um %H:%M Uhr")
    return _seite_html("Termin bestätigen", f'<h1 style="font-size:22px">Ich komme</h1><p><b>{escape(g["firma"])}</b><br>{escape(lokal)}</p><p>Bitte bestätigen Sie kurz, dass Sie den Termin wahrnehmen. Dann wissen wir Bescheid und planen mit Ihnen.</p>'
                       f'<form method="post" action="/site/{site_id}/termin/komme"><input type="hidden" name="t" value="{escape(t)}"><button style="{_knopf_stil()}">Ja, ich komme</button></form>'
                       f'<p><a href="/site/{site_id}/termin/absagen?t={escape(t)}">Ich kann doch nicht</a></p>')


@router.post("/site/{site_id}/termin/komme", include_in_schema=False)
async def oeff_komme_senden(site_id: str, request: Request):
    rate_limit("sitekommep:" + client_ip(request), 20, 600)
    form = await request.form()
    t = str(form.get("t") or "")
    if not re.match(r"^mk_[A-Za-z0-9]{6,40}$", site_id) or not 10 <= len(t) <= 100:
        return _seite("Dieser Link ist ungültig.")
    with tx("system") as cur:
        g = _buchung_zu_token(cur, site_id, t)
        if not g or g["status"] != "gebucht" or g["start"] <= datetime.now(timezone.utc):
            return _seite("Dieser Termin ist nicht mehr aktiv.")
        cur.execute("UPDATE website_termine SET kommt_am=coalesce(kommt_am, now()) WHERE id=%s", (g["id"],))
    return _seite("Danke, wir haben Ihre Bestätigung. Bis bald!")


AKTIONEN = {"bestaetigen": "bestätigen", "ablehnen": "ablehnen"}


def aktionstoken(bid, aktion: str) -> str:
    """Unterschriebener Schlüssel für den Knopf in der Mail an den Betrieb (nicht erratbar, an Buchung und Aktion gebunden)."""
    mac = hmac.new(config.PORTAL_SECRET.encode(), f"mailaktion:{aktion}:{bid}".encode(), hashlib.sha256).digest()
    return base64.urlsafe_b64encode(mac).decode().rstrip("=")[:32]


def aktionslink(site_id: str, bid, aktion: str) -> str:
    return f"{config.BASE_URL}/site/{site_id}/buchung/{bid}/{aktion}?k={aktionstoken(bid, aktion)}"


def _aktion_laden(cur, site_id: str, bid: str, aktion: str, k: str):
    """Buchung zum Knopf oder eine Fehlermeldung (Text)."""
    if aktion not in AKTIONEN or not re.match(r"^mk_[A-Za-z0-9]{6,40}$", site_id) or not re.match(r"^[0-9a-f-]{36}$", bid) or not hmac.compare_digest(k, aktionstoken(bid, aktion)):
        return None, "Dieser Link ist ungültig."
    cur.execute("""SELECT w.id, w.kunden_id, w.status, w.start, w.name, w.email, w.telefon, w.nachricht, w.ref, k.firma, (k.logo IS NOT NULL) AS hat_logo,
                          l.name AS leistung, m.name AS mitarbeiter
                   FROM website_termine w JOIN website_dienste d ON d.kunden_id=w.kunden_id JOIN kunden k ON k.id=w.kunden_id
                        LEFT JOIN website_leistungen l ON l.id=w.leistung_id LEFT JOIN website_mitarbeiter m ON m.id=w.mitarbeiter_id
                   WHERE w.id=%s AND d.site_id=%s""", (bid, site_id))
    g = cur.fetchone()
    if not g:
        return None, "Diese Buchung gibt es nicht mehr."
    if g["status"] == "gebucht" and aktion == "bestaetigen":
        return None, "Dieser Termin ist schon bestätigt."
    if g["status"] == "abgesagt":
        return None, "Dieser Termin ist schon abgesagt oder abgelehnt."
    if aktion == "bestaetigen" and g["status"] != "reserviert":
        return None, "Diese Buchung wartet nicht mehr auf Bestätigung (abgelaufen oder bearbeitet)."
    if g["start"] <= datetime.now(timezone.utc):
        return None, "Dieser Termin liegt in der Vergangenheit."
    return g, ""


@router.get("/site/{site_id}/buchung/{bid}/{aktion}", include_in_schema=False)
def oeff_aktion(site_id: str, bid: str, aktion: str, k: str, request: Request):
    """Zeigt die Buchung und fragt nach. Ausgeführt wird erst nach dem Klick (Mail-Programme rufen Links manchmal von selbst auf)."""
    rate_limit("siteaktion:" + client_ip(request), 30, 600)
    with tx("system") as cur:
        g, fehler = _aktion_laden(cur, site_id, bid, aktion, k)
    if not g:
        return _seite(fehler)
    lokal = g["start"].astimezone(TZ).strftime("%d.%m.%Y um %H:%M Uhr")
    zeilen = "".join(f"<p style=\"margin:2px 0\">{escape(x)}</p>" for x in (g["leistung"], ("bei " + g["mitarbeiter"]) if g["mitarbeiter"] else "", g["name"],
                                                                        " · ".join(y for y in (g["email"], g["telefon"]) if y), g["nachricht"] or "") if x)
    titel = "Termin bestätigen" if aktion == "bestaetigen" else "Termin ablehnen"
    extra = "" if aktion == "bestaetigen" else '<label style="display:block;margin:12px 0 4px;font-size:14px">Kurze Begründung für den Gast (freiwillig)</label><textarea name="grund" maxlength="600" rows="3" style="width:100%;box-sizing:border-box;font:inherit;padding:8px"></textarea>'
    frage = "Der Gast bekommt eine Bestätigung mit Kalender-Einladung." if aktion == "bestaetigen" else "Der Gast bekommt eine freundliche Absage."
    return _seite_html(titel, f'<h1 style="font-size:22px;margin:0 0 10px">{titel}</h1><p><b>{escape(lokal)}</b></p>{zeilen}<p style="margin-top:12px">{frage}</p>'
                       f'<form method="post" action="/site/{escape(site_id)}/buchung/{escape(bid)}/{aktion}"><input type="hidden" name="k" value="{escape(k)}">{extra}'
                       f'<button style="{_knopf_stil()}">Ja, {AKTIONEN[aktion]}</button></form>', _seitenkopf(site_id, g["firma"], g["hat_logo"]))


@router.post("/site/{site_id}/buchung/{bid}/{aktion}", include_in_schema=False)
async def oeff_aktion_senden(site_id: str, bid: str, aktion: str, request: Request, bg: BackgroundTasks):
    rate_limit("siteaktionp:" + client_ip(request), 20, 600)
    form = await request.form()
    k = str(form.get("k") or "")
    with tx("system") as cur:
        g, fehler = _aktion_laden(cur, site_id, bid, aktion, k)
    if not g:
        return _seite(fehler)
    try:
        if aktion == "bestaetigen":
            _bestaetigen(g["kunden_id"], bid, "", "", bg, None)
            return _seite("Der Termin ist bestätigt. Der Gast bekommt die Bestätigung per E-Mail.")
        _ablehnen(g["kunden_id"], bid, str(form.get("grund") or "")[:600], bg, None)
        return _seite("Der Termin ist abgelehnt. Der Gast wurde per E-Mail informiert.")
    except HTTPException as e:
        return _seite(str(e.detail))


@router.get("/site/{site_id}/termin/absagen", include_in_schema=False)
def oeff_absagen(site_id: str, t: str, request: Request):
    """Zeigt die Absageregel und fragt nach. Absagen passiert erst nach dem Klick (Mail-Programme rufen Links manchmal von selbst auf)."""
    rate_limit("siteabsagen:" + client_ip(request), 30, 600)
    if not re.match(r"^mk_[A-Za-z0-9]{6,40}$", site_id) or not 10 <= len(t) <= 100:
        return _seite("Dieser Link ist ungültig.")
    with tx("system") as cur:
        g = _buchung_zu_token(cur, site_id, t)
    if not g or g["status"] not in ("gebucht", "reserviert") or g["start"] <= datetime.now(timezone.utc):
        return _seite("Dieser Termin ist nicht mehr aktiv.")
    lokal = g["start"].astimezone(TZ).strftime("%d.%m.%Y um %H:%M Uhr")
    frei, _ = _regel(g)
    anz = _anz_text(g)
    if frei:
        regel = "Die kostenlose Absage ist noch möglich." + (f" Ihre Anzahlung von {escape(anz)} € wird erstattet." if g["anzahlung_erhalten_am"] and anz else "")
    else:
        regel = (f"Die Absage liegt innerhalb der Frist von {g['storno_std']} Stunden vor dem Termin. Dann ist der volle Betrag zu zahlen" + (f", die Anzahlung von {escape(anz)} € wird nicht erstattet" if anz else "") +
                 ". Wenn Sie den Termin nicht wahrnehmen können, sprechen Sie bitte direkt mit uns.")
    umbuchen = f'<p><a href="/site/{site_id}/termin/verschieben?t={escape(t)}">Lieber einen anderen Termin wählen</a></p>' if frei else ""
    return _seite_html("Termin absagen", f'<h1 style="font-size:22px;margin:0 0 10px">Termin absagen</h1><p><b>{escape(lokal)}</b></p><p>{regel}</p>'
                       f'<form method="post" action="/site/{site_id}/termin/absagen"><input type="hidden" name="t" value="{escape(t)}"><button style="{_knopf_stil()}">Ja, Termin absagen</button></form>{umbuchen}', _seitenkopf(site_id, g["firma"], g["hat_logo"]))


@router.post("/site/{site_id}/termin/absagen", include_in_schema=False)
async def oeff_absagen_senden(site_id: str, request: Request, bg: BackgroundTasks):
    rate_limit("siteabsagenp:" + client_ip(request), 20, 600)
    form = await request.form()
    t = str(form.get("t") or "")
    if not re.match(r"^mk_[A-Za-z0-9]{6,40}$", site_id) or not 10 <= len(t) <= 100:
        return _seite("Dieser Link ist ungültig.")
    with tx("system") as cur:
        g = _buchung_zu_token(cur, site_id, t)
        if not g or g["status"] not in ("gebucht", "reserviert") or g["start"] <= datetime.now(timezone.utc):
            return _seite("Dieser Termin ist nicht mehr aktiv.")
        frei, _ = _regel(g)
        cur.execute("UPDATE website_termine SET status='abgesagt', abgesagt_am=now(), abgesagt_von='besucher', storno_kostenlos=%s WHERE id=%s", (frei, g["id"]))
        if g["termin_id"]:
            cur.execute("UPDATE termine SET status='abgesagt' WHERE id=%s", (g["termin_id"],))
        ids = push.kunde(cur, g["kunden_id"])
    wann = g["start"].astimezone(TZ).strftime("%d.%m.%Y %H:%M")
    anz = _anz_text(g)
    erstatten = bool(frei and g["anzahlung_erhalten_am"] and anz)
    kosten = "" if frei else (f"\nDie Absage kam später als {g['storno_std']} Stunden vor dem Termin: Der volle Betrag ist fällig" + (f", die Anzahlung von {anz} € wird einbehalten." if g["anzahlung_erhalten_am"] and anz else "."))
    hinweis = f"\nBitte die Anzahlung von {anz} € zurückzahlen und im Portal beim Termin „Erstattet“ klicken." if erstatten else ""
    bg.add_task(sd.warteliste_benachrichtigen, g["kunden_id"], g["start"].astimezone(TZ).date())
    push.senden(ids, "Termin abgesagt" + (" (Erstattung offen)" if erstatten else "" if frei else " (zu spät)"), f"{g['name']} hat den Termin am {wann} Uhr abgesagt.", "/app#dienste", "termin")
    for a in [x for x in _empfaenger(g["kal_empfaenger"] or g["mail_empfaenger"]) if EMAIL_RE.match(x)][:5]:
        bg.add_task(sd.kundenmail, g["kunden_id"], a, f"Termin abgesagt: {g['name']}", f"{g['name']} hat den Termin am {wann} Uhr abgesagt.{kosten}{hinweis}", None, None, "termin")
    gast = (f"Guten Tag {g['name']},\n\nIhr Termin am {wann} Uhr ist abgesagt." + ("\nIhre Anzahlung von " + anz + " € wird erstattet." if erstatten else "") + (kosten if not frei else "") +
            f"\n\nViele Grüße\n{g['firma']}")
    bg.add_task(sd.kundenmail, g["kunden_id"], g["email"], f"Absage bestätigt: Ihr Termin bei {g['firma']}", gast, None, None, "termin")
    return _seite("Ihr Termin wurde abgesagt." + (" Die Anzahlung wird erstattet." if erstatten else "") + (" Da die Absage nach der Frist kam, ist der volle Betrag fällig." if not frei else "") + " Sie können jederzeit einen neuen buchen.")


@router.get("/site/{site_id}/termin/verschieben", include_in_schema=False)
def oeff_verschieben(site_id: str, t: str, request: Request):
    rate_limit("siteverschieben:" + client_ip(request), 30, 600)
    if not re.match(r"^mk_[A-Za-z0-9]{6,40}$", site_id) or not 10 <= len(t) <= 100:
        return _seite("Dieser Link ist ungültig.")
    with tx("system") as cur:
        g = _buchung_zu_token(cur, site_id, t)
        if not g or g["status"] != "gebucht" or g["start"] <= datetime.now(timezone.utc):
            return _seite("Dieser Termin ist nicht mehr aktiv.")
        frei, _ = _regel(g)
        if not frei:
            return _seite(f"Umbuchen ist bis {g['storno_std']} Stunden vor dem Termin möglich. Bitte sprechen Sie direkt mit uns.")
        z = sd.zeile(cur, g["kunden_id"])
        lz = None
        if g["leistung_id"]:
            cur.execute("SELECT * FROM website_leistungen WHERE id=%s", (g["leistung_id"],))
            lz = cur.fetchone()
        plan = sd.slotplan(cur, z, leistung=lz, mitarbeiter=g["mitarbeiter_id"])
    tage = []
    for tag, slots in sorted(plan.items()):
        fr = [x for x in slots if x[2]]
        if fr:
            tage.append((tag, fr))
        if len(tage) >= 14:
            break
    if not tage:
        return _seite("Im Moment sind keine anderen Zeiten frei. Bitte sprechen Sie direkt mit uns.")
    WT = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]
    teile = []
    for tag, fr in tage:
        knoepfe = "".join(f'<button name="neu" value="{escape(u.isoformat())}" style="{_knopf_stil()}">{escape(zt)}</button>' for zt, u, _f, _r in fr)
        teile.append(f'<details style="margin:6px 0"><summary><b>{WT[tag.weekday()]}, {tag.strftime("%d.%m.%Y")}</b></summary>{knoepfe}</details>')
    lokal = g["start"].astimezone(TZ).strftime("%d.%m.%Y um %H:%M Uhr")
    return _seite_html("Termin verschieben", f'<h1 style="font-size:22px">Termin verschieben</h1><p>Ihr aktueller Termin: <b>{escape(lokal)}</b>. Bitte wählen Sie eine neue Zeit.</p>'
                       f'<form method="post" action="/site/{site_id}/termin/verschieben"><input type="hidden" name="t" value="{escape(t)}">{"".join(teile)}</form>')


@router.post("/site/{site_id}/termin/verschieben", include_in_schema=False)
async def oeff_verschieben_senden(site_id: str, request: Request, bg: BackgroundTasks):
    rate_limit("siteverschiebenp:" + client_ip(request), 20, 600)
    form = await request.form()
    t, neu = str(form.get("t") or ""), str(form.get("neu") or "")
    if not re.match(r"^mk_[A-Za-z0-9]{6,40}$", site_id) or not 10 <= len(t) <= 100:
        return _seite("Dieser Link ist ungültig.")
    try:
        start = datetime.fromisoformat(neu).astimezone(timezone.utc)
    except ValueError:
        return _seite("Diese Zeit gibt es nicht. Bitte noch einmal wählen.")
    with tx("system") as cur:
        g = _buchung_zu_token(cur, site_id, t)
        if not g or g["status"] != "gebucht" or g["start"] <= datetime.now(timezone.utc):
            return _seite("Dieser Termin ist nicht mehr aktiv.")
        if not _regel(g)[0]:
            return _seite(f"Umbuchen ist bis {g['storno_std']} Stunden vor dem Termin möglich. Bitte sprechen Sie direkt mit uns.")
        cur.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s, 0))", (str(g["kunden_id"]),))
        z = sd.zeile(cur, g["kunden_id"])
        lz = None
        if g["leistung_id"]:
            cur.execute("SELECT * FROM website_leistungen WHERE id=%s", (g["leistung_id"],))
            lz = cur.fetchone()
        tag = start.astimezone(TZ).date()
        treffer = [(u, rid) for zt, u, f, rid in sd.slotplan(cur, z, leistung=lz, mitarbeiter=g["mitarbeiter_id"]).get(tag, []) if u == start and f]
        if not treffer:
            return _seite("Diese Zeit ist leider nicht mehr frei. Bitte gehen Sie zurück und wählen Sie eine andere.")
        dauer = g["ende"] - g["start"]
        rid = treffer[0][1] or g["mitarbeiter_id"]
        try:
            cur.execute("UPDATE website_termine SET start=%s, ende=%s, mitarbeiter_id=%s, erinnerung=0, verschoben_am=now() WHERE id=%s ", (start, start + dauer, rid, g["id"]))
        except Exception:
            return _seite("Diese Zeit wurde gerade vergeben. Bitte wählen Sie eine andere.")
        if g["termin_id"]:
            cur.execute("UPDATE termine SET beginn=%s, ende=%s WHERE id=%s", (start, start + dauer, g["termin_id"]))
        ids = push.kunde(cur, g["kunden_id"])
    bg.add_task(sd.warteliste_benachrichtigen, g["kunden_id"], g["start"].astimezone(TZ).date())
    alt = g["start"].astimezone(TZ).strftime("%d.%m.%Y %H:%M")
    neu_txt = start.astimezone(TZ).strftime("%d.%m.%Y um %H:%M Uhr")
    push.senden(ids, "Termin verschoben", f"{g['name']}: von {alt} auf {neu_txt}.", "/app#heute", "termin")
    for a in [x for x in _empfaenger(g["kal_empfaenger"] or g["mail_empfaenger"]) if EMAIL_RE.match(x)][:5]:
        bg.add_task(sd.kundenmail, g["kunden_id"], a, f"Termin verschoben: {g['name']}", f"{g['name']} hat den Termin von {alt} Uhr auf {neu_txt} verschoben.", None, None, "termin")
    absagen = f"{config.BASE_URL}/site/{site_id}/termin/absagen?t={t}"
    bg.add_task(sd.kundenmail, g["kunden_id"], g["email"], f"Ihr neuer Termin bei {g['firma']} am {neu_txt}",
                f"Guten Tag {g['name']},\n\nIhr Termin wurde verschoben auf: {neu_txt}.\n{g['kal_ort'] or ''}\n\nDie neue Kalendereinladung hängt an. Falls Sie nicht können, sagen Sie bitte hier ab:\n{absagen}\n\nBis bald!\n\nViele Grüße\n{g['firma']}",
                None, ("termin.ics", sd.ics(str(g["id"]), start, start + dauer, "Termin bei " + g["firma"], g["kal_ort"] or ""), "text/calendar"), "termin")
    return _seite(f"Ihr Termin ist jetzt am {neu_txt}. Eine Bestätigung kommt per E-Mail.")


@router.get("/site/{site_id}/bewertung/abmelden", include_in_schema=False)
def oeff_bewertung_abmelden(site_id: str, t: str, request: Request):
    """Abmeldelink in der Bewertungsanfrage: keine weitere Bewertungsmail an diese Adresse."""
    rate_limit("sitebewab:" + client_ip(request), 20, 600)
    if not re.match(r"^mk_[A-Za-z0-9]{6,40}$", site_id) or not 10 <= len(t) <= 100:
        return _seite("Dieser Link ist ungültig.")
    with tx("system") as cur:
        cur.execute("SELECT w.kunden_id, w.email FROM website_termine w JOIN website_dienste d ON d.kunden_id=w.kunden_id WHERE w.token=%s AND d.site_id=%s", (t, site_id))
        g = cur.fetchone()
        if not g:
            return _seite("Dieser Link ist nicht mehr gültig.")
        cur.execute("INSERT INTO website_bewertung_sperre (kunden_id, email) VALUES (%s, lower(%s)) ON CONFLICT DO NOTHING", (g["kunden_id"], g["email"]))
        cur.execute("UPDATE website_termine SET bewertung_faellig=NULL WHERE kunden_id=%s AND lower(email)=lower(%s)", (g["kunden_id"], g["email"]))
    return _seite("Sie erhalten keine Bewertungsanfragen mehr von uns.")


@router.get("/site/{site_id}/bewertungen", include_in_schema=False)
def oeff_bewertungen(site_id: str, request: Request):
    z = _site(site_id)
    kopf = _cors(request, z)
    rate_limit("sitebew:" + client_ip(request), 60, 300)
    if not z["bew_aktiv"]:
        return _antwort({"aktiv": False}, kopf)
    return _antwort({"aktiv": True, **sd.bewertungen_fuer(z)}, kopf)


# =============================================================== Kunde (nur Inhaber)
kunde = APIRouter(prefix="/api/kunde/website-dienste")


def require_termine(request: Request, user: dict = Depends(current_user)) -> dict:
    """Kalender, Termine von Hand und Stammkunden: Inhaber, Mitarbeiter und Empfang (nur lesen darf zusehen). Ist der Zugang mit einer Person verknüpft, sieht er nur deren Termine."""
    if user["rolle"] != "kunde" or not user["kunden_id"]:
        raise HTTPException(403, "Nur für Kunden.")
    if user["kunden_rolle"] == "nur_lesen" and request.method not in ("GET", "HEAD"):
        raise HTTPException(403, "Dein Zugang ist nur zum Ansehen freigegeben.")
    return user


def _status(cur, kid) -> dict:
    cur.execute("""SELECT art, count(*) FILTER (WHERE ok) AS ok, count(*) FILTER (WHERE NOT ok) AS fehler, max(zeit) AS zuletzt
                   FROM website_ereignisse WHERE kunden_id=%s AND zeit > now() - interval '30 days' GROUP BY art""", (kid,))
    out = {r["art"]: {"ok": r["ok"], "fehler": r["fehler"], "zuletzt": r["zuletzt"]} for r in cur.fetchall()}
    cur.execute("SELECT DISTINCT ON (art) art, ok, hinweis, zeit FROM website_ereignisse WHERE kunden_id=%s ORDER BY art, zeit DESC", (kid,))
    for r in cur.fetchall():
        out.setdefault(r["art"], {})["letzte"] = {"ok": r["ok"], "hinweis": r["hinweis"], "zeit": r["zeit"]}
    return out


def _zustand(cur, kid) -> dict:
    z = sd.zeile(cur, kid)
    cur.execute("SELECT count(*) AS n FROM website_termine WHERE kunden_id=%s AND status IN ('gebucht','reserviert') AND start > now()", (kid,))
    offen = cur.fetchone()["n"]
    places = bool(einstellungen.lesen(cur, "places").get("api_key"))
    hosts = sorted(h for h in sd.erlaubte_hosts(cur, kid, ROOT) if not h.startswith("demo.") and not h.startswith("www."))
    cur.execute("SELECT firma, strasse, plz, ort, telefon, email FROM kunden WHERE id=%s", (kid,))
    kz = cur.fetchone() or {}
    vorschlag = "\n".join(x for x in ("Mit freundlichen Grüßen", kz.get("firma") or "", kz.get("strasse") or "", " ".join(y for y in (kz.get("plz"), kz.get("ort")) if y),
                                      ("Tel. " + kz["telefon"]) if kz.get("telefon") else "", kz.get("email") or "") if x)
    smtp = {}
    if z["smtp_enc"]:
        c = json.loads(einstellungen.entschluesseln(z["smtp_enc"]) or "{}")
        smtp = {"gesetzt": True, "absender": z["smtp_absender"], "host": c.get("host", ""), "port": c.get("port", 587), "sicherung": c.get("sicherung", "starttls"),
                "benutzer": c.get("benutzer", ""), "geprueft_am": z["smtp_geprueft_am"]}
    return {
        "site_id": z["site_id"], "domains": hosts, "basis": config.BASE_URL, "places_verfuegbar": places,
        "mail": {"aktiv": z["mail_aktiv"], "empfaenger": z["mail_empfaenger"] or "", "betreff": z["mail_betreff"] or "", "name": z["mail_name"] or "", "signatur": z["mail_signatur"] or "", "signatur_vorschlag": vorschlag, "bewertung": z["mail_bewertung"], "bewertung_link_da": bool(z["bew_link"] or z["bew_place_id"]), "facebook": z["mail_facebook"] or "", "instagram": z["mail_instagram"] or "", "hinweis": z["mail_hinweis"] or "", "smtp": smtp},
        "fotos": {"aktiv": z["foto_aktiv"], "tage": z["foto_tage"]},
        "kalender": {"aktiv": z["kal_aktiv"], "ical_gesetzt": bool(z["kal_ical_enc"]), "ical_geprueft_am": z["kal_ical_geprueft_am"], "dauer": z["kal_dauer"],
                     "tage": z["kal_tage"], "von": z["kal_von"], "bis": z["kal_bis"], "wochentage": z["kal_wochentage"], "vorlauf_std": z["kal_vorlauf_std"],
                     "ort": z["kal_ort"] or "", "empfaenger": z["kal_empfaenger"] or "", "offene_termine": offen,
                     "zeiten": z["kal_zeiten"] or "", "kapazitaet": z["kal_kapazitaet"], "freigabe": z["kal_freigabe"], "halte_std": z["kal_halte_std"],
                     "plan": sd.plan_text(z.get("kal_wochenplan")), "erinnerung": z["erinnerung_aktiv"], "erinnerung_std1": z["erinnerung_std1"], "erinnerung_std2": z["erinnerung_std2"]},
        "paypal": {"aktiv": z["pp_aktiv"], "client_id": z["pp_client_id"] or "", "waehrung": z["pp_waehrung"], "me_link": z["pp_me_link"] or ""},
        "analytics": {"aktiv": z["ga_aktiv"], "id": z["ga_id"] or "", "statistik": z["stat_aktiv"]},
        "bewertungen": {"aktiv": z["bew_aktiv"], "place_id": z["bew_place_id"] or "", "sterne": float(z["bew_sterne"]) if z["bew_sterne"] is not None else None,
                        "anzahl": z["bew_anzahl"], "link": z["bew_link"] or "", "stand": z["bew_stand"],
                        "nach_termin": z["bew_nach_termin"], "verzoegerung_std": z["bew_verzoegerung_std"]},
        "status": _status(cur, kid),
        "start": webbetrieb.start_checkliste(cur, kid),
    }


@kunde.get("")
def k_lesen(user: dict = Depends(require_inhaber)):
    with tx("system") as cur:
        return _zustand(cur, user["kunden_id"])


@kunde.get("/anbieter")
def k_anbieter(email: str, user: dict = Depends(require_inhaber)):
    return sd.vorschlag(email)


def _empf_pruefen(text: str, pflicht: bool) -> str:
    liste = _empfaenger(text)
    if pflicht and not liste:
        raise HTTPException(400, "Bitte trage mindestens eine E-Mail-Adresse für den Empfang ein.")
    if len(liste) > 5 or any(not EMAIL_RE.match(a) for a in liste):
        raise HTTPException(400, "Eine der Empfänger-Adressen ist ungültig (höchstens 5, mit Komma getrennt).")
    return ", ".join(liste)


class MailIn(BaseModel):
    aktiv: bool = False
    empfaenger: str = Field(default="", max_length=400)
    betreff: str = Field(default="", max_length=120)
    name: str = Field(default="", max_length=80)
    signatur: str = Field(default="", max_length=600)
    bewertung: bool = False
    facebook: str = Field(default="", max_length=300)
    instagram: str = Field(default="", max_length=300)
    hinweis: str = Field(default="", max_length=600)


def _link_pruefen(u: str, name: str) -> str | None:
    u = u.strip()
    if not u:
        return None
    if not re.fullmatch(r"https://[^\s<>\"']{4,290}", u):
        raise HTTPException(400, f"Der Link zu {name} muss mit https:// beginnen.")
    return u


@kunde.put("/mail")
def k_mail(data: MailIn, user: dict = Depends(require_inhaber)):
    emp = _empf_pruefen(data.empfaenger, data.aktiv)
    with tx("system") as cur:
        sd.zeile(cur, user["kunden_id"])
        cur.execute("UPDATE website_dienste SET mail_aktiv=%s, mail_empfaenger=%s, mail_betreff=%s, mail_name=%s, mail_signatur=%s, mail_bewertung=%s, mail_facebook=%s, mail_instagram=%s, mail_hinweis=%s, geaendert_am=now() WHERE kunden_id=%s",
                    (data.aktiv, emp or None, data.betreff.strip() or None, data.name.strip() or None, "\n".join(x.rstrip() for x in data.signatur.strip().splitlines()) or None,
                     data.bewertung, _link_pruefen(data.facebook, "Facebook"), _link_pruefen(data.instagram, "Instagram"), data.hinweis.strip() or None, user["kunden_id"]))
        audit(cur, user, "website_dienste_mail", user["kunden_id"], {"aktiv": data.aktiv})
    return {"ok": True}


class FotoIn(BaseModel):
    aktiv: bool = False
    tage: int = Field(default=fotos.STANDARD_TAGE, ge=7, le=365)


@kunde.put("/fotos")
def k_fotos(data: FotoIn, user: dict = Depends(require_inhaber)):
    """Foto-Upload für Anfragen und Termine ein/aus und Löschfrist (Tage). Gilt für neue Fotos; bereits gespeicherte behalten ihr Löschdatum."""
    with tx("system") as cur:
        sd.zeile(cur, user["kunden_id"])
        cur.execute("UPDATE website_dienste SET foto_aktiv=%s, foto_tage=%s, geaendert_am=now() WHERE kunden_id=%s", (data.aktiv, data.tage, user["kunden_id"]))
        audit(cur, user, "website_dienste_fotos", user["kunden_id"], {"aktiv": data.aktiv, "tage": data.tage})
    return {"ok": True}


class SmtpIn(BaseModel):
    absender: str = Field(max_length=200)
    host: str = Field(min_length=3, max_length=253)
    port: int = Field(ge=1, le=65535)
    sicherung: str = Field(default="starttls", pattern="^(starttls|ssl|keine)$")
    benutzer: str = Field(default="", max_length=200)
    passwort: str = Field(default="", max_length=300)


@kunde.put("/smtp")
def k_smtp(data: SmtpIn, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("smtppruef:" + str(kid), 12, 600)
    if not EMAIL_RE.match(data.absender.strip()):
        raise HTTPException(400, "Die Absender-Adresse ist ungültig.")
    if data.sicherung == "keine" and not config.SITE_LOKAL_ERLAUBT:
        raise HTTPException(400, "Bitte wähle STARTTLS oder SSL, unverschlüsselt wird nicht versendet.")
    with tx("system") as cur:
        z = sd.zeile(cur, kid)
        alt = json.loads(einstellungen.entschluesseln(z["smtp_enc"]) or "{}")
    c = {"host": data.host.strip().lower(), "port": data.port, "sicherung": data.sicherung, "benutzer": data.benutzer.strip() if data.sicherung == "keine" else (data.benutzer.strip() or data.absender.strip()),
         "passwort": data.passwort or alt.get("passwort", ""), "absender": data.absender.strip()}
    if not c["passwort"]:
        raise HTTPException(400, "Bitte trage das Passwort ein.")
    fehler = sd.smtp_pruefen(c)
    if fehler:
        raise HTTPException(400, fehler)
    with tx("system") as cur:
        cur.execute("UPDATE website_dienste SET smtp_enc=%s, smtp_absender=%s, smtp_geprueft_am=now(), geaendert_am=now() WHERE kunden_id=%s",
                    (einstellungen.verschluesseln(json.dumps(c)), c["absender"], kid))
        audit(cur, user, "website_dienste_smtp", kid, {"host": c["host"]})
    return {"ok": True}


@kunde.delete("/smtp")
def k_smtp_weg(user: dict = Depends(require_inhaber)):
    with tx("system") as cur:
        cur.execute("UPDATE website_dienste SET smtp_enc=NULL, smtp_absender=NULL, smtp_geprueft_am=NULL, geaendert_am=now() WHERE kunden_id=%s", (user["kunden_id"],))
        audit(cur, user, "website_dienste_smtp_entfernt", user["kunden_id"], {})
    return {"ok": True}


class TestIn(BaseModel):
    an: str = Field(max_length=200)


@kunde.post("/mail/test")
def k_mail_test(data: TestIn, user: dict = Depends(require_inhaber)):
    """Schickt eine Testmail auf dem Weg, den auch die Website nutzt (eigenes Postfach, sonst Markentrix)."""
    kid = user["kunden_id"]
    rate_limit("mailtest:" + str(kid), 8, 600)
    if not EMAIL_RE.match(data.an.strip()):
        raise HTTPException(400, "Die E-Mail-Adresse ist ungültig.")
    ok, via = sd.kundenmail(kid, data.an.strip(), "Testnachricht von Ihrer Website",
                          "Guten Tag,\n\ndas ist eine Testnachricht. Das Kontaktformular Ihrer Website kann E-Mails versenden. So sehen Mails an Ihre Kunden aus: mit Ihrem Logo (falls hinterlegt) und Ihrer Signatur.\n\n"
                          "Beispiel für Angaben:\nTermin: 27.10.2026 um 08:00 Uhr\nLeistung: Beispielleistung\n\nMit freundlichen Grüßen\n" + "Ihr Team", art="test")
    if not ok:
        raise HTTPException(400, "Der Versand ist fehlgeschlagen: " + via)
    return {"ok": True, "ueber": via}


class VorschauIn(BaseModel):
    art: str = Field(default="bestaetigung", pattern="^(bestaetigung|eingang)$")
    signatur: str = Field(default="", max_length=600)
    bewertung: bool = False
    facebook: str = Field(default="", max_length=300)
    instagram: str = Field(default="", max_length=300)
    hinweis: str = Field(default="", max_length=600)


@kunde.post("/mail/vorschau")
def k_mail_vorschau(data: VorschauIn, user: dict = Depends(require_inhaber)):
    """So sieht eine Mail an deine Kunden aus (mit den Werten aus dem Formular, auch ungespeichert). Es wird nichts verschickt."""
    kid = user["kunden_id"]
    rate_limit("mailvorschau:" + str(kid), 60, 600)
    ueber = {"signatur": "\n".join(x.rstrip() for x in data.signatur.strip().splitlines()) or None, "bewertung": data.bewertung,
             "facebook": _link_pruefen(data.facebook, "Facebook"), "instagram": _link_pruefen(data.instagram, "Instagram"), "hinweis": data.hinweis.strip() or None}
    vorlage, z, _name = sd.vorlage_fuer(kid, ueber)
    firma, wann = vorlage["firma"], (datetime.now(TZ) + timedelta(days=5)).replace(hour=9, minute=0).strftime("%d.%m.%Y um %H:%M Uhr")
    if data.art == "bestaetigung":
        betreff = f"Ihr Termin bei {firma} am {wann} ist bestätigt"
        text = (f"Guten Tag Maria Beispiel,\n\nIhr Termin ist bestätigt: {wann}.\n{z.get('kal_ort') or ''}\n\nDie Kalendereinladung hängt an. Falls Sie nicht können, sagen Sie bitte hier ab:\n"
                f"{config.BASE_URL}/site/{z['site_id']}/termin/absagen?t=beispiel1234567890\n\nBis bald!\n\nViele Grüße\n{firma}")
    else:
        betreff = "Eingangsbestätigung: Terminanfrage (T261024-0000)"
        text = (f"Guten Tag Maria Beispiel,\n\nvielen Dank, Ihre Terminanfrage ist bei uns eingegangen. Ihre Vorgangsnummer: T261024-0000.\nIhr Wunschtermin {wann} ist für Sie vorgemerkt. Wir prüfen ihn und melden uns mit der Bestätigung.\n\n"
                "Ihre Angaben:\nLeistung: Beispielleistung\nTelefon: 0511 123456\n\nMit freundlichen Grüßen\n" + firma)
    logo = vorlage.get("logo")
    html = mailvorlage.html_mail(firma, text, vorlage.get("signatur"), vorlage.get("farbe"), "logo@mail" if logo else None, vorlage.get("extras"))
    if logo:
        html = html.replace("cid:logo@mail", f"data:{logo[1]}/{logo[2]};base64," + base64.b64encode(logo[0]).decode())
    return {"betreff": betreff, "html": html, "logo": bool(logo)}


class KalenderIn(BaseModel):
    aktiv: bool = False
    ical_url: str = Field(default="", max_length=1500)          # leer = unverändert
    dauer: int = Field(default=30, ge=10, le=240)
    tage: int = Field(default=14, ge=1, le=90)
    von: str = Field(default="09:00", pattern=r"^\d{2}:\d{2}$")
    bis: str = Field(default="17:00", pattern=r"^\d{2}:\d{2}$")
    wochentage: list[int] = Field(default=[1, 2, 3, 4, 5], max_length=7)
    vorlauf_std: int = Field(default=12, ge=0, le=240)
    ort: str = Field(default="", max_length=300)
    empfaenger: str = Field(default="", max_length=400)
    zeiten: str = Field(default="", max_length=300)             # feste Uhrzeiten "08:00, 10:00"; leer = Raster
    kapazitaet: int = Field(default=0, ge=0, le=200)
    freigabe: bool = False
    halte_std: int = Field(default=72, ge=1, le=336)
    erinnerung: bool = True                                      # Erinnerungs-Mails an die Besucher
    erinnerung_std1: int = Field(default=24, ge=0, le=168)       # 0 = aus
    erinnerung_std2: int = Field(default=2, ge=0, le=48)
    plan: dict[str, str] = Field(default_factory=dict, max_length=7)   # Arbeitszeiten je Wochentag, z. B. {"1": "09:00-12:00, 13:00-18:00"}; leer = von/bis/Wochentage


@kunde.put("/kalender")
def k_kalender(data: KalenderIn, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("kalpruef:" + str(kid), 12, 600)
    try:
        plan = sd.plan_lesen(data.plan) if data.plan else {}
    except ValueError as e:
        raise HTTPException(400, str(e))
    if data.plan and not plan:
        raise HTTPException(400, "Bitte mindestens einen Tag mit Arbeitszeiten eintragen.")
    if plan:
        data.wochentage, data.von, data.bis = sd.plan_grenzen(plan)
    if data.von >= data.bis:
        raise HTTPException(400, "Das Ende muss nach dem Beginn liegen.")
    if not data.wochentage or any(w < 1 or w > 7 for w in data.wochentage):
        raise HTTPException(400, "Bitte wähle mindestens einen Wochentag.")
    emp = _empf_pruefen(data.empfaenger, False)
    zeiten = [x.strip() for x in re.split(r"[,;\s]+", data.zeiten) if x.strip()]
    if any(not sd.ZEIT_RE.match(x) for x in zeiten) or len(zeiten) > 24:
        raise HTTPException(400, "Die Uhrzeiten bitte so eintragen: 08:00, 10:00, 13:30.")
    enc, geprueft = None, None
    if data.ical_url.strip():
        try:
            u = sd.ical_url_pruefen(data.ical_url)
            sd.ical_belegt(sd.ical_laden(u), datetime.now(timezone.utc), datetime.now(timezone.utc) + timedelta(days=30))
        except ValueError as e:
            raise HTTPException(400, str(e))
        enc, geprueft = einstellungen.verschluesseln(u), datetime.now(timezone.utc)
    with tx("system") as cur:
        sd.zeile(cur, kid)
        cur.execute("""UPDATE website_dienste SET kal_aktiv=%s, kal_dauer=%s, kal_tage=%s, kal_von=%s, kal_bis=%s, kal_wochentage=%s, kal_vorlauf_std=%s,
                       kal_ort=%s, kal_empfaenger=%s, kal_ical_enc=coalesce(%s, kal_ical_enc), kal_ical_geprueft_am=coalesce(%s, kal_ical_geprueft_am),
                       kal_zeiten=%s, kal_kapazitaet=%s, kal_freigabe=%s, kal_halte_std=%s, kal_wochenplan=%s::jsonb,
                       erinnerung_aktiv=%s, erinnerung_std1=%s, erinnerung_std2=%s, geaendert_am=now()
                       WHERE kunden_id=%s""",
                    (data.aktiv, data.dauer, data.tage, data.von, data.bis, sorted(set(data.wochentage)), data.vorlauf_std, data.ort.strip() or None, emp or None, enc, geprueft,
                     ", ".join(sorted(set(zeiten))) or None, data.kapazitaet, data.freigabe, data.halte_std, json.dumps(plan) if plan else None,
                     data.erinnerung, data.erinnerung_std1, data.erinnerung_std2, kid))
        audit(cur, user, "website_dienste_kalender", kid, {"aktiv": data.aktiv})
    sd._ical_cache.pop(str(kid), None)
    return {"ok": True}


@kunde.delete("/kalender-adresse")
def k_kal_weg(user: dict = Depends(require_inhaber)):
    with tx("system") as cur:
        cur.execute("UPDATE website_dienste SET kal_ical_enc=NULL, kal_ical_geprueft_am=NULL, geaendert_am=now() WHERE kunden_id=%s", (user["kunden_id"],))
        audit(cur, user, "website_dienste_kalender_entfernt", user["kunden_id"], {})
    sd._ical_cache.pop(str(user["kunden_id"]), None)
    return {"ok": True}


@kunde.post("/kalender/test")
def k_kal_test(user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("kaltest:" + str(kid), 12, 600)
    with tx("system") as cur:
        z = sd.zeile(cur, kid)
        if z["kal_ical_enc"]:
            try:
                n = len(sd.kalender_belegt(z, datetime.now(timezone.utc), datetime.now(timezone.utc) + timedelta(days=z["kal_tage"])))
            except ValueError as e:
                raise HTTPException(400, str(e))
        else:
            n = None
        slots = sd.freie_slots(cur, z)
    erste = [s.astimezone(TZ).strftime("%d.%m. %H:%M") for s in slots[:4]]
    return {"ok": True, "belegt": n, "freie": len(slots), "beispiele": erste}


class PayPalIn(BaseModel):
    aktiv: bool = False
    client_id: str = Field(default="", max_length=140)
    waehrung: str = Field(default="EUR", pattern="^(EUR|CHF|USD|GBP|PLN|CZK|DKK|SEK|NOK)$")
    me_link: str = Field(default="", max_length=200)


@kunde.put("/paypal")
def k_paypal(data: PayPalIn, user: dict = Depends(require_inhaber)):
    cid, link = data.client_id.strip(), data.me_link.strip()
    if cid and not re.match(r"^[A-Za-z0-9_-]{20,140}$", cid):
        raise HTTPException(400, "Die Client-ID sieht ungültig aus. Sie ist eine lange Zeichenfolge, die oft mit „A“ beginnt.")
    if link and not re.match(r"^https://(www\.)?(paypal\.me|paypal\.com/paypalme)/[A-Za-z0-9._/-]{2,80}$", link):
        raise HTTPException(400, "Der PayPal.Me-Link muss so aussehen: https://paypal.me/deinname")
    if data.aktiv and not (cid or link):
        raise HTTPException(400, "Bitte trage die Client-ID oder deinen PayPal.Me-Link ein.")
    with tx("system") as cur:
        sd.zeile(cur, user["kunden_id"])
        cur.execute("UPDATE website_dienste SET pp_aktiv=%s, pp_client_id=%s, pp_waehrung=%s, pp_me_link=%s, geaendert_am=now() WHERE kunden_id=%s",
                    (data.aktiv, cid or None, data.waehrung, link or None, user["kunden_id"]))
        audit(cur, user, "website_dienste_paypal", user["kunden_id"], {"aktiv": data.aktiv})
    return {"ok": True}


class AnalyticsIn(BaseModel):
    aktiv: bool = False
    id: str = Field(default="", max_length=20)


@kunde.put("/analytics")
def k_analytics(data: AnalyticsIn, user: dict = Depends(require_inhaber)):
    gid = data.id.strip().upper()
    if gid and not re.match(r"^G-[A-Z0-9]{6,14}$", gid):
        raise HTTPException(400, "Die Mess-ID sieht ungültig aus. Sie beginnt mit „G-“, zum Beispiel G-AB12CD34EF.")
    if data.aktiv and not gid:
        raise HTTPException(400, "Bitte trage die Mess-ID ein.")
    with tx("system") as cur:
        sd.zeile(cur, user["kunden_id"])
        cur.execute("UPDATE website_dienste SET ga_aktiv=%s, ga_id=%s, geaendert_am=now() WHERE kunden_id=%s", (data.aktiv, gid or None, user["kunden_id"]))
        audit(cur, user, "website_dienste_analytics", user["kunden_id"], {"aktiv": data.aktiv})
    return {"ok": True}


class BewertungIn(BaseModel):
    aktiv: bool = False
    place_id: str = Field(default="", max_length=200)
    sterne: float | None = Field(default=None, ge=1, le=5)
    anzahl: int | None = Field(default=None, ge=0, le=1000000)
    link: str = Field(default="", max_length=400)


@kunde.put("/bewertungen")
def k_bewertungen(data: BewertungIn, user: dict = Depends(require_inhaber)):
    pid, link = data.place_id.strip(), data.link.strip()
    if pid and not sd.PLACE_ID.match(pid):
        raise HTTPException(400, "Die Place-ID sieht ungültig aus. Sie beginnt meist mit „ChIJ“.")
    if link and not re.match(r"^https://[^\s]{4,380}$", link):
        raise HTTPException(400, "Der Bewertungs-Link muss mit https:// beginnen.")
    if data.aktiv and not (pid or data.sterne):
        raise HTTPException(400, "Bitte trage die Place-ID oder Sterne und Anzahl ein.")
    with tx("system") as cur:
        sd.zeile(cur, user["kunden_id"])
        cur.execute("""UPDATE website_dienste SET bew_aktiv=%s, bew_place_id=%s, bew_sterne=%s, bew_anzahl=%s, bew_link=%s, bew_cache=NULL, bew_stand=NULL, geaendert_am=now()
                       WHERE kunden_id=%s""", (data.aktiv, pid or None, data.sterne, data.anzahl, link or None, user["kunden_id"]))
        audit(cur, user, "website_dienste_bewertungen", user["kunden_id"], {"aktiv": data.aktiv})
    return {"ok": True}


class BewAnfrageIn(BaseModel):
    aktiv: bool = False
    verzoegerung_std: int = Field(default=3, ge=1, le=72)


@kunde.put("/bewertungen/anfrage")
def k_bew_anfrage(data: BewAnfrageIn, user: dict = Depends(require_inhaber)):
    """Opt-in: nach einem als „erschienen“ markierten Termin bittet die Website den Besucher per Mail um eine Google-Bewertung."""
    kid = user["kunden_id"]
    rate_limit("bewanfrage:" + str(kid), 30, 600)
    with tx("system") as cur:
        z = sd.zeile(cur, kid)
        if data.aktiv and not (z["bew_link"] or z["bew_place_id"]):
            raise HTTPException(400, "Bitte trage zuerst die Place-ID oder den Link „Rezension schreiben“ ein und speichere.")
        cur.execute("UPDATE website_dienste SET bew_nach_termin=%s, bew_verzoegerung_std=%s, geaendert_am=now() WHERE kunden_id=%s", (data.aktiv, data.verzoegerung_std, kid))
        audit(cur, user, "website_dienste_bewertungsanfrage", kid, {"aktiv": data.aktiv})
    return {"ok": True}


@kunde.post("/bewertungen/test")
def k_bew_test(user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("bewtest:" + str(kid), 10, 600)
    with tx("system") as cur:
        z = sd.zeile(cur, kid)
        key = einstellungen.lesen(cur, "places").get("api_key", "")
    if not z["bew_place_id"]:
        raise HTTPException(400, "Bitte speichere zuerst die Place-ID.")
    if not key:
        raise HTTPException(400, "Die automatische Anzeige ist bei Markentrix noch nicht freigeschaltet. Trage solange Sterne und Anzahl von Hand ein.")
    try:
        r = sd.places_holen(key, z["bew_place_id"])
    except ValueError as e:
        raise HTTPException(400, str(e))
    with tx("system") as cur:
        cur.execute("UPDATE website_dienste SET bew_cache=%s::jsonb, bew_stand=now() WHERE kunden_id=%s", (json.dumps(r), kid))
    return {"ok": True, "sterne": r["sterne"], "anzahl": r["anzahl"]}


# =============================================================== Admin: Status ohne Geheimnisse, Google-Schlüssel
admin = APIRouter(prefix="/api/admin")


# =============================================================== Buchungen, Sperrtage, Statistik (Inhaber)
@kunde.get("/buchungen")
def k_buchungen(user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    with tx("system") as cur:
        z = sd.zeile(cur, kid)
        sd.reservierungen_verfallen(cur, kid, z["kal_halte_std"])
        cur.execute("""SELECT w.id, w.ref, w.status, w.start, w.ende, w.name, w.email, w.telefon, w.nachricht, w.daten, w.anzahlung, w.anzahlung_betrag, w.preis_betrag, w.anzahlung_erhalten_am, w.bezahlt_am, w.quittung_am, w.erstellt_am, w.ergebnis,
                              l.name AS leistung, m.name AS mitarbeiter,
                              (SELECT count(*) FROM website_termine x WHERE x.kunden_id=w.kunden_id AND lower(x.email)=lower(w.email) AND x.ergebnis='nicht_erschienen') AS fehlt_bisher
                       FROM website_termine w LEFT JOIN website_leistungen l ON l.id=w.leistung_id LEFT JOIN website_mitarbeiter m ON m.id=w.mitarbeiter_id
                       WHERE w.kunden_id=%s AND ((w.status IN ('reserviert','gebucht') AND w.start > now() - interval '1 day')
                                                 OR (w.status='gebucht' AND w.ergebnis IS NULL AND w.start <= now() AND w.start > now() - interval '14 days'))
                       ORDER BY (w.status='reserviert') DESC, w.start LIMIT 150""", (kid,))
        zeilen = cur.fetchall()
        bilder = fotos.liste(cur, kid, buchung_ids=[r["id"] for r in zeilen])["buchung"]
    out, vorbei = [], []
    jetzt = datetime.now(timezone.utc)
    for r in zeilen:
        if r["start"] <= jetzt and r["ergebnis"] is not None:
            continue
        d = r["daten"] or {}
        details = {k: v for k, v in d.items() if k not in ("typ", "name", "telefon", "email", "nachricht", "datum", "datum2", "uhrzeit") and v}
        eintrag = {"id": str(r["id"]), "ref": r["ref"] or "", "status": r["status"], "start": r["start"], "name": r["name"], "email": r["email"], "telefon": r["telefon"] or "",
                   "nachricht": r["nachricht"] or "", "details": details, "anzahlung": r["anzahlung"] or "", "preis": f"{r['preis_betrag']:.2f}" if r["preis_betrag"] is not None else "", "anzahlung_erhalten_am": r["anzahlung_erhalten_am"], "bezahlt_am": r["bezahlt_am"], "quittung_am": r["quittung_am"], "leistung": r["leistung"] or "", "mitarbeiter": r["mitarbeiter"] or "",
                   "ende": r["ende"], "fehlt_bisher": r["fehlt_bisher"], "fotos": bilder.get(str(r["id"]), [])}
        (vorbei if r["start"] <= jetzt and r["status"] == "gebucht" and r["ergebnis"] is None else out).append(eintrag)
    return {"buchungen": out, "vergangen": vorbei, "freigabe": z["kal_freigabe"]}


# ---------------------------------------------------------------- Excel-Übersicht
def _zeitraum(art: str, von: str, bis: str) -> tuple[date, date, str]:
    heute = datetime.now(TZ).date()
    if art == "woche":
        a = heute - timedelta(days=heute.weekday())
        return a, a + timedelta(days=6), "Diese Woche"
    if art == "vorwoche":
        a = heute - timedelta(days=heute.weekday() + 7)
        return a, a + timedelta(days=6), "Letzte Woche"
    if art in ("monat", "vormonat"):
        a = heute.replace(day=1)
        if art == "vormonat":
            a = (a - timedelta(days=1)).replace(day=1)
        n = (a + timedelta(days=32)).replace(day=1)
        return a, n - timedelta(days=1), a.strftime("%m/%Y")
    try:
        a, b = date.fromisoformat(von), date.fromisoformat(bis)
    except ValueError:
        raise HTTPException(400, "Bitte Von- und Bis-Datum angeben (Format JJJJ-MM-TT).")
    if b < a or (b - a).days > 400:
        raise HTTPException(400, "Der Zeitraum ist ungültig (höchstens 400 Tage).")
    return a, b, f"{a.strftime('%d.%m.%Y')} bis {b.strftime('%d.%m.%Y')}"


@kunde.get("/uebersicht.xlsx")
def k_uebersicht(zeitraum: str = "monat", von: str = "", bis: str = "", user: dict = Depends(require_inhaber)):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    kid = user["kunden_id"]
    rate_limit("uebersicht:" + str(kid), 30, 600)
    a, b, titel = _zeitraum(zeitraum, von, bis)
    ab, bis_ = datetime.combine(a, datetime.min.time(), TZ), datetime.combine(b + timedelta(days=1), datetime.min.time(), TZ)
    with tx("system") as cur:
        cur.execute("SELECT firma FROM kunden WHERE id=%s", (kid,))
        firma = cur.fetchone()["firma"]
        cur.execute("""SELECT w.*, l.name AS leistung, m.name AS mitarbeiter FROM website_termine w LEFT JOIN website_leistungen l ON l.id=w.leistung_id
                       LEFT JOIN website_mitarbeiter m ON m.id=w.mitarbeiter_id WHERE w.kunden_id=%s AND w.start >= %s AND w.start < %s ORDER BY w.start LIMIT 5000""", (kid, ab, bis_))
        termine = cur.fetchall()
        cur.execute("SELECT * FROM website_gutscheine WHERE kunden_id=%s AND erstellt_am >= %s AND erstellt_am < %s ORDER BY erstellt_am LIMIT 2000", (kid, ab, bis_))
        gutscheine = cur.fetchall()
    kopf_stil, fuell = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")
    wb = Workbook()
    ws = wb.active
    ws.title = "Termine"
    ws.append(["Datum", "Uhrzeit", "Name", "E-Mail", "Telefon", "Leistung", "Mitarbeiter", "Status", "Preis (€)", "Anzahlung (€)", "Anzahlung erhalten", "Komplett bezahlt", "Rest offen (€)", "Erstattung", "Vorgang"])
    summe = {"gebucht": 0, "abgesagt": 0, "nicht_erschienen": 0, "anz_soll": 0.0, "anz_da": 0.0, "preis": 0.0, "bezahlt": 0}
    nach_leistung: dict = {}
    for t in termine:
        lokal = t["start"].astimezone(TZ)
        status = {"gebucht": "gebucht", "reserviert": "wartet auf Bestätigung", "abgesagt": "abgesagt"}.get(t["status"], t["status"])
        if t["ergebnis"] == "nicht_erschienen":
            status = "nicht erschienen"
        if t["status"] == "abgesagt":
            status = "abgesagt" if not t.get("abgesagt_von") else ("abgesagt (kostenlos)" if t.get("storno_kostenlos") else "zu spät abgesagt (voller Betrag)")
        erstattung = ""
        if t["status"] == "abgesagt" and t["anzahlung_erhalten_am"] and t.get("storno_kostenlos"):
            erstattung = "erstattet" if t.get("erstattet_am") else "offen"
        preis = float(t["preis_betrag"]) if t["preis_betrag"] is not None else None
        anz = float(t["anzahlung_betrag"]) if t["anzahlung_betrag"] is not None else None
        rest = None
        if preis is not None and t["status"] in ("gebucht", "reserviert"):
            rest = 0.0 if t["bezahlt_am"] else max(preis - (anz or 0.0 if t["anzahlung_erhalten_am"] else 0.0), 0.0)
        ws.append([lokal.date(), lokal.strftime("%H:%M"), t["name"], t["email"], t["telefon"] or "", t["leistung"] or "", t["mitarbeiter"] or "", status, preis, anz,
                   t["anzahlung_erhalten_am"].astimezone(TZ).date() if t["anzahlung_erhalten_am"] else ("" if anz else "–"),
                   t["bezahlt_am"].astimezone(TZ).date() if t["bezahlt_am"] else "", rest, erstattung, t["ref"] or ""])
        if t["status"] == "abgesagt":
            summe["abgesagt"] += 1
            continue
        if t["ergebnis"] == "nicht_erschienen":
            summe["nicht_erschienen"] += 1
        summe["gebucht"] += 1
        if anz:
            summe["anz_soll"] += anz
            summe["anz_da"] += anz if t["anzahlung_erhalten_am"] else 0.0
        if preis:
            summe["preis"] += preis
            e = nach_leistung.setdefault(t["leistung"] or "(ohne Leistung)", [0, 0.0])
            e[0] += 1
            e[1] += preis
        summe["bezahlt"] += 1 if t["bezahlt_am"] else 0
    for c in ws[1]:
        c.font, c.fill = kopf_stil, fuell
    for sp, br in zip("ABCDEFGHIJKLMNO", (12, 8, 22, 28, 16, 24, 16, 22, 11, 14, 17, 16, 14, 12, 16)):
        ws.column_dimensions[sp].width = br
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD.MM.YYYY"
        for c in (row[8], row[9], row[12]):
            c.number_format = '#,##0.00 "€"'
        for c in (row[10], row[11]):
            if hasattr(c.value, "year"):
                c.number_format = "DD.MM.YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    zs = wb.create_sheet("Zusammenfassung")
    zs.append([f"{firma}: Übersicht {titel}"])
    zs["A1"].font = Font(bold=True, size=14)
    zs.append([f"Zeitraum: {a.strftime('%d.%m.%Y')} bis {b.strftime('%d.%m.%Y')}. Zahlungen sind nur erfasst, wenn sie im Portal als erhalten markiert wurden."])
    zs.append([])
    zeilen = [("Termine (nicht abgesagt)", summe["gebucht"]), ("Abgesagt", summe["abgesagt"]), ("Nicht erschienen", summe["nicht_erschienen"]), ("Komplett bezahlt (Termine)", summe["bezahlt"]),
              ("Anzahlungen fällig (€)", round(summe["anz_soll"], 2)), ("davon erhalten (€)", round(summe["anz_da"], 2)), ("noch offen (€)", round(summe["anz_soll"] - summe["anz_da"], 2)),
              ("Preise der Termine zusammen (€)", round(summe["preis"], 2))]
    for k, v in zeilen:
        zs.append([k, v])
        if "(€)" in k:
            zs.cell(zs.max_row, 2).number_format = '#,##0.00 "€"'
    zs.append([])
    zs.append(["Leistung", "Anzahl", "Preise zusammen (€)"])
    for c in zs[zs.max_row]:
        c.font, c.fill = kopf_stil, fuell
    for k, (n, summ) in sorted(nach_leistung.items(), key=lambda x: -x[1][1]):
        zs.append([k, n, round(summ, 2)])
        zs.cell(zs.max_row, 3).number_format = '#,##0.00 "€"'
    zs.column_dimensions["A"].width = 44
    zs.column_dimensions["B"].width = 16
    zs.column_dimensions["C"].width = 22
    zs["A2"].alignment = Alignment(wrap_text=False)

    gs = wb.create_sheet("Gutscheine")
    gs.append(["Angefragt am", "Name", "E-Mail", "Betrag (€)", "Status", "Code", "Gesendet am", "Vorgang"])
    for c in gs[1]:
        c.font, c.fill = kopf_stil, fuell
    namen = {"angefragt": "neu", "zahlungslink_gesendet": "wartet auf Zahlung", "gesendet": "Code gesendet", "storniert": "storniert"}
    for g in gutscheine:
        gs.append([g["erstellt_am"].astimezone(TZ).date(), g["name"], g["email"] or "", float(g["betrag"]) if g["betrag"] is not None else None, namen.get(g["status"], g["status"]),
                   g["code"] or "", g["gesendet_am"].astimezone(TZ).date() if g["gesendet_am"] else "", g["ref"]])
        gs.cell(gs.max_row, 1).number_format = "DD.MM.YYYY"
        gs.cell(gs.max_row, 4).number_format = '#,##0.00 "€"'
        if g["gesendet_am"]:
            gs.cell(gs.max_row, 7).number_format = "DD.MM.YYYY"
    for sp, br in zip("ABCDEFGH", (14, 24, 28, 12, 22, 20, 14, 16)):
        gs.column_dimensions[sp].width = br
    ausgabe = BytesIO()
    wb.save(ausgabe)
    with tx("system") as cur:
        audit(cur, user, "website_uebersicht_export", kid, {"von": str(a), "bis": str(b)})
    name = f"Uebersicht-{a.isoformat()}-{b.isoformat()}.xlsx"
    return Response(ausgabe.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ---------------------------------------------------------------- Gutscheine
CODE_RE = re.compile(r"^[A-Za-z0-9ÄÖÜäöüß][A-Za-z0-9ÄÖÜäöüß _.\-/]{2,39}$")


def _gutschein_json(r: dict) -> dict:
    return {"id": str(r["id"]), "ref": r["ref"], "name": r["name"], "email": r["email"] or "", "telefon": r["telefon"] or "",
            "betrag": f"{r['betrag']:.2f}" if r["betrag"] is not None else "", "empfaenger": r["empfaenger"] or "", "widmung": r["widmung"] or "",
            "status": r["status"], "code": r["code"] or "", "erstellt_am": r["erstellt_am"], "zahlungslink_am": r["zahlungslink_am"], "gesendet_am": r["gesendet_am"]}


@kunde.get("/gutscheine")
def k_gutscheine(user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    with tx("system") as cur:
        z = sd.zeile(cur, kid)
        cur.execute("""SELECT * FROM website_gutscheine WHERE kunden_id=%s AND (status<>'storniert' OR erstellt_am > now() - interval '30 days')
                       ORDER BY (status IN ('angefragt','zahlungslink_gesendet')) DESC, erstellt_am DESC LIMIT 100""", (kid,))
        zeilen = cur.fetchall()
    me = (z["pp_me_link"] or "").strip()
    return {"gutscheine": [_gutschein_json(r) for r in zeilen], "paypal_bereit": bool(z["pp_aktiv"] and me.startswith("https://")), "waehrung": z["pp_waehrung"]}


class ZahlungslinkIn(BaseModel):
    betrag: str = Field(default="", max_length=20)


def _gutschein_holen(cur, gid: str, kid):
    try:
        cur.execute("SELECT * FROM website_gutscheine WHERE id=%s AND kunden_id=%s FOR UPDATE", (gid, kid))
    except Exception:
        raise HTTPException(400, "Ungültiger Gutschein.")
    g = cur.fetchone()
    if not g:
        raise HTTPException(404, "Gutschein nicht gefunden.")
    if g["status"] == "storniert":
        raise HTTPException(409, "Diese Anfrage ist storniert.")
    if not g["email"] or not EMAIL_RE.match(g["email"]):
        raise HTTPException(409, "Zu dieser Anfrage gibt es keine E-Mail-Adresse. Bitte den Kunden anrufen.")
    return g


@kunde.post("/gutscheine/{gid}/zahlungslink")
def k_gutschein_link(gid: str, data: ZahlungslinkIn, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("gutlink:" + str(kid), 60, 600)
    with tx("system") as cur:
        z = sd.zeile(cur, kid)
        me = (z["pp_me_link"] or "").strip().rstrip("/")
        if not (z["pp_aktiv"] and me.startswith("https://")):
            raise HTTPException(409, "Bitte zuerst oben bei PayPal deinen PayPal.Me-Link eintragen und PayPal einschalten.")
        g = _gutschein_holen(cur, gid, kid)
        if data.betrag.strip():
            v = sd.betrag_lesen(data.betrag)
        else:
            v = float(g["betrag"]) if g["betrag"] is not None else None
        if v is None or not 1 <= v <= 10000:
            raise HTTPException(400, "Bitte den Betrag als Zahl eintragen, z. B. 50 oder 49,90.")
        betrag = f"{v:.2f}"
        cur.execute("SELECT firma FROM kunden WHERE id=%s", (kid,))
        firma = cur.fetchone()["firma"]
    link = f"{me}/{betrag}{z['pp_waehrung']}"
    fuer = f"\nDer Gutschein ist für: {g['empfaenger']}" if g["empfaenger"] else ""
    text = (f"Guten Tag {g['name']},\n\nvielen Dank für Ihre Gutschein-Anfrage ({g['ref']}) über {betrag.replace('.', ',')} €.{fuer}\n\n"
            f"Sie können den Betrag bequem per PayPal bezahlen:\n{link}\n\n"
            f"Sobald die Zahlung bei uns eingegangen ist, senden wir Ihnen den Gutschein-Code per E-Mail.\n\nViele Grüße\n{firma}")
    ok, via = sd.kundenmail(kid, g["email"], f"Ihr Gutschein bei {firma}: Zahlungslink ({g['ref']})", text, art="kontakt")
    if not ok:
        raise HTTPException(502, "Die E-Mail konnte nicht gesendet werden: " + str(via)[:200])
    with tx("system") as cur:
        cur.execute("""UPDATE website_gutscheine SET betrag=%s, status=CASE WHEN status='angefragt' THEN 'zahlungslink_gesendet' ELSE status END, zahlungslink_am=now()
                       WHERE id=%s AND kunden_id=%s""", (betrag, gid, kid))
        audit(cur, user, "website_gutschein_zahlungslink", kid, {"ref": g["ref"], "betrag": betrag})
    return {"ok": True}


class GutscheinCodeIn(BaseModel):
    code: str = Field(min_length=3, max_length=40)
    hinweis: str = Field(default="", max_length=600)


@kunde.post("/gutscheine/{gid}/code")
def k_gutschein_code(gid: str, data: GutscheinCodeIn, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("gutcode:" + str(kid), 60, 600)
    code = data.code.strip()
    if not CODE_RE.match(code):
        raise HTTPException(400, "Bitte einen Code mit 3 bis 40 Zeichen eintragen (Buchstaben, Zahlen, Bindestrich).")
    with tx("system") as cur:
        g = _gutschein_holen(cur, gid, kid)
        cur.execute("SELECT 1 FROM website_gutscheine WHERE kunden_id=%s AND id<>%s AND lower(code)=lower(%s) AND status='gesendet'", (kid, gid, code))
        if cur.fetchone():
            raise HTTPException(409, "Diesen Code hast du schon an jemand anderen gesendet. Bitte einen neuen Code wählen.")
        cur.execute("SELECT firma FROM kunden WHERE id=%s", (kid,))
        firma = cur.fetchone()["firma"]
    betrag = f" über {g['betrag']:.2f}".replace(".", ",") + " €" if g["betrag"] is not None else ""
    fuer = f"\nFür: {g['empfaenger']}" if g["empfaenger"] else ""
    widmung = f"\nWidmung: {g['widmung']}" if g["widmung"] else ""
    hinweis = f"\n{data.hinweis.strip()}\n" if data.hinweis.strip() else ""
    text = (f"Guten Tag {g['name']},\n\nvielen Dank, Ihre Zahlung ist bei uns eingegangen. Hier ist Ihr Gutschein{betrag}:{fuer}{widmung}\n\n"
            f"Gutschein-Code: {code}\n{hinweis}\nBitte nennen Sie den Code bei Ihrem Besuch oder bei der Buchung.\n\nViele Grüße\n{firma}")
    ok, via = sd.kundenmail(kid, g["email"], f"Ihr Gutschein bei {firma} ({g['ref']})", text, art="kontakt")
    if not ok:
        raise HTTPException(502, "Die E-Mail konnte nicht gesendet werden: " + str(via)[:200])
    with tx("system") as cur:
        cur.execute("UPDATE website_gutscheine SET status='gesendet', code=%s, gesendet_am=now() WHERE id=%s AND kunden_id=%s", (code, gid, kid))
        audit(cur, user, "website_gutschein_code", kid, {"ref": g["ref"]})
    return {"ok": True}


@kunde.post("/gutscheine/{gid}/storno")
def k_gutschein_storno(gid: str, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("gutstorno:" + str(kid), 60, 600)
    with tx("system") as cur:
        try:
            cur.execute("UPDATE website_gutscheine SET status='storniert' WHERE id=%s AND kunden_id=%s AND status IN ('angefragt','zahlungslink_gesendet') RETURNING ref", (gid, kid))
        except Exception:
            raise HTTPException(400, "Ungültiger Gutschein.")
        r = cur.fetchone()
        if not r:
            raise HTTPException(409, "Diese Anfrage kann nicht mehr storniert werden.")
        audit(cur, user, "website_gutschein_storno", kid, {"ref": r["ref"]})
    return {"ok": True}


class ErgebnisIn(BaseModel):
    ergebnis: str = Field(pattern=r"^(erschienen|nicht_erschienen)$")


@kunde.post("/buchungen/{bid}/ergebnis")
def k_ergebnis(bid: str, data: ErgebnisIn, user: dict = Depends(require_termine)):
    """Hat der Termin stattgefunden? Die Antwort steuert den Nichterschienen-Zähler und den Status des Termins im Portal."""
    kid = user["kunden_id"]
    rate_limit("buchergebnis:" + str(kid), 120, 600)
    with tx("system") as cur:
        try:
            cur.execute("""UPDATE website_termine SET ergebnis=%s WHERE id=%s AND kunden_id=%s AND status='gebucht' AND start <= now()
                           AND (%s::uuid IS NULL OR mitarbeiter_id = %s::uuid) RETURNING termin_id""", (data.ergebnis, bid, kid, user.get("website_mitarbeiter_id"), user.get("website_mitarbeiter_id")))
        except Exception:
            raise HTTPException(400, "Ungültige Buchung.")
        g = cur.fetchone()
        if not g:
            raise HTTPException(409, "Diese Buchung lässt sich nicht mehr bewerten (noch nicht vorbei oder abgesagt).")
        if g["termin_id"]:
            cur.execute("UPDATE termine SET status=%s WHERE id=%s", ("erledigt" if data.ergebnis == "erschienen" else "nicht_erschienen", g["termin_id"]))
        if data.ergebnis == "erschienen":
            cur.execute("""UPDATE website_termine w SET bewertung_faellig = now() + make_interval(hours => d.bew_verzoegerung_std)
                           FROM website_dienste d WHERE w.id=%s AND d.kunden_id=w.kunden_id AND d.bew_nach_termin AND w.bewertung_gesendet IS NULL
                             AND w.email IS NOT NULL AND (d.bew_link IS NOT NULL OR d.bew_place_id IS NOT NULL)
                             AND NOT EXISTS (SELECT 1 FROM website_bewertung_sperre s WHERE s.kunden_id=w.kunden_id AND s.email=lower(w.email))""", (bid,))
    return {"ok": True}


class BestaetigenIn(BaseModel):
    anzahlung: str = Field(default="", max_length=20)
    hinweis: str = Field(default="", max_length=600)


def _betrag(text: str) -> str:
    t = text.strip().replace("€", "").replace(",", ".").strip()
    if not t:
        return ""
    try:
        v = float(t)
    except ValueError:
        raise HTTPException(400, "Bitte die Anzahlung als Zahl eintragen, z. B. 30 oder 29,90.")
    if not 0 < v <= 100000:
        raise HTTPException(400, "Bitte die Anzahlung als Zahl eintragen, z. B. 30 oder 29,90.")
    return f"{v:.2f}"


@kunde.post("/buchungen/{bid}/bestaetigen")
def k_bestaetigen(bid: str, data: BestaetigenIn, bg: BackgroundTasks, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("buchbest:" + str(kid), 60, 600)
    return _bestaetigen(kid, bid, _betrag(data.anzahlung), data.hinweis.strip(), bg, user)


def _bestaetigen(kid, bid: str, betrag: str, hinweis: str, bg: BackgroundTasks, user: dict | None):
    """Bestätigt eine vorgemerkte Buchung und schickt dem Gast die Bestätigung. Gemeinsam für das Portal und den Knopf in der Mail (user=None)."""
    with tx("system") as cur:
        z = sd.zeile(cur, kid)
        try:
            cur.execute("""UPDATE website_termine SET status='gebucht', anzahlung=coalesce(%s, anzahlung), anzahlung_betrag=coalesce(%s::numeric, anzahlung_betrag)
                           WHERE id=%s AND kunden_id=%s AND status='reserviert' AND start > now()
                           RETURNING id, termin_id, start, ende, name, email, telefon, nachricht, token, ref, anzahlung""", (betrag or None, betrag or None, bid, kid))
        except Exception:
            raise HTTPException(400, "Ungültige Buchung.")
        g = cur.fetchone()
        if not g:
            raise HTTPException(409, "Diese Buchung ist nicht mehr offen (bestätigt, abgelehnt oder abgelaufen).")
        if g["termin_id"]:
            cur.execute("UPDATE termine SET status='bestaetigt' WHERE id=%s", (g["termin_id"],))
        cur.execute("SELECT firma FROM kunden WHERE id=%s", (kid,))
        firma = cur.fetchone()["firma"]
        audit(cur, user, "website_buchung_bestaetigt", kid, {"ref": g["ref"]})
    bg.add_task(_bestaetigung_mail, kid, g, z, firma, g["anzahlung"] or "", hinweis)
    return {"ok": True}


def _bestaetigung_mail(kid, g: dict, z: dict, firma: str, betrag: str, hinweis: str):
    lokal = g["start"].astimezone(TZ).strftime("%d.%m.%Y um %H:%M Uhr")
    absagen = f"{config.BASE_URL}/site/{z['site_id']}/termin/absagen?t={g['token']}"
    zahlung = sd.zahlungsabsatz(z, betrag or None)
    text = (f"Guten Tag {g['name']},\n\nIhr Termin ist bestätigt: {lokal}.{(chr(10) + z['kal_ort']) if z['kal_ort'] else ''}\n{zahlung}\n"
            f"{(chr(10) + hinweis + chr(10)) if hinweis else ''}\nDie Kalendereinladung hängt an. Falls Sie nicht können, sagen Sie bitte hier ab:\n{absagen}\n\nBis bald!\n\nViele Grüße\n{firma}")
    sd.kundenmail(kid, g["email"], f"Ihr Termin bei {firma} am {lokal} ist bestätigt", text,
                  anhang=("termin.ics", sd.ics(str(g["id"]), g["start"], g["ende"], "Termin bei " + firma, z["kal_ort"] or ""), "text/calendar"), art="termin")


@kunde.post("/buchungen/{bid}/anzahlung-erhalten")
def k_anzahlung_erhalten(bid: str, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("buchanz:" + str(kid), 120, 600)
    with tx("system") as cur:
        try:
            cur.execute("""UPDATE website_termine SET anzahlung_erhalten_am=coalesce(anzahlung_erhalten_am, now()) WHERE id=%s AND kunden_id=%s AND status IN ('gebucht','reserviert') RETURNING ref""", (bid, kid))
        except Exception:
            raise HTTPException(400, "Ungültige Buchung.")
        r = cur.fetchone()
        if not r:
            raise HTTPException(409, "Diese Buchung gibt es nicht mehr oder sie ist abgesagt.")
        audit(cur, user, "website_anzahlung_erhalten", kid, {"ref": r["ref"]})
    return {"ok": True}


@kunde.post("/buchungen/{bid}/bezahlt")
def k_bezahlt(bid: str, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("buchbez:" + str(kid), 120, 600)
    with tx("system") as cur:
        try:
            cur.execute("""UPDATE website_termine SET bezahlt_am=coalesce(bezahlt_am, now()), anzahlung_erhalten_am=coalesce(anzahlung_erhalten_am, now())
                           WHERE id=%s AND kunden_id=%s AND status='gebucht' RETURNING ref""", (bid, kid))
        except Exception:
            raise HTTPException(400, "Ungültige Buchung.")
        r = cur.fetchone()
        if not r:
            raise HTTPException(409, "Diese Buchung gibt es nicht oder sie ist nicht bestätigt.")
        audit(cur, user, "website_bezahlt", kid, {"ref": r["ref"]})
    return {"ok": True}


class AblehnenIn(BaseModel):
    grund: str = Field(default="", max_length=600)


@kunde.post("/buchungen/{bid}/ablehnen")
def k_ablehnen(bid: str, data: AblehnenIn, bg: BackgroundTasks, user: dict = Depends(require_termine)):
    kid = user["kunden_id"]
    rate_limit("buchabl:" + str(kid), 60, 600)
    return _ablehnen(kid, bid, data.grund, bg, user, user.get("website_mitarbeiter_id"))


def _ablehnen(kid, bid: str, grund: str, bg: BackgroundTasks, user: dict | None, mitarbeiter_id=None):
    """Lehnt eine Buchung ab bzw. sagt sie ab und informiert den Gast. Gemeinsam für das Portal und den Knopf in der Mail (user=None)."""
    with tx("system") as cur:
        try:
            cur.execute("""UPDATE website_termine SET status='abgesagt', abgesagt_am=now(), abgesagt_von='betrieb', storno_kostenlos=true WHERE id=%s AND kunden_id=%s AND status IN ('reserviert','gebucht') AND start > now()
                           AND (%s::uuid IS NULL OR mitarbeiter_id = %s::uuid)
                           RETURNING termin_id, start, name, email, ref, anzahlung_erhalten_am, anzahlung_betrag""", (bid, kid, mitarbeiter_id, mitarbeiter_id))
        except Exception:
            raise HTTPException(400, "Ungültige Buchung.")
        g = cur.fetchone()
        if not g:
            raise HTTPException(409, "Diese Buchung ist nicht mehr offen.")
        if g["termin_id"]:
            cur.execute("UPDATE termine SET status='abgesagt' WHERE id=%s", (g["termin_id"],))
        cur.execute("SELECT firma FROM kunden WHERE id=%s", (kid,))
        firma = cur.fetchone()["firma"]
        audit(cur, user, "website_buchung_abgelehnt", kid, {"ref": g["ref"]})
    bg.add_task(sd.warteliste_benachrichtigen, kid, g["start"].astimezone(TZ).date())
    lokal = g["start"].astimezone(TZ).strftime("%d.%m.%Y um %H:%M Uhr")
    if not g["email"]:
        return {"ok": True, "erstatten": bool(g["anzahlung_erhalten_am"] and g["anzahlung_betrag"])}
    bg.add_task(sd.kundenmail, kid, g["email"], f"Zu Ihrem Termin am {lokal}",
                f"Guten Tag {g['name']},\n\nleider können wir den gewünschten Termin am {lokal} nicht anbieten.\n{(grund.strip() + chr(10)) if grund.strip() else ''}"
                f"{('Ihre Anzahlung von ' + str(g['anzahlung_betrag']).replace('.', ',') + ' € erstatten wir Ihnen.' + chr(10)) if g['anzahlung_erhalten_am'] and g['anzahlung_betrag'] else ''}"
                f"Auf unserer Website können Sie jederzeit einen anderen Termin wählen.\n\nViele Grüße\n{firma}", None, None, "termin")
    return {"ok": True, "erstatten": bool(g["anzahlung_erhalten_am"] and g["anzahlung_betrag"])}


@kunde.get("/erstattungen")
def k_erstattungen(user: dict = Depends(require_inhaber)):
    """Rechtzeitig oder vom Betrieb abgesagte Termine, deren Anzahlung noch zurückgezahlt werden muss."""
    with tx("system") as cur:
        cur.execute("""SELECT id, ref, name, email, start, anzahlung_betrag FROM website_termine WHERE kunden_id=%s AND status='abgesagt' AND storno_kostenlos
                       AND anzahlung_erhalten_am IS NOT NULL AND erstattet_am IS NULL AND coalesce(anzahlung_betrag,0) > 0 ORDER BY abgesagt_am DESC LIMIT 100""", (user["kunden_id"],))
        return {"erstattungen": [{"id": str(r["id"]), "ref": r["ref"] or "", "name": r["name"], "email": r["email"], "start": r["start"], "betrag": f"{r['anzahlung_betrag']:.2f}"} for r in cur.fetchall()]}


@kunde.post("/buchungen/{bid}/erstattet")
def k_erstattet(bid: str, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("buchers:" + str(kid), 120, 600)
    with tx("system") as cur:
        try:
            cur.execute("""UPDATE website_termine SET erstattet_am=coalesce(erstattet_am, now()) WHERE id=%s AND kunden_id=%s AND status='abgesagt' AND anzahlung_erhalten_am IS NOT NULL RETURNING ref""", (bid, kid))
        except Exception:
            raise HTTPException(400, "Ungültige Buchung.")
        r = cur.fetchone()
        if not r:
            raise HTTPException(409, "Zu dieser Buchung gibt es keine Erstattung.")
        audit(cur, user, "website_anzahlung_erstattet", kid, {"ref": r["ref"]})
    return {"ok": True}


def _quittung_daten(kid, bid: str):
    with tx("system") as cur:
        try:
            cur.execute("""SELECT w.id, w.ref, w.name, w.email, w.start, w.anzahlung_betrag, w.anzahlung_erhalten_am, w.status, l.name AS leistung
                           FROM website_termine w LEFT JOIN website_leistungen l ON l.id=w.leistung_id WHERE w.id=%s AND w.kunden_id=%s""", (bid, kid))
        except Exception:
            raise HTTPException(400, "Ungültige Buchung.")
        g = cur.fetchone()
        if not g or not g["anzahlung_erhalten_am"] or not g["anzahlung_betrag"]:
            raise HTTPException(409, "Für diesen Termin ist keine erhaltene Anzahlung eingetragen.")
        cur.execute("SELECT firma, strasse, plz, ort FROM kunden WHERE id=%s", (kid,))
        return g, cur.fetchone()


@kunde.get("/buchungen/{bid}/quittung.pdf")
def k_quittung(bid: str, user: dict = Depends(require_inhaber)):
    g, b = _quittung_daten(user["kunden_id"], bid)
    pdfb = quittung.quittung_pdf(b, g["name"], g["start"].astimezone(TZ).strftime("%d.%m.%Y um %H:%M Uhr"), g["leistung"] or "", float(g["anzahlung_betrag"]), g["anzahlung_erhalten_am"].astimezone(TZ), g["ref"] or "")
    return Response(pdfb, media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="quittung-{g["ref"] or "anzahlung"}.pdf"'})


@kunde.post("/buchungen/{bid}/quittung-senden")
def k_quittung_senden(bid: str, bg: BackgroundTasks, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("buchquitt:" + str(kid), 60, 600)
    g, b = _quittung_daten(kid, bid)
    if not EMAIL_RE.match(g["email"] or ""):
        raise HTTPException(409, "Zu diesem Termin gibt es keine gültige E-Mail-Adresse.")
    pdfb = quittung.quittung_pdf(b, g["name"], g["start"].astimezone(TZ).strftime("%d.%m.%Y um %H:%M Uhr"), g["leistung"] or "", float(g["anzahlung_betrag"]), g["anzahlung_erhalten_am"].astimezone(TZ), g["ref"] or "")
    with tx("system") as cur:
        cur.execute("UPDATE website_termine SET quittung_am=now() WHERE id=%s", (bid,))
        audit(cur, user, "website_quittung_gesendet", kid, {"ref": g["ref"]})
    bg.add_task(sd.kundenmail, kid, g["email"], f"Quittung über Ihre Anzahlung bei {b['firma']}",
                f"Guten Tag {g['name']},\n\nanbei die Quittung über Ihre Anzahlung von {str(g['anzahlung_betrag']).replace('.', ',')} €.\n\nViele Grüße\n{b['firma']}",
                None, (f"quittung-{g['ref'] or 'anzahlung'}.pdf", pdfb, "application/pdf"), "termin")
    return {"ok": True}


# ---------------------------------------------------------------- Kalenderansicht und Termine von Hand
@kunde.get("/kalender-ansicht")
def k_kalender_ansicht(von: str = "", tage: int = 7, user: dict = Depends(require_termine)):
    """Termine eines Zeitraums (1 bis 14 Tage) für die Tages- und Wochenansicht."""
    try:
        a = date.fromisoformat(von) if von else datetime.now(TZ).date()
    except ValueError:
        raise HTTPException(400, "Das Datum ist ungültig.")
    tage = max(1, min(14, tage))
    ab, bis_ = datetime.combine(a, datetime.min.time(), TZ), datetime.combine(a + timedelta(days=tage), datetime.min.time(), TZ)
    kid = user["kunden_id"]
    with tx("system") as cur:
        z = sd.zeile(cur, kid)
        sd.reservierungen_verfallen(cur, kid, z["kal_halte_std"])
        cur.execute("""SELECT w.id, w.ref, w.status, w.start, w.ende, w.name, w.telefon, w.email, w.mitarbeiter_id, w.anzahlung_betrag, w.anzahlung_erhalten_am, w.bezahlt_am, w.nachricht, w.kommt_am, w.ergebnis,
                              l.name AS leistung, m.name AS mitarbeiter, (w.daten->>'quelle') = 'manuell' AS manuell
                       FROM website_termine w LEFT JOIN website_leistungen l ON l.id=w.leistung_id LEFT JOIN website_mitarbeiter m ON m.id=w.mitarbeiter_id
                       WHERE w.kunden_id=%s AND w.status IN ('gebucht','reserviert') AND w.start >= %s AND w.start < %s
                             AND (%s::uuid IS NULL OR w.mitarbeiter_id = %s::uuid) ORDER BY w.start LIMIT 800""", (kid, ab, bis_, user.get("website_mitarbeiter_id"), user.get("website_mitarbeiter_id")))
        termine = cur.fetchall()
        cur.execute("SELECT id, name FROM website_mitarbeiter WHERE kunden_id=%s AND aktiv AND (%s::uuid IS NULL OR id = %s::uuid) ORDER BY sort, name", (kid, user.get("website_mitarbeiter_id"), user.get("website_mitarbeiter_id")))
        team = cur.fetchall()
        cur.execute("""SELECT s.tag, s.grund, s.mitarbeiter_id FROM website_sperrtage s WHERE s.kunden_id=%s AND s.tag >= %s AND s.tag < %s""", (kid, a, a + timedelta(days=tage)))
        sperr = cur.fetchall()
        plan = {}
        for r in sd.ressourcen(cur, z):
            plan[str(r["id"] or "")] = {str(t): f for t, f in r["plan"].items()}
    return {"von": a.isoformat(), "tage": tage, "mitarbeiter": [{"id": str(m["id"]), "name": m["name"]} for m in team],
            "termine": [{"id": str(t["id"]), "ref": t["ref"] or "", "status": t["status"], "start": t["start"], "ende": t["ende"], "name": t["name"], "telefon": t["telefon"] or "",
                         "email": t["email"] or "", "mitarbeiter_id": str(t["mitarbeiter_id"]) if t["mitarbeiter_id"] else "", "mitarbeiter": t["mitarbeiter"] or "", "leistung": t["leistung"] or "",
                         "anzahlung": f"{t['anzahlung_betrag']:.2f}" if t["anzahlung_betrag"] else "", "anzahlung_erhalten": bool(t["anzahlung_erhalten_am"]), "bezahlt": bool(t["bezahlt_am"]),
                         "notiz": t["nachricht"] or "", "manuell": bool(t["manuell"]), "kommt": bool(t["kommt_am"]), "ergebnis": t["ergebnis"] or ""} for t in termine],
            "sperr": [{"tag": s["tag"].isoformat(), "grund": s["grund"] or "", "mitarbeiter_id": str(s["mitarbeiter_id"]) if s["mitarbeiter_id"] else ""} for s in sperr],
            "arbeitszeit": plan}


class ManuellIn(BaseModel):
    name: str = Field(min_length=2, max_length=120)
    telefon: str = Field(default="", max_length=40)
    email: str = Field(default="", max_length=200)
    start: str = Field(pattern=r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}$")      # Ortszeit, z. B. 2026-11-03T14:30
    leistung_id: str = Field(default="", max_length=40)
    dauer_min: int = Field(default=0, ge=0, le=720)                     # 0 = Dauer der Leistung
    personen: int = Field(default=1, ge=1, le=5)                        # mehrere Personen nacheinander: Dauer mal Personenzahl
    mitarbeiter_id: str = Field(default="", max_length=40)
    notiz: str = Field(default="", max_length=600)
    wiederholen: int = Field(default=0, ge=0, le=51)                    # so viele weitere Termine
    abstand_tage: int = Field(default=7, ge=1, le=90)


@kunde.get("/kalender-optionen")
def k_kalender_optionen(user: dict = Depends(require_termine)):
    """Auswahllisten für „Termin eintragen“ (Leistungen und Personen)."""
    kid, eigene = user["kunden_id"], user.get("website_mitarbeiter_id")
    with tx("system") as cur:
        cur.execute("SELECT id, name, dauer_min FROM website_leistungen WHERE kunden_id=%s AND aktiv ORDER BY sort, name", (kid,))
        ls = cur.fetchall()
        cur.execute("SELECT id, name FROM website_mitarbeiter WHERE kunden_id=%s AND aktiv AND (%s::uuid IS NULL OR id = %s::uuid) ORDER BY sort, name", (kid, eigene, eigene))
        ms = cur.fetchall()
    return {"leistungen": [{"id": str(x["id"]), "name": x["name"], "dauer_min": x["dauer_min"]} for x in ls], "mitarbeiter": [{"id": str(x["id"]), "name": x["name"]} for x in ms], "eigene": str(eigene) if eigene else ""}


@kunde.post("/buchungen/manuell")
def k_manuell(data: ManuellIn, user: dict = Depends(require_termine)):
    """Trägt einen Termin von Hand ein (Telefon, Laufkundschaft, Stammkunde). Er sperrt die Zeit wie eine Buchung über die Website; auch als Serie."""
    kid = user["kunden_id"]
    rate_limit("buchmanuell:" + str(kid), 60, 600)
    if data.email.strip() and not EMAIL_RE.match(data.email.strip()):
        raise HTTPException(400, "Die E-Mail-Adresse ist ungültig.")
    try:
        start0 = datetime.fromisoformat(data.start).replace(tzinfo=TZ)
    except ValueError:
        raise HTTPException(400, "Die Zeit ist ungültig.")
    if start0 < datetime.now(TZ) - timedelta(days=1):
        raise HTTPException(400, "Der Termin liegt zu weit in der Vergangenheit.")
    mid = _uuid(data.mitarbeiter_id) if data.mitarbeiter_id else None
    if user.get("website_mitarbeiter_id"):
        mid = str(user["website_mitarbeiter_id"])                     # verknüpfter Zugang trägt nur für sich selbst ein
    angelegt, uebersprungen = [], []
    with tx("system") as cur:
        cur.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s, 0))", (str(kid),))
        lz = None
        if data.leistung_id:
            cur.execute("SELECT * FROM website_leistungen WHERE id=%s AND kunden_id=%s", (_uuid(data.leistung_id), kid))
            lz = cur.fetchone()
            if not lz:
                raise HTTPException(400, "Diese Leistung gibt es nicht.")
        if mid:
            cur.execute("SELECT 1 FROM website_mitarbeiter WHERE id=%s AND kunden_id=%s", (mid, kid))
            if not cur.fetchone():
                raise HTTPException(400, "Diesen Mitarbeiter gibt es nicht.")
        z = sd.zeile(cur, kid)
        minuten = (data.dauer_min or (lz or {}).get("dauer_min") or z["kal_dauer"]) * data.personen
        puffer = (lz or {}).get("puffer_min") or 0
        preis = sd.preis_zahl((lz or {}).get("preis"))
        for i in range(data.wiederholen + 1):
            start = (start0.replace(tzinfo=None) + timedelta(days=i * data.abstand_tage)).replace(tzinfo=TZ).astimezone(timezone.utc)
            ende = start + timedelta(minutes=minuten)
            cur.execute("""SELECT count(*) AS n FROM website_termine WHERE kunden_id=%s AND status IN ('gebucht','reserviert') AND start < %s AND ende + make_interval(mins => coalesce(puffer_min,0)) > %s
                           AND (%s::uuid IS NULL OR mitarbeiter_id = %s::uuid OR mitarbeiter_id IS NULL)""", (kid, ende + timedelta(minutes=puffer), start, mid, mid))
            lokal = start.astimezone(TZ).strftime("%d.%m.%Y %H:%M")
            if cur.fetchone()["n"]:
                uebersprungen.append({"wann": lokal, "grund": "Zu dieser Zeit ist schon ein Termin eingetragen."})
                continue
            titel = " – ".join(x for x in ((lz or {}).get("name"), data.name.strip()) if x)[:120]
            kontakt = ", ".join(x for x in (data.email.strip(), data.telefon.strip()) if x)
            cur.execute("""INSERT INTO termine (kunden_id, titel, beginn, ende, kontakt_name, kontakt_info, status, notiz) VALUES (%s,%s,%s,%s,%s,%s,'bestaetigt',%s) RETURNING id""",
                        (kid, titel, start, ende, data.name.strip(), kontakt, data.notiz.strip()[:1500] or None))
            tid = cur.fetchone()["id"]
            ref = "M" + datetime.now(TZ).strftime("%y%m%d") + "-" + secrets.token_hex(2).upper()
            try:
                cur.execute("SAVEPOINT s1")
                cur.execute("""INSERT INTO website_termine (kunden_id, termin_id, start, ende, name, email, telefon, nachricht, token, status, ref, daten, mitarbeiter_id, leistung_id, puffer_min, preis_betrag)
                               VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,'gebucht',%s,%s,%s,%s,%s,%s) RETURNING id""",
                            (kid, tid, start, ende, data.name.strip(), data.email.strip(), data.telefon.strip() or None, data.notiz.strip() or None, secrets.token_urlsafe(24), ref,
                             json.dumps({"quelle": "manuell", "personen": data.personen}), mid, (lz or {}).get("id"), puffer, preis * data.personen if preis else None))
                angelegt.append({"id": str(cur.fetchone()["id"]), "wann": lokal})
            except Exception:
                cur.execute("ROLLBACK TO SAVEPOINT s1")
                cur.execute("DELETE FROM termine WHERE id=%s", (tid,))
                uebersprungen.append({"wann": lokal, "grund": "Diese Zeit ist schon vergeben."})
        audit(cur, user, "website_termin_manuell", kid, {"anzahl": len(angelegt)})
    if not angelegt:
        raise HTTPException(409, "Kein Termin angelegt: " + (uebersprungen[0]["grund"] if uebersprungen else "unbekannter Grund"))
    return {"ok": True, "angelegt": angelegt, "uebersprungen": uebersprungen}


# ---------------------------------------------------------------- Stammkunden
def _stamm_schluessel(email, tel) -> str:
    e = (email or "").strip().lower()
    if e:
        return e
    z = re.sub(r"\D", "", tel or "")
    return "tel:" + z if z else ""


@kunde.get("/stamm")
def k_stamm(q: str = "", user: dict = Depends(require_termine)):
    """Wer kommt wie oft? Aus den Terminen zusammengestellt (E-Mail oder Telefon gleicht Personen ab), mit Notiz."""
    if user.get("website_mitarbeiter_id"):
        raise HTTPException(403, "Die Stammkundenliste sehen nur Inhaber, Empfang und Mitarbeiter ohne Personenbindung.")
    kid = user["kunden_id"]
    with tx("system") as cur:
        cur.execute("SELECT name, email, telefon, start, status, ergebnis FROM website_termine WHERE kunden_id=%s ORDER BY start DESC LIMIT 6000", (kid,))
        zeilen = cur.fetchall()
        cur.execute("SELECT schluessel, notiz FROM website_stamm_notizen WHERE kunden_id=%s", (kid,))
        notizen = {r["schluessel"]: r["notiz"] for r in cur.fetchall()}
    jetzt = datetime.now(timezone.utc)
    gruppen: dict = {}
    for r in zeilen:                                                        # neueste zuerst: der Name der letzten Buchung gilt
        k = _stamm_schluessel(r["email"], r["telefon"])
        if not k:
            continue
        g = gruppen.setdefault(k, {"schluessel": k, "name": r["name"], "email": r["email"] or "", "telefon": r["telefon"] or "", "besuche": 0, "kommend": 0,
                                   "nicht_erschienen": 0, "abgesagt": 0, "letzter": None, "naechster": None})
        g["email"] = g["email"] or (r["email"] or "")
        g["telefon"] = g["telefon"] or (r["telefon"] or "")
        if r["status"] == "abgesagt":
            g["abgesagt"] += 1
        elif r["start"] > jetzt:
            g["kommend"] += 1
            g["naechster"] = r["start"] if g["naechster"] is None or r["start"] < g["naechster"] else g["naechster"]
        else:
            if r["ergebnis"] == "nicht_erschienen":
                g["nicht_erschienen"] += 1
            else:
                g["besuche"] += 1
                g["letzter"] = g["letzter"] or r["start"]
    ql = q.strip().lower()
    liste = [dict(g, notiz=notizen.get(g["schluessel"], "")) for g in gruppen.values()
             if not ql or ql in (g["name"] or "").lower() or ql in g["email"].lower() or re.sub(r"\D", "", ql) and re.sub(r"\D", "", ql) in re.sub(r"\D", "", g["telefon"])]
    liste.sort(key=lambda g: (g["naechster"] is None, g["naechster"] or datetime.max.replace(tzinfo=timezone.utc), -(g["letzter"] or datetime.min.replace(tzinfo=timezone.utc)).timestamp()))
    return {"personen": liste[:200], "gesamt": len(liste)}


class StammNotizIn(BaseModel):
    schluessel: str = Field(min_length=3, max_length=200)
    notiz: str = Field(default="", max_length=1000)


@kunde.put("/stamm/notiz")
def k_stamm_notiz(data: StammNotizIn, user: dict = Depends(require_termine)):
    if user.get("website_mitarbeiter_id"):
        raise HTTPException(403, "Nicht für diesen Zugang.")
    kid = user["kunden_id"]
    rate_limit("stammnotiz:" + str(kid), 120, 600)
    with tx("system") as cur:
        if data.notiz.strip():
            cur.execute("""INSERT INTO website_stamm_notizen (kunden_id, schluessel, notiz) VALUES (%s,%s,%s)
                           ON CONFLICT (kunden_id, schluessel) DO UPDATE SET notiz=EXCLUDED.notiz, geaendert_am=now()""", (kid, data.schluessel.strip().lower()[:200], data.notiz.strip()))
        else:
            cur.execute("DELETE FROM website_stamm_notizen WHERE kunden_id=%s AND schluessel=%s", (kid, data.schluessel.strip().lower()))
    return {"ok": True}


@kunde.get("/sperrtage")
def k_sperrtage(user: dict = Depends(require_inhaber)):
    with tx("system") as cur:
        cur.execute("""SELECT s.tag, s.grund, s.mitarbeiter_id, m.name AS mitarbeiter FROM website_sperrtage s LEFT JOIN website_mitarbeiter m ON m.id=s.mitarbeiter_id
                       WHERE s.kunden_id=%s AND s.tag >= current_date ORDER BY s.tag, m.name NULLS FIRST LIMIT 600""", (user["kunden_id"],))
        return {"tage": [{"tag": r["tag"].isoformat(), "grund": r["grund"] or "", "mitarbeiter_id": str(r["mitarbeiter_id"]) if r["mitarbeiter_id"] else "",
                          "mitarbeiter": r["mitarbeiter"] or ""} for r in cur.fetchall()]}


class SperrIn(BaseModel):
    von: str = Field(pattern=r"^\d{4}-\d{2}-\d{2}$")
    bis: str = Field(default="", pattern=r"^(\d{4}-\d{2}-\d{2})?$")
    grund: str = Field(default="", max_length=80)
    mitarbeiter_id: str = Field(default="", max_length=40)      # leer = ganzer Betrieb


def _uuid(w: str, text: str = "Ungültige Angabe."):
    try:
        return str(uuid.UUID(str(w)))
    except ValueError:
        raise HTTPException(400, text)


@kunde.post("/sperrtage")
def k_sperre(data: SperrIn, user: dict = Depends(require_inhaber)):
    try:
        a = date.fromisoformat(data.von)
        b = date.fromisoformat(data.bis) if data.bis else a
    except ValueError:
        raise HTTPException(400, "Das Datum ist ungültig.")
    if b < a or (b - a).days > 120:
        raise HTTPException(400, "Der Zeitraum ist ungültig (höchstens 120 Tage am Stück).")
    mid = _uuid(data.mitarbeiter_id) if data.mitarbeiter_id else None
    with tx("system") as cur:
        if mid:
            cur.execute("SELECT 1 FROM website_mitarbeiter WHERE id=%s AND kunden_id=%s", (mid, user["kunden_id"]))
            if not cur.fetchone():
                raise HTTPException(400, "Diesen Mitarbeiter gibt es nicht.")
        for i in range((b - a).days + 1):
            cur.execute("""INSERT INTO website_sperrtage (kunden_id, tag, grund, mitarbeiter_id) VALUES (%s,%s,%s,%s)
                           ON CONFLICT (kunden_id, tag, coalesce(mitarbeiter_id, '00000000-0000-0000-0000-000000000000'::uuid)) DO UPDATE SET grund=excluded.grund""",
                        (user["kunden_id"], a + timedelta(days=i), data.grund.strip() or None, mid))
    return {"ok": True}


@kunde.delete("/sperrtage/{tag}")
def k_sperre_weg(tag: str, mitarbeiter: str = "", user: dict = Depends(require_inhaber)):
    try:
        t = date.fromisoformat(tag)
    except ValueError:
        raise HTTPException(400, "Das Datum ist ungültig.")
    mid = _uuid(mitarbeiter) if mitarbeiter else None
    with tx("system") as cur:
        cur.execute("DELETE FROM website_sperrtage WHERE kunden_id=%s AND tag=%s AND mitarbeiter_id IS NOT DISTINCT FROM %s", (user["kunden_id"], t, mid))
    return {"ok": True}


# =============================================================== Leistungen und Mitarbeiter (Inhaber)
def _slug(text: str) -> str:
    t = text.lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    return re.sub(r"[^a-z0-9]+", "-", t).strip("-")[:40]


def _team(cur, kid) -> dict:
    cur.execute("SELECT id, schluessel, name, beschreibung, preis, dauer_min, puffer_min, aktiv, sort, freigabe, anzahlung_art, anzahlung_wert, felder FROM website_leistungen WHERE kunden_id=%s ORDER BY sort, name", (kid,))
    leist = [{**r, "id": str(r["id"]), "beschreibung": r["beschreibung"] or "", "preis": r["preis"] or "", "anzahlung_wert": float(r["anzahlung_wert"]) if r["anzahlung_wert"] is not None else None,
              "anzahlung": f"{x:.2f}" if (x := sd.anzahlung_berechnen(r)) else ""} for r in cur.fetchall()]
    cur.execute("SELECT * FROM website_mitarbeiter WHERE kunden_id=%s ORDER BY sort, name", (kid,))
    ms = cur.fetchall()
    cur.execute("SELECT mitarbeiter_id, leistung_id FROM website_mitarbeiter_leistungen WHERE mitarbeiter_id IN (SELECT id FROM website_mitarbeiter WHERE kunden_id=%s)", (kid,))
    bietet: dict = {}
    for r in cur.fetchall():
        bietet.setdefault(r["mitarbeiter_id"], []).append(str(r["leistung_id"]))
    z = sd.zeile(cur, kid)
    return {"leistungen": leist, "einstellung": {"frist_std": z["anz_frist_std"], "storno_std": z["storno_std"], "leistungen_freigabe": z["leistungen_freigabe"], "paypal_bereit": bool(z["pp_aktiv"] and (z["pp_me_link"] or "").startswith("https://"))},
            "mitarbeiter": [{"id": str(m["id"]), "schluessel": m["schluessel"], "name": m["name"], "aktiv": m["aktiv"], "wochentage": m["wochentage"] or [], "plan": sd.plan_text(m.get("wochenplan")), "von": m["von"] or "",
                             "bis": m["bis"] or "", "zeiten": m["zeiten"] or "", "leistungen": bietet.get(m["id"], []), "ical_gesetzt": bool(m["ical_enc"]),
                             "ical_geprueft_am": m["ical_geprueft_am"]} for m in ms]}


@kunde.get("/team")
def k_team(user: dict = Depends(require_inhaber)):
    with tx("system") as cur:
        return _team(cur, user["kunden_id"])


class FeldIn(BaseModel):
    label: str = Field(min_length=2, max_length=60)
    key: str = Field(default="", max_length=30)
    pflicht: bool = False


class LeistungIn(BaseModel):
    name: str = Field(min_length=2, max_length=80)
    schluessel: str = Field(default="", max_length=40)
    beschreibung: str = Field(default="", max_length=300)
    preis: str = Field(default="", max_length=40)
    dauer_min: int = Field(default=30, ge=5, le=720)
    puffer_min: int = Field(default=0, ge=0, le=240)
    aktiv: bool = True
    freigabe: bool = False                                         # Termine dieser Leistung bestätigt der Betrieb selbst
    anzahlung_art: str = Field(default="keine", pattern=r"^(keine|betrag|prozent)$")
    anzahlung_wert: float | None = Field(default=None, ge=0, le=100000)
    felder: list[FeldIn] = Field(default_factory=list, max_length=12)      # Zusatzfragen im Buchungsformular, z. B. „Haarlänge“

    def pruefen(self):
        belegt = {"name", "telefon", "email", "nachricht", "datum", "datum2", "uhrzeit", "leistung", "mitarbeiter", "typ", "website", "hp", "einwilligung"}
        for f in self.felder:
            f.label = f.label.strip()
            f.key = (re.sub(r"[^a-z0-9]+", "_", (f.key or f.label).lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")).strip("_")[:30]) or ""
            if not re.match(r"^[a-z][a-z0-9_]*$", f.key) or f.key in belegt:
                raise HTTPException(400, f"Die Zusatzfrage „{f.label}“ hat einen ungültigen oder reservierten Namen. Bitte anders benennen.")
        if len({f.key for f in self.felder}) != len(self.felder):
            raise HTTPException(400, "Zwei Zusatzfragen haben denselben Namen.")
        if self.anzahlung_art != "keine":
            if not self.anzahlung_wert or self.anzahlung_wert <= 0:
                raise HTTPException(400, "Bitte die Anzahlung als Zahl eintragen (Betrag in € oder Prozent).")
            if self.anzahlung_art == "prozent" and self.anzahlung_wert > 100:
                raise HTTPException(400, "Mehr als 100 Prozent sind nicht möglich.")
            if self.anzahlung_art == "prozent" and sd.preis_zahl(self.preis) is None:
                raise HTTPException(400, "Für eine Anzahlung in Prozent braucht die Leistung einen Preis mit Zahl, z. B. 45 €.")
        return self


def _schluessel(roh: str, name: str) -> str:
    k = _slug(roh or name)
    if not k:
        raise HTTPException(400, "Bitte einen Namen mit Buchstaben oder Zahlen angeben.")
    return k


def _leistung_neu(cur, kid, data: LeistungIn) -> dict:
    k = _schluessel(data.schluessel, data.name)
    cur.execute("SELECT count(*) AS n FROM website_leistungen WHERE kunden_id=%s", (kid,))
    if cur.fetchone()["n"] >= 60:
        raise HTTPException(400, "Mehr als 60 Leistungen sind nicht möglich.")
    cur.execute("SELECT 1 FROM website_leistungen WHERE kunden_id=%s AND schluessel=%s", (kid, k))
    if cur.fetchone():
        raise HTTPException(400, "Diesen Kurznamen gibt es schon. Bitte einen anderen wählen.")
    art = data.anzahlung_art
    cur.execute("""INSERT INTO website_leistungen (kunden_id, schluessel, name, beschreibung, preis, dauer_min, puffer_min, aktiv, freigabe, anzahlung_art, anzahlung_wert, felder, sort)
                   VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s::jsonb,(SELECT coalesce(max(sort),0)+1 FROM website_leistungen WHERE kunden_id=%s)) RETURNING id""",
                (kid, k, data.name.strip(), data.beschreibung.strip() or None, data.preis.strip() or None, data.dauer_min, data.puffer_min, data.aktiv, data.freigabe,
                 art, data.anzahlung_wert if art != "keine" else None, json.dumps([f.model_dump() for f in data.felder]), kid))
    neue_id = str(cur.fetchone()["id"])
    sd.leistungen_wissen_sync(cur, kid)
    return {"ok": True, "id": neue_id, "schluessel": k}


def _leistung_aendern(cur, kid, lid, data: LeistungIn) -> dict:
    cur.execute("SELECT schluessel FROM website_leistungen WHERE id=%s AND kunden_id=%s", (lid, kid))
    alt = cur.fetchone()
    if not alt:
        raise HTTPException(404, "Diese Leistung gibt es nicht.")
    k = _slug(data.schluessel) if data.schluessel.strip() else alt["schluessel"]
    if not k:
        raise HTTPException(400, "Der Kurzname ist ungültig.")
    cur.execute("SELECT 1 FROM website_leistungen WHERE kunden_id=%s AND schluessel=%s AND id<>%s", (kid, k, lid))
    if cur.fetchone():
        raise HTTPException(400, "Diesen Kurznamen gibt es schon. Bitte einen anderen wählen.")
    art = data.anzahlung_art
    cur.execute("""UPDATE website_leistungen SET schluessel=%s, name=%s, beschreibung=%s, preis=%s, dauer_min=%s, puffer_min=%s, aktiv=%s, freigabe=%s, anzahlung_art=%s, anzahlung_wert=%s, felder=%s::jsonb
                   WHERE id=%s AND kunden_id=%s""",
                (k, data.name.strip(), data.beschreibung.strip() or None, data.preis.strip() or None, data.dauer_min, data.puffer_min, data.aktiv, data.freigabe,
                 art, data.anzahlung_wert if art != "keine" else None, json.dumps([f.model_dump() for f in data.felder]), lid, kid))
    sd.leistungen_wissen_sync(cur, kid)
    return {"ok": True}


def _leistung_weg(cur, kid, lid) -> dict:
    cur.execute("DELETE FROM website_leistungen WHERE id=%s AND kunden_id=%s", (lid, kid))
    sd.leistungen_wissen_sync(cur, kid)
    return {"ok": True}


def _vormerken(cur, user, aktion: str, lid, name: str, daten: dict) -> dict:
    """Muss das Team Änderungen an Leistungen freigeben, wird die Änderung nur vorgemerkt."""
    kid = user["kunden_id"]
    cur.execute("SELECT count(*) AS n FROM website_leistung_aenderungen WHERE kunden_id=%s AND status='offen'", (kid,))
    if cur.fetchone()["n"] >= 40:
        raise HTTPException(409, "Es warten schon viele Änderungen auf Freigabe. Bitte warte, bis das Team sie bearbeitet hat.")
    cur.execute("INSERT INTO website_leistung_aenderungen (kunden_id, aktion, leistung_id, name, daten, von) VALUES (%s,%s,%s,%s,%s,%s) RETURNING id",
                (kid, aktion, lid, name[:80], json.dumps(daten, ensure_ascii=False), user["id"]))
    aid = cur.fetchone()["id"]
    cur.execute("SELECT firma FROM kunden WHERE id=%s", (kid,))
    firma = cur.fetchone()["firma"]
    cur.execute("SELECT email FROM benutzer WHERE rolle='admin' AND aktiv ORDER BY erstellt_am LIMIT 1")
    a = cur.fetchone()
    cur.execute("""INSERT INTO aufgaben (kunden_id, titel, zustaendig, faellig_am, quelle, auto_schluessel) VALUES (%s,%s,%s,current_date + 2,'auto',%s) ON CONFLICT DO NOTHING""",
                (kid, f"{firma}: Leistungs-Änderung prüfen ({name})"[:200], a and a["email"], f"leistung-{aid}"))
    ids = push.admins(cur)
    push.senden(ids, "Leistungen: Änderung wartet", f"{firma}: {name}", "/admin#kunden", "leistung")
    return {"ok": True, "wartet": True, "id": str(aid)}


@kunde.post("/leistungen")
def k_leistung_neu(data: LeistungIn, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    data.pruefen()
    with tx("system") as cur:
        if sd.zeile(cur, kid)["leistungen_freigabe"]:
            _schluessel(data.schluessel, data.name)
            return _vormerken(cur, user, "neu", None, data.name, data.model_dump())
        return _leistung_neu(cur, kid, data)


@kunde.put("/leistungen/{lid}")
def k_leistung_aendern(lid: str, data: LeistungIn, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    lid = _uuid(lid)
    data.pruefen()
    with tx("system") as cur:
        if sd.zeile(cur, kid)["leistungen_freigabe"]:
            cur.execute("SELECT 1 FROM website_leistungen WHERE id=%s AND kunden_id=%s", (lid, kid))
            if not cur.fetchone():
                raise HTTPException(404, "Diese Leistung gibt es nicht.")
            return _vormerken(cur, user, "aendern", lid, data.name, data.model_dump())
        return _leistung_aendern(cur, kid, lid, data)


@kunde.delete("/leistungen/{lid}")
def k_leistung_weg(lid: str, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    lid = _uuid(lid)
    with tx("system") as cur:
        if sd.zeile(cur, kid)["leistungen_freigabe"]:
            cur.execute("SELECT name FROM website_leistungen WHERE id=%s AND kunden_id=%s", (lid, kid))
            r = cur.fetchone()
            if not r:
                raise HTTPException(404, "Diese Leistung gibt es nicht.")
            return _vormerken(cur, user, "loeschen", lid, r["name"], {})
        return _leistung_weg(cur, kid, lid)


@kunde.get("/leistungen/aenderungen")
def k_leistung_aenderungen(user: dict = Depends(require_inhaber)):
    with tx("system") as cur:
        cur.execute("""SELECT id, aktion, name, status, notiz, erstellt_am FROM website_leistung_aenderungen WHERE kunden_id=%s
                       AND (status='offen' OR bearbeitet_am > now() - interval '14 days') ORDER BY (status='offen') DESC, erstellt_am DESC LIMIT 40""", (user["kunden_id"],))
        return {"aenderungen": [{**r, "id": str(r["id"])} for r in cur.fetchall()]}


@kunde.delete("/leistungen/aenderungen/{aid}")
def k_leistung_zurueck(aid: str, user: dict = Depends(require_inhaber)):
    with tx("system") as cur:
        cur.execute("UPDATE website_leistung_aenderungen SET status='zurueckgezogen', bearbeitet_am=now() WHERE id=%s AND kunden_id=%s AND status='offen' RETURNING id", (_uuid(aid), user["kunden_id"]))
        if not cur.fetchone():
            raise HTTPException(409, "Diese Änderung ist nicht mehr offen.")
        cur.execute("UPDATE aufgaben SET status='erledigt', erledigt_am=now(), erledigt_von='Kunde' WHERE auto_schluessel=%s AND status='offen'", (f"leistung-{aid}",))
    return {"ok": True}


class AnzahlungEinstellungIn(BaseModel):
    frist_std: int = Field(default=0, ge=0, le=720)
    storno_std: int = Field(default=24, ge=0, le=168)


@kunde.put("/anzahlung")
def k_anzahlung_einstellung(data: AnzahlungEinstellungIn, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    with tx("system") as cur:
        sd.zeile(cur, kid)
        cur.execute("UPDATE website_dienste SET anz_frist_std=%s, storno_std=%s WHERE kunden_id=%s", (data.frist_std, data.storno_std, kid))
        sd.leistungen_wissen_sync(cur, kid)
        audit(cur, user, "website_anzahlung_einstellung", kid, data.model_dump())
    return {"ok": True}


class MitarbeiterIn(BaseModel):
    name: str = Field(min_length=2, max_length=80)
    schluessel: str = Field(default="", max_length=40)
    aktiv: bool = True
    wochentage: list[int] = Field(default=[], max_length=7)      # leer = wie der Betrieb
    von: str = Field(default="", pattern=r"^(\d{2}:\d{2})?$")
    bis: str = Field(default="", pattern=r"^(\d{2}:\d{2})?$")
    zeiten: str = Field(default="", max_length=300)
    leistungen: list[str] = Field(default=[], max_length=60)       # leer = alle Leistungen
    ical_url: str = Field(default="", max_length=1500)             # leer = unverändert
    plan: dict[str, str] = Field(default_factory=dict, max_length=7)   # Arbeitszeiten je Wochentag; leer = Wochentage/von/bis oder wie der Betrieb


def _ma_pruefen(cur, kid, data: MitarbeiterIn):
    try:
        plan = sd.plan_lesen(data.plan) if data.plan else {}
    except ValueError as e:
        raise HTTPException(400, str(e))
    if data.plan and not plan:
        raise HTTPException(400, "Bitte mindestens einen Tag mit Arbeitszeiten eintragen.")
    if plan:
        data.wochentage, data.von, data.bis = sd.plan_grenzen(plan)
    if any(w < 1 or w > 7 for w in data.wochentage):
        raise HTTPException(400, "Die Wochentage sind ungültig.")
    if bool(data.von) != bool(data.bis) or (data.von and data.von >= data.bis):
        raise HTTPException(400, "Bitte Beginn und Ende angeben, das Ende nach dem Beginn.")
    zeiten = [x.strip() for x in re.split(r"[,;\s]+", data.zeiten) if x.strip()]
    if any(not sd.ZEIT_RE.match(x) for x in zeiten) or len(zeiten) > 24:
        raise HTTPException(400, "Die Uhrzeiten bitte so eintragen: 08:00, 10:00, 13:30.")
    ids = [_uuid(x) for x in data.leistungen]
    if ids:
        cur.execute("SELECT count(*) AS n FROM website_leistungen WHERE kunden_id=%s AND id = ANY(%s::uuid[])", (kid, ids))
        if cur.fetchone()["n"] != len(set(ids)):
            raise HTTPException(400, "Eine der Leistungen gibt es nicht.")
    enc, geprueft = None, None
    if data.ical_url.strip():
        try:
            u = sd.ical_url_pruefen(data.ical_url)
            sd.ical_belegt(sd.ical_laden(u), datetime.now(timezone.utc), datetime.now(timezone.utc) + timedelta(days=30))
        except ValueError as e:
            raise HTTPException(400, str(e))
        enc, geprueft = einstellungen.verschluesseln(u), datetime.now(timezone.utc)
    return ", ".join(sorted(set(zeiten))) or None, sorted(set(ids)), enc, geprueft, plan


def _ma_leistungen(cur, mid, ids):
    cur.execute("DELETE FROM website_mitarbeiter_leistungen WHERE mitarbeiter_id=%s", (mid,))
    for lid in ids:
        cur.execute("INSERT INTO website_mitarbeiter_leistungen (mitarbeiter_id, leistung_id) VALUES (%s,%s)", (mid, lid))


@kunde.post("/mitarbeiter")
def k_ma_neu(data: MitarbeiterIn, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("mapruef:" + str(kid), 30, 600)
    k = _schluessel(data.schluessel, data.name)
    with tx("system") as cur:
        cur.execute("SELECT count(*) AS n FROM website_mitarbeiter WHERE kunden_id=%s", (kid,))
        if cur.fetchone()["n"] >= 30:
            raise HTTPException(400, "Mehr als 30 Mitarbeiter sind nicht möglich.")
        cur.execute("SELECT 1 FROM website_mitarbeiter WHERE kunden_id=%s AND schluessel=%s", (kid, k))
        if cur.fetchone():
            raise HTTPException(400, "Diesen Kurznamen gibt es schon. Bitte einen anderen wählen.")
        zeiten, ids, enc, geprueft, plan = _ma_pruefen(cur, kid, data)
        cur.execute("""INSERT INTO website_mitarbeiter (kunden_id, schluessel, name, aktiv, sort, wochentage, von, bis, zeiten, ical_enc, ical_geprueft_am, wochenplan)
                       VALUES (%s,%s,%s,%s,(SELECT coalesce(max(sort),0)+1 FROM website_mitarbeiter WHERE kunden_id=%s),%s,%s,%s,%s,%s,%s,%s::jsonb) RETURNING id""",
                    (kid, k, data.name.strip(), data.aktiv, kid, sorted(set(data.wochentage)) or None, data.von or None, data.bis or None, zeiten, enc, geprueft, json.dumps(plan) if plan else None))
        mid = cur.fetchone()["id"]
        _ma_leistungen(cur, mid, ids)
        return {"ok": True, "id": str(mid), "schluessel": k}


@kunde.put("/mitarbeiter/{mid}")
def k_ma_aendern(mid: str, data: MitarbeiterIn, user: dict = Depends(require_inhaber)):
    kid = user["kunden_id"]
    rate_limit("mapruef:" + str(kid), 30, 600)
    mid = _uuid(mid)
    with tx("system") as cur:
        cur.execute("SELECT schluessel FROM website_mitarbeiter WHERE id=%s AND kunden_id=%s", (mid, kid))
        alt = cur.fetchone()
        if not alt:
            raise HTTPException(404, "Diesen Mitarbeiter gibt es nicht.")
        k = _slug(data.schluessel) if data.schluessel.strip() else alt["schluessel"]
        cur.execute("SELECT 1 FROM website_mitarbeiter WHERE kunden_id=%s AND schluessel=%s AND id<>%s", (kid, k, mid))
        if cur.fetchone():
            raise HTTPException(400, "Diesen Kurznamen gibt es schon. Bitte einen anderen wählen.")
        zeiten, ids, enc, geprueft, plan = _ma_pruefen(cur, kid, data)
        cur.execute("""UPDATE website_mitarbeiter SET schluessel=%s, name=%s, aktiv=%s, wochentage=%s, von=%s, bis=%s, zeiten=%s,
                       ical_enc=coalesce(%s, ical_enc), ical_geprueft_am=coalesce(%s, ical_geprueft_am), wochenplan=%s::jsonb WHERE id=%s AND kunden_id=%s""",
                    (k, data.name.strip(), data.aktiv, sorted(set(data.wochentage)) or None, data.von or None, data.bis or None, zeiten, enc, geprueft, json.dumps(plan) if plan else None, mid, kid))
        _ma_leistungen(cur, mid, ids)
    sd._ical_cache.pop("m:" + mid, None)
    return {"ok": True}


@kunde.delete("/mitarbeiter/{mid}/kalender-adresse")
def k_ma_kal_weg(mid: str, user: dict = Depends(require_inhaber)):
    mid = _uuid(mid)
    with tx("system") as cur:
        cur.execute("UPDATE website_mitarbeiter SET ical_enc=NULL, ical_geprueft_am=NULL WHERE id=%s AND kunden_id=%s", (mid, user["kunden_id"]))
    sd._ical_cache.pop("m:" + mid, None)
    return {"ok": True}


@kunde.delete("/mitarbeiter/{mid}")
def k_ma_weg(mid: str, user: dict = Depends(require_inhaber)):
    with tx("system") as cur:
        cur.execute("DELETE FROM website_mitarbeiter WHERE id=%s AND kunden_id=%s", (_uuid(mid), user["kunden_id"]))
    return {"ok": True}


class StatIn(BaseModel):
    aktiv: bool


@kunde.put("/statistik-schalter")
def k_stat_schalter(data: StatIn, user: dict = Depends(require_inhaber)):
    with tx("system") as cur:
        sd.zeile(cur, user["kunden_id"])
        cur.execute("UPDATE website_dienste SET stat_aktiv=%s, geaendert_am=now() WHERE kunden_id=%s", (data.aktiv, user["kunden_id"]))
    return {"ok": True}


@kunde.get("/wacht")
def k_wacht(user: dict = Depends(require_inhaber)):
    with tx("system") as cur:
        return {"seiten": webbetrieb.status(cur, user["kunden_id"])}


@kunde.get("/domains")
def k_domains(user: dict = Depends(require_inhaber)):
    """Stand der Domain(s) des Kunden in einfachen Worten: wartet auf DNS (mit Anleitung), Zertifikat wird eingerichtet, live oder gestört."""
    with tx("system") as cur:
        cur.execute("""SELECT w.domain, w.letzte_veroeffentlichung, p.ok AS dns_ok, p.geprueft_am AS dns_am, p.ergebnis, t.ok AS seite_ok, t.seit, t.fehler
                       FROM websites w LEFT JOIN domain_pruefung p ON p.domain = w.domain LEFT JOIN website_wacht t ON t.domain = w.domain
                       WHERE w.kunden_id = %s AND w.art = 'live' AND w.status = 'live' ORDER BY w.domain""", (user["kunden_id"],))
        zeilen = cur.fetchall()
    out = []
    for z in zeilen:
        erg = z["ergebnis"] or {}
        if z["seite_ok"]:
            stand, titel, text = "live", "Deine Website ist online", "Deine Website ist unter dieser Adresse erreichbar und wird laufend überwacht."
        elif z["dns_ok"] is None:
            stand, titel, text = "neu", "Wird gerade eingerichtet", "Wir richten deine Website ein. Sobald sie veröffentlicht ist, siehst du hier den Stand."
        elif not z["dns_ok"]:
            stand, titel, text = "wartet_dns", "Wartet auf die Umstellung deiner Domain", "Deine Website ist fertig. Damit sie unter deiner Adresse erscheint, muss die Domain noch auf unseren Server zeigen. Die nötigen Schritte stehen unten, wir helfen gern dabei."
        elif z["fehler"] and "Zertifikat" in z["fehler"]:
            stand, titel, text = "zertifikat", "Sicheres HTTPS wird eingerichtet", "Die Domain zeigt auf unseren Server. Das Sicherheitszertifikat (HTTPS) wird gerade ausgestellt. Das dauert meist wenige Minuten."
        else:
            stand, titel, text = "gestoert", "Wir prüfen deine Website", "Die Website ist im Moment nicht erreichbar. Wir sind informiert und kümmern uns darum."
        out.append({"domain": z["domain"], "stand": stand, "titel": titel, "text": text, "dns_am": z["dns_am"], "seit": z["seit"] if stand == "live" else None,
                    "schritte": [x for x in (erg.get("schritte") or []) if not x.startswith("Danach")] if stand == "wartet_dns" else [],
                    "dns_anbieter": erg.get("dns_anbieter") if stand == "wartet_dns" else None,
                    "server_ipv4": (erg.get("server_ipv4") or []) if stand == "wartet_dns" else []})
    return {"domains": out}


@kunde.get("/statistik")
def k_statistik(tage: int = 30, user: dict = Depends(require_inhaber)):
    tage = max(1, min(tage, 365))
    kid = user["kunden_id"]
    with tx("system") as cur:
        q = lambda sql, *a: (cur.execute(sql, a), cur.fetchall())[1]  # noqa: E731
        ges = q("""SELECT count(*) FILTER (WHERE art='pv') AS aufrufe, count(DISTINCT tag::text || besucher) FILTER (WHERE art='pv') AS besucher,
                          count(*) FILTER (WHERE art='ev' AND name ILIKE 'WhatsApp%%') AS whatsapp, count(*) FILTER (WHERE art='ev' AND name ILIKE 'Anruf%%') AS anrufe
                   FROM website_stat WHERE kunden_id=%s AND tag > current_date - %s""", kid, tage)[0]
        cur.execute("SELECT count(*) AS n FROM anfragen WHERE kunden_id=%s AND quelle='website' AND erstellt_am > now() - make_interval(days => %s)", (kid, tage))
        anfragen = cur.fetchone()["n"]
        cur.execute("""SELECT count(*) FILTER (WHERE ergebnis='erschienen') AS da, count(*) FILTER (WHERE ergebnis='nicht_erschienen') AS weg, count(*) FILTER (WHERE status='gebucht') AS gebucht
                       FROM website_termine WHERE kunden_id=%s AND start > now() - make_interval(days => %s) AND start < now()""", (kid, tage))
        tm = cur.fetchone()
        erreich = webbetrieb.erreichbarkeit(cur, kid, date.today() - timedelta(days=tage), date.today() + timedelta(days=1))
        seiten = q("SELECT pfad, count(*) AS aufrufe FROM website_stat WHERE kunden_id=%s AND art='pv' AND tag > current_date - %s GROUP BY pfad ORDER BY 2 DESC LIMIT 10", kid, tage)
        quellen = q("""SELECT coalesce(quelle, 'Direkt / unbekannt') AS quelle, count(DISTINCT tag::text || besucher) AS besucher FROM website_stat
                       WHERE kunden_id=%s AND art='pv' AND tag > current_date - %s GROUP BY 1 ORDER BY 2 DESC LIMIT 8""", kid, tage)
        verlauf = q("SELECT tag, count(*) AS aufrufe FROM website_stat WHERE kunden_id=%s AND art='pv' AND tag > current_date - %s GROUP BY tag ORDER BY tag", kid, tage)
    return {"tage": tage, "besucher": ges["besucher"], "aufrufe": ges["aufrufe"], "whatsapp": ges["whatsapp"], "anrufe": ges["anrufe"], "anfragen": anfragen,
            "erschienen": tm["da"], "nicht_erschienen": tm["weg"], "termine_vergangen": tm["gebucht"], "erreichbarkeit": erreich,
            "seiten": seiten, "quellen": quellen, "verlauf": [{"tag": r["tag"].isoformat(), "aufrufe": r["aufrufe"]} for r in verlauf]}


@admin.get("/kunden/{kid}/website-dienste")
def a_status(kid: str, user: dict = Depends(require_admin)):
    try:
        with tx("admin") as cur:
            cur.execute("SELECT 1 FROM kunden WHERE id=%s", (kid,))
            if not cur.fetchone():
                raise HTTPException(404, "Kunde nicht gefunden.")
            d = _zustand(cur, kid)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(404, "Kunde nicht gefunden.")
    d["mail"]["smtp"].pop("benutzer", None)
    return d


class PlacesIn(BaseModel):
    api_key: str = Field(default="", max_length=300)


@admin.put("/einstellungen/places")
def a_places(data: PlacesIn, user: dict = Depends(require_admin)):
    if not config.PORTAL_SECRET:
        raise HTTPException(500, "PORTAL_SECRET fehlt in der .env. Bitte install.sh erneut ausführen.")
    with tx("admin") as cur:
        alt = einstellungen.lesen(cur, "places")
        key = data.api_key.strip() or alt.get("api_key", "")
        einstellungen.schreiben(cur, "places", {"api_key": key})
        audit(cur, user, "einstellungen_places", None, {"gesetzt": bool(key)})
    return {"ok": True}


@admin.get("/einstellungen/places")
def a_places_lesen(user: dict = Depends(require_admin)):
    with tx("admin") as cur:
        return {"key_gesetzt": bool(einstellungen.lesen(cur, "places").get("api_key"))}
